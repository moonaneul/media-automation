from __future__ import annotations

import argparse
from pathlib import Path

import yaml
from pptx import Presentation

from media_automation.bible import (
    BiblePassage,
    BiblePassageNotFoundError,
    BibleVerse,
    InMemoryBibleProvider,
)
from media_automation.planning import (
    BlockKind,
    build_wednesday_plan,
)
from media_automation.ppt import (
    FilePreServiceSlideProvider,
    MissingSongAssetError,
    build_presentation_file_from_plan,
    build_wednesday_structure,
    create_platform_slide_merger,
    load_song_asset_provider,
    validate_wednesday_structure,
)
from media_automation.weekly_data.models import (
    WednesdayData,
    parse_weekly_data,
)


def load_bible_provider(
    path: Path,
) -> InMemoryBibleProvider:
    """
    이번 주에 사용할 개역개정 본문 YAML을 읽는다.

    형식:

    passages:
      "삼상 16:6-7":
        reference: "삼상 16:6~7"
        verses:
          6: "본문..."
          7: "본문..."
    """

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:
        raw = yaml.safe_load(file)

    if not isinstance(raw, dict):
        raise ValueError(
            "Bible YAML의 최상위 값은 객체여야 합니다."
        )

    passages_raw = raw.get("passages")

    if not isinstance(
        passages_raw,
        dict,
    ):
        raise ValueError(
            "Bible YAML에 passages 객체가 필요합니다."
        )

    passages: dict[
        str,
        BiblePassage,
    ] = {}

    for lookup_reference, data in passages_raw.items():
        if not isinstance(
            lookup_reference,
            str,
        ):
            raise ValueError(
                "Bible passage key는 문자열이어야 합니다."
            )

        if not isinstance(
            data,
            dict,
        ):
            raise ValueError(
                f"{lookup_reference}의 데이터 형식이 "
                "잘못되었습니다."
            )

        display_reference = data.get(
            "reference"
        )

        if not isinstance(
            display_reference,
            str,
        ):
            raise ValueError(
                f"{lookup_reference}의 reference가 "
                "필요합니다."
            )

        verses_raw = data.get(
            "verses"
        )

        if not isinstance(
            verses_raw,
            dict,
        ):
            raise ValueError(
                f"{lookup_reference}의 verses가 "
                "필요합니다."
            )

        verses: list[BibleVerse] = []

        for number_raw, text in verses_raw.items():
            try:
                number = int(
                    number_raw
                )
            except (
                TypeError,
                ValueError,
            ) as error:
                raise ValueError(
                    f"{lookup_reference}의 절 번호가 "
                    f"잘못되었습니다: {number_raw}"
                ) from error

            if not isinstance(
                text,
                str,
            ):
                raise ValueError(
                    f"{lookup_reference} {number}절 "
                    "본문은 문자열이어야 합니다."
                )

            verses.append(
                BibleVerse(
                    number=number,
                    text=text,
                )
            )

        verses.sort(
            key=lambda verse: verse.number
        )

        passages[
            lookup_reference
        ] = BiblePassage(
            reference=display_reference,
            verses=tuple(verses),
        )

    return InMemoryBibleProvider(
        passages
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "하늘빛기쁨성서침례교회 "
            "수요예배 PPT 제작"
        )
    )

    parser.add_argument(
        "--weekly",
        type=Path,
        required=True,
        help="이번 주 수요예배 YAML",
    )

    parser.add_argument(
        "--songs",
        type=Path,
        required=True,
        help="검수된 악보 PPT manifest YAML",
    )

    parser.add_argument(
        "--bible",
        type=Path,
        required=True,
        help="이번 주 개역개정 본문 YAML",
    )

    parser.add_argument(
        "--source",
        type=Path,
        required=True,
        help=(
            "예배 전 안내 1~4장을 가져올 "
            "기존 수요예배 PPT"
        ),
    )

    parser.add_argument(
        "--output",
        type=Path,
        required=True,
        help="생성할 수요예배 PPT 경로",
    )
    parser.add_argument(
        "--validate-only",
        action="store_true",
        help=(
            "PPT를 생성하지 않고 "
            "Weekly Data, 성경 본문, 찬양 자료를 검증합니다."
        ),
    )

    parser.add_argument(
        "--structure-only",
        action="store_true",
        help=(
            "Mac에서도 확인 가능한 수요예배 구조 PPT를 생성합니다. "
            "찬양은 실제 악보 PPT의 장수만 반영한 "
            "빈 placeholder로 표시합니다."
        ),
    )

    args = parser.parse_args()

    if (
        args.validate_only
        and args.structure_only
    ):
        parser.error(
            "--validate-only와 --structure-only는 "
            "동시에 사용할 수 없습니다."
        )

    # -------------------------
    # 입력 파일 존재 확인
    # -------------------------

    for name, path in {
        "weekly": args.weekly,
        "songs": args.songs,
        "bible": args.bible,
        "source": args.source,
    }.items():
        if not path.exists():
            raise FileNotFoundError(
                f"{name} 파일을 찾을 수 없습니다: "
                f"{path}"
            )

    # -------------------------
    # Weekly Data
    # -------------------------

    with args.weekly.open(
        "r",
        encoding="utf-8",
    ) as file:
        weekly_raw = yaml.safe_load(
            file
        )

    weekly = parse_weekly_data(
        weekly_raw
    )

    if not isinstance(
        weekly,
        WednesdayData,
    ):
        raise TypeError(
            "--weekly에는 수요예배 데이터가 "
            "필요합니다."
        )

    plan = build_wednesday_plan(
        weekly
    )

    # -------------------------
    # Providers
    # -------------------------

    song_provider = (
        load_song_asset_provider(
            args.songs
        )
    )

    bible_provider = (
        load_bible_provider(
            args.bible
        )
    )

    pre_service_provider = (
        FilePreServiceSlideProvider(
            {
                "pre_service": (
                    args.source
                ),
            },
            slide_counts={
                "pre_service": 4,
            },
        )
    )


    # -------------------------
    # 찬양 자료 사전 확인
    # -------------------------


    missing_songs: list[str] = []

    for block in plan:
        if block.kind != BlockKind.SONG:
            continue

        title = block.value.title

        try:
            song_provider.get_song_asset(
                title
            )
        except MissingSongAssetError:
            missing_songs.append(
                title
            )
    missing_bible_passages: list[str] = []

    for block in plan:
        if block.kind != BlockKind.SCRIPTURE:
            continue

        reference = block.value.reference

        try:
            bible_provider.get_passage(
                reference
            )
        except BiblePassageNotFoundError:
            missing_bible_passages.append(
                reference
            )

    print()
    print("=== Wednesday Build ===")

    print(
        f"weekly : {args.weekly}"
    )
    print(
        f"songs  : {args.songs}"
    )
    print(
        f"bible  : {args.bible}"
    )
    print(
        f"source : {args.source}"
    )
    print(
        f"output : {args.output}"
    )

    if missing_songs:
        print()
        print(
            "=== Missing Song Assets ==="
        )

        for title in missing_songs:
            print(
                f"- {title}"
            )

        print(
            "누락된 찬양은 PPT에 "
            "임의 대체하지 않고 생략합니다."
        )
    if missing_bible_passages:
        print()
        print(
            "=== Missing Bible Passages ==="
        )

        for reference in missing_bible_passages:
            print(
                f"- {reference}"
            )

        raise RuntimeError(
            "필요한 개역개정 성경 본문이 "
            "Bible YAML에 없습니다."
        )

    if args.validate_only:
        print()
        print("=== Validation Result ===")
        print(
            f"plan blocks : {len(plan)}"
        )
        print(
            f"missing songs : {len(missing_songs)}"
        )
        print(
            "bible passages : OK"
        )

        print()
        print(
            "입력 검증 완료. "
            "PPT는 생성하지 않았습니다."
        )

        return
    if args.structure_only:
        result = build_wednesday_structure(
            plan,
            args.output,
            bible_provider=bible_provider,
            pre_service_provider=(
                pre_service_provider
            ),
            song_asset_provider=(
                song_provider
            ),
            skip_missing_songs=True,
        )

        reopened = Presentation(
            result.output_path
        )
        scripture_passages = {}

        for block in plan:
            if block.kind != BlockKind.SCRIPTURE:
                continue

            scripture_passages[
                block.key
            ] = bible_provider.get_passage(
                block.value.reference
            )
        qa_result = validate_wednesday_structure(
            result.output_path,
            slide_ranges=result.slide_ranges,
            scripture_passages=(
                scripture_passages
            ),
        )
        print()
        print(
            "=== Wednesday Structure Preview ==="
        )
        print(
            f"slides : {len(reopened.slides)}"
        )
        print(
            f"output : {result.output_path}"
        )

        print()
        print(
            "=== Slide Ranges ==="
        )

        for key, slide_range in sorted(
            result.slide_ranges.items(),
            key=lambda item: item[1].start,
        ):
            if (
                slide_range.start
                == slide_range.end
            ):
                location = str(
                    slide_range.start
                )
            else:
                location = (
                    f"{slide_range.start}"
                    f"-{slide_range.end}"
                )

            print(
                f"{location:>7} | {key}"
            )

        print()
        print("=== Wednesday QA ===")

        if qa_result.ok:
            print("QA PASSED")
            print("- aspect ratio : OK")
            print("- transition blanks : OK")
            print("- baptism term : OK")
            print("- scripture slides : OK")
            print("- standalone alignment : OK")
            print("- scripture layout : OK")
            print("- blue background shapes : OK")
            print("- hyperlinks / URLs : OK")
        else:
            print(
                f"QA FAILED : {len(qa_result.issues)} issue(s)"
            )

            for issue in qa_result.issues:
                location = (
                    f"slide {issue.slide_number}"
                    if issue.slide_number is not None
                    else "presentation"
                )

                print(
                    f"- [{issue.code}] "
                    f"{location} | "
                    f"{issue.message}"
                )

            raise RuntimeError(
                "수요예배 구조 QA에서 오류가 발견되었습니다."
            )
        print()
        print(
            "수요예배 구조 프리뷰 생성 완료."
        )
        print(
            "찬양 화면은 실제 악보가 아니라 "
            "악보 PPT의 장수만 반영한 빈 화면입니다."
        )
        print(
            "최종 악보 병합은 Windows PowerPoint "
            "제작 환경에서 별도 수행합니다."
        )

        return
    merger = (
        create_platform_slide_merger()
    )
    # -------------------------
    # 실제 PPT 생성
    # -------------------------

    try:
        result = (
            build_presentation_file_from_plan(
                plan,
                args.output,
                bible_provider=(
                    bible_provider
                ),
                pre_service_provider=(
                    pre_service_provider
                ),
                song_asset_provider=(
                    song_provider
                ),
                slide_merger=merger,
                skip_missing_songs=True,
            )
        )
    except PermissionError as error:
        raise RuntimeError(
            f"{args.output} 파일을 저장할 수 없습니다. "
            "PowerPoint에서 결과 파일이 열려 있다면 "
            "닫은 뒤 다시 실행해주세요."
        ) from error

    # -------------------------
    # 결과 기본 검사
    # -------------------------

    reopened = Presentation(
        result
    )

    print()
    print("=== Result ===")
    print(
        f"slides : {len(reopened.slides)}"
    )
    print(
        f"output : {result}"
    )

    print()
    print(
        "PPT 생성 완료."
    )
    print(
        "실제 PowerPoint에서 "
        "악보·본문·빈 화면·잘림을 "
        "최종 확인해주세요."
    )

if __name__ == "__main__":
    main()