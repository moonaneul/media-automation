"""Read-only compare a known song PPTX with its logical slides in the generated deck.

No files are modified or uploaded. Output contains only structural counts and hashes,
not church text or artwork. Uses presentation order (not slideN.xml numbering).
"""
from __future__ import annotations

import argparse
import hashlib
from collections import Counter
from pathlib import Path

from pptx import Presentation


def digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()[:16]


def image_hashes(part) -> tuple[str, ...]:
    """Collect images linked to this part, independent of relationship IDs/names."""
    return tuple(sorted(digest(rel.target_part.blob)
                        for rel in part.rels.values()
                        if not rel.is_external and rel.reltype.endswith("/image")))


def fingerprint(slide) -> dict:
    shapes = tuple(slide.shapes)
    native_text = "\n".join(shape.text for shape in shapes if shape.has_text_frame)
    geometry = tuple((int(shape.shape_type), shape.left, shape.top,
                      shape.width, shape.height) for shape in shapes)
    layout = slide.slide_layout
    master = layout.slide_master
    return {
        "shapes": len(shapes),
        "geometry": digest(repr(geometry).encode()),
        "text": digest(native_text.encode()),
        "slide_images": image_hashes(slide.part),
        "layout_images": image_hashes(layout.part),
        "master_images": image_hashes(master.part),
    }


def compare(source: Path, generated: Path, start: int) -> list[dict]:
    original = Presentation(str(source))
    produced = Presentation(str(generated))
    if (original.slide_width, original.slide_height) != (
            produced.slide_width, produced.slide_height):
        raise ValueError("원본과 생성본의 화면 크기가 다릅니다.")
    end = start + len(original.slides) - 1
    if start < 1 or end > len(produced.slides):
        raise ValueError("생성본의 시작 슬라이드 범위를 확인하세요.")
    rows = []
    keys = ("shapes", "geometry", "text", "slide_images",
            "layout_images", "master_images")
    for index, slide in enumerate(original.slides):
        target = produced.slides[start - 1 + index]
        source_fp, output_fp = fingerprint(slide), fingerprint(target)
        rows.append({
            "source_slide": index + 1,
            "output_slide": start + index,
            "same": {key: source_fp[key] == output_fp[key] for key in keys},
            "source_images": len(source_fp["slide_images"]),
            "output_images": len(output_fp["slide_images"]),
            "shared_images": sum((Counter(source_fp["slide_images"]) &
                                  Counter(output_fp["slide_images"])).values()),
            "source_shapes": source_fp["shapes"],
            "output_shapes": output_fp["shapes"],
        })
    print(f"source_slides={len(original.slides)} output_slides={len(produced.slides)} "
          f"compared_output={start}-{end}")
    print("모든 비교는 읽기 전용입니다. false는 표시가 다를 가능성이지 원인 확정은 아닙니다.")
    for row in rows:
        changed = [key for key, same in row["same"].items() if not same]
        print(f"source {row['source_slide']:2} -> output {row['output_slide']:3}: "
              f"{'CORE-MATCH' if not any(k in changed for k in ('shapes','geometry','text','slide_images')) else 'CORE-DIFF'} "
              f"differ={','.join(changed) if changed else '-'} "
              f"shapes={row['source_shapes']}/{row['output_shapes']} "
              f"direct_images={row['source_images']}/{row['output_images']} "
              f"shared_images={row['shared_images']}")
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("original", type=Path, help="정상 찬양 원본 PPTX")
    parser.add_argument("generated", type=Path, help="전체 예배 생성 PPTX")
    parser.add_argument("--start", type=int, default=6, help="생성본에서 찬양 첫 슬라이드 (기본 6)")
    args = parser.parse_args()
    compare(args.original, args.generated, args.start)


if __name__ == "__main__":
    main()
