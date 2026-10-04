from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from media_automation.ppt import (
    create_4x3_presentation,
    create_platform_slide_merger,
)


def main() -> None:
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "source",
        type=Path,
        help="기존 수요예배 PPT",
    )

    args = parser.parse_args()

    source = args.source.resolve()

    if not source.exists():
        raise FileNotFoundError(
            f"원본 PPT를 찾을 수 없습니다: {source}"
        )

    output = Path(
        "output/mac_slide_merge_check.pptx"
    )

    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    # -------------------------
    # 병합 대상 PPT 준비
    # -------------------------

    prs = create_4x3_presentation()

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    box = slide.shapes.add_textbox(
        Inches(1),
        Inches(2.5),
        Inches(8),
        Inches(2),
    )

    paragraph = (
        box.text_frame
        .paragraphs[0]
    )

    paragraph.text = (
        "MAC SLIDE MERGE TEST"
    )

    paragraph.runs[0].font.size = Pt(32)

    prs.save(
        output
    )

    # -------------------------
    # 실제 악보 22~26장 삽입
    # -------------------------

    merger = (
        create_platform_slide_merger()
    )

    print(
        "merger:",
        type(merger).__name__,
    )

    merger.insert_range(
        output,
        source,
        after_slide=1,
        start_slide=22,
        end_slide=26,
    )

    # -------------------------
    # 결과 확인
    # -------------------------

    reopened = Presentation(
        output
    )

    print()
    print("=== Result ===")
    print(
        "output:",
        output,
    )
    print(
        "slides:",
        len(reopened.slides),
    )

    for index, slide in enumerate(
        reopened.slides,
        start=1,
    ):
        texts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    texts.append(text)

        summary = " | ".join(
            texts
        )

        print(
            f"{index:02d}: {summary}"
        )

    print()
    print(
        "OK: Mac slide merge POC finished"
    )


if __name__ == "__main__":
    main()