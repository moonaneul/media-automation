from pathlib import Path

from pptx import Presentation
from pptx_slide_copier import SlideCopier


SOURCE = Path("/Users/moonaneul/Downloads/수요예배.pptx")
OUTPUT = Path("output/python_slide_copy_blank_check.pptx")

def get_slide_text(slide) -> str:
    texts = []

    for shape in slide.shapes:
        if hasattr(shape, "text"):
            text = shape.text.strip()

            if text:
                texts.append(text.replace("\n", " / "))

    return " | ".join(texts)


def remove_all_slides(prs: Presentation) -> None:
    while len(prs.slides) > 0:
        slide_id = prs.slides._sldIdLst[0]
        r_id = slide_id.rId

        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[0]


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(
            f"원본 PPT가 없습니다: {SOURCE}"
        )

    OUTPUT.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    source = Presentation(SOURCE)

    print("source slides:", len(source.slides))
    print()
    print("=== source 22~26 ===")

    for slide_number in range(22, 27):
        slide = source.slides[slide_number - 1]

        print(
            f"{slide_number}:",
            get_slide_text(slide),
        )

    target = Presentation()

    print()
    print("=== copying ===")

    for slide_number in range(22, 27):
        source_index = slide_number - 1

        print("copying:", slide_number)

        SlideCopier.copy_slide(
            source,
            source_index,
            target,
        )

    target.save(OUTPUT)

    reopened = Presentation(OUTPUT)

    print()
    print("=== result ===")
    print("output:", OUTPUT)
    print("slides:", len(reopened.slides))

    for i, slide in enumerate(
        reopened.slides,
        start=1,
    ):
        print(
            f"{i}:",
            get_slide_text(slide),
        )


if __name__ == "__main__":
    main()