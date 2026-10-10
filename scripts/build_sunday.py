from __future__ import annotations

import argparse
from pathlib import Path

import yaml
from pptx import Presentation

from media_automation.ppt.hyperlinks import remove_presentation_hyperlinks

from media_automation.bible import (
    BiblePassage,
    BiblePassageNotFoundError,
    BibleVerse,
    InMemoryBibleProvider,
)
from media_automation.planning import (
    BlockKind,
    build_sunday_plan,
)
from media_automation.ppt import (
    FilePreServiceSlideProvider,
    MissingSongAssetError,
    build_presentation_file_from_plan,
    create_platform_slide_merger,
    load_song_asset_provider,
    build_in_person_structure,
    validate_in_person_structure,
)
from media_automation.weekly_data.models import (
    SundayData,
    parse_weekly_data,
)


def load_bible_provider(
    path: Path,
) -> InMemoryBibleProvider:
    with path.open(
        "r",
        encoding="utf-8",
    ) as file:
        raw = yaml.safe_load(file)

    if not isinstance(raw, dict):
        raise ValueError(
            "Bible YAML의 최상위 값은 객체여야 합니다."
        )

    passages_raw = raw.get("passages")

    if not isinstance(passages_raw, dict):
        raise ValueError(
            "Bible YAML에 passages 객체가 필요합니다."
        )

    passages: dict[str, BiblePassage] = {}

    for lookup_reference, data in passages_raw.items():
        if not isinstance(lookup_reference, str):
            raise ValueError(
                "Bible passage key는 문자열이어야 합니다."
            )

        if not isinstance(data, dict):
            raise ValueError(
                f"{lookup_reference}의 데이터 형식이 잘못되었습니다."
            )

        display_reference = data.get("reference")
        verses_raw = data.get("verses")

        if not isinstance(display_reference, str):
            raise ValueError(
                f"{lookup_reference}의 reference가 필요합니다."
            )

        if not isinstance(verses_raw, dict):
            raise ValueError(
                f"{lookup_reference}의 verses가 필요합니다."
            )

        verses: list[BibleVerse] = []

        for number_raw, text in verses_raw.items():
            try:
                number = int(number_raw)
            except (TypeError, ValueError) as error:
                raise ValueError(
                    f"{lookup_reference}의 절 번호가 "
                    f"잘못되었습니다: {number_raw}"
                ) from error

            if not isinstance(text, str):
                raise ValueError(
                    f"{lookup_reference} {number}절 "
                    "본문은 문자열이어야 합니다."
                )

            verses.append(
                BibleVerse(
                    number=number,
                    text=text,
                    end_number=(data.get("verse_ends", {}).get(number)
                                or data.get("verse_ends", {}).get(str(number))),
                )
            )

        verses.sort(
            key=lambda verse: verse.number
        )

        passages[lookup_reference] = (
            BiblePassage(
                reference=display_reference,
                verses=tuple(verses),
            )
        )

    return InMemoryBibleProvider(passages)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "하늘빛기쁨성서침례교회 "
            "주일 오전 2부 PPT 제작"
        )
    )

    parser.add_argument(
        "--weekly",
        type=Path,
        required=True,
        help="이번 주 주일예배 YAML",
    )

    parser.add_argument(
        "--songs",
        type=Path,
        required=True,
        help="검수된 악보 PPT manifest YAML",
    )

    parser.add_argument(
        "--bible",
        type=Path,
        required=True,
        help="이번 주 개역개정 본문 YAML",
    )

    parser.add_argument(
        "--source",
        type=Path,
        help="예배 전 안내를 가져올 기존 주일예배 PPT",
    )

    parser.add_argument(
        "--pre-service-slides",
        type=int,
        help="기존 PPT에서 가져올 예배 전 안내 슬라이드 수",
    )

    parser.add_argument(
        "--output",
        type=Path,
        help="생성할 주일예배 PPT 경로",
    )

    parser.add_argument(
        "--validate-only",
        action="store_true",
        help=(
            "PPT를 생성하지 않고 "
            "Weekly Data, 성경 본문, 찬양 자료를 검증합니다."
        ),
    )
    parser.add_argument(
        "--structure-only",
        action="store_true",
        help=(
            "Mac에서도 확인 가능한 "
            "주일예배 구조 프리뷰를 생성합니다."
        ),
    )
    args = parser.parse_args()

    if (
        args.validate_only
        and args.structure_only
    ):
        raise ValueError(
            "--validate-only와 "
            "--structure-only는 함께 사용할 수 없습니다."
        )
    # 항상 필요한 입력
    for name, path in {
        "weekly": args.weekly,
        "songs": args.songs,
        "bible": args.bible,
    }.items():
        if not path.exists():
            raise FileNotFoundError(
                f"{name} 파일을 찾을 수 없습니다: {path}"
            )

    # Weekly Data
    with args.weekly.open(
        "r",
        encoding="utf-8",
    ) as file:
        weekly_raw = yaml.safe_load(file)

    weekly = parse_weekly_data(weekly_raw)

    if not isinstance(weekly, SundayData):
        raise TypeError(
            "--weekly에는 주일예배 데이터가 필요합니다."
        )

    plan = build_sunday_plan(weekly)

    # Providers
    song_provider = load_song_asset_provider(
        args.songs
    )

    bible_provider = load_bible_provider(
        args.bible
    )

    # 찬양 자료 확인
    missing_songs: list[str] = []

    for block in plan:
        if block.kind != BlockKind.SONG:
            continue

        title = block.value.title

        try:
            song_provider.get_song_asset(title)
        except MissingSongAssetError:
            missing_songs.append(title)

    # 성경 본문 확인
    missing_bible_passages: list[str] = []

    for block in plan:
        if block.kind != BlockKind.SCRIPTURE:
            continue

        reference = block.value.reference

        try:
            bible_provider.get_passage(reference)
        except BiblePassageNotFoundError:
            missing_bible_passages.append(reference)

    print()
    print("=== Sunday Build ===")
    print(f"weekly : {args.weekly}")
    print(f"songs  : {args.songs}")
    print(f"bible  : {args.bible}")

    if missing_songs:
        print()
        print("=== Missing Song Assets ===")

        for title in missing_songs:
            print(f"- {title}")

        print(
            "누락된 찬양은 PPT에 "
            "임의 대체하지 않고 생략합니다."
        )

    if missing_bible_passages:
        print()
        print("=== Missing Bible Passages ===")

        for reference in missing_bible_passages:
            print(f"- {reference}")

        raise RuntimeError(
            "필요한 개역개정 성경 본문이 "
            "Bible YAML에 없습니다."
        )

    # Mac에서도 여기까지 검증 가능
    if args.validate_only:
        print()
        print("=== Validation Result ===")
        print(f"plan blocks : {len(plan)}")
        print(f"missing songs : {len(missing_songs)}")
        print("bible passages : OK")

        print()
        print(
            "입력 검증 완료. "
            "PPT는 생성하지 않았습니다."
        )

        return

    # 실제 PPT 제작에는 원본/출력 정보 필요
    if args.source is None:
        raise ValueError(
            "실제 PPT 생성에는 --source가 필요합니다."
        )

    if args.output is None:
        raise ValueError(
            "실제 PPT 생성에는 --output이 필요합니다."
        )

    if args.pre_service_slides is None:
        raise ValueError(
            "실제 PPT 생성에는 "
            "--pre-service-slides가 필요합니다."
        )

    if args.pre_service_slides < 1:
        raise ValueError(
            "--pre-service-slides는 1 이상이어야 합니다."
        )

    if not args.source.exists():
        raise FileNotFoundError(
            f"source 파일을 찾을 수 없습니다: "
            f"{args.source}"
        )

    pre_service_provider = (
        FilePreServiceSlideProvider(
            {
                "pre_service": args.source,
            },
            slide_counts={
                "pre_service": (
                    args.pre_service_slides
                ),
            },
        )
    )
    
    if args.structure_only:
        result = build_in_person_structure(
            plan,
            args.output,
            bible_provider=bible_provider,
            pre_service_provider=(
                pre_service_provider
            ),
            song_asset_provider=(
                song_provider
            ),
            skip_missing_songs=True,
            preserve_blank_after_pre_service=True,
            preserve_trailing_blank=True,
        )

        reopened = Presentation(
            result.output_path
        )
        scripture_passages = {}

        for block in plan:
            if block.kind != BlockKind.SCRIPTURE:
                continue

            scripture_passages[
                block.key
            ] = bible_provider.get_passage(
                block.value.reference
            )

        qa = validate_in_person_structure(
            result.output_path,
            slide_ranges=result.slide_ranges,
            scripture_passages=(
                scripture_passages
            ),
            standalone_keys=(
                (
                    "serving.this_week."
                    "second_service.prayer"
                ),
                (
                    "serving.this_week."
                    "second_service.offering_prayer"
                ),
                "worship.sermon_title",
            ),
            generated_keys=(
                (
                    "serving.this_week."
                    "second_service.prayer"
                ),
                (
                    "serving.this_week."
                    "second_service.offering_prayer"
                ),
                "worship.sermon_title",
                "worship.scripture",
                "worship.additional_scripture",
            ),
        )
        print()
        print(
            "=== Sunday Structure Preview ==="
        )
        print(
            f"slides : {len(reopened.slides)}"
        )

        for key, slide_range in (
            result.slide_ranges.items()
        ):
            print(
                f"{slide_range.start}-"
                f"{slide_range.end} "
                f"{key}"
            )

        print()
        print("=== Sunday QA ===")

        if qa.ok:
            print("PASSED")
        else:
            for issue in qa.issues:
                location = (
                    f"slide {issue.slide_number}"
                    if issue.slide_number
                    is not None
                    else "presentation"
                )

                print(
                    f"- [{issue.code}] "
                    f"{location}: "
                    f"{issue.message}"
                )

            raise RuntimeError(
                "Sunday structure QA에 실패했습니다."
            )
        print()
        print(
            f"output : {result.output_path}"
        )

        return
    # Windows는 PowerPoint COM, macOS/Linux는 Open XML 원본 병합
    merger = create_platform_slide_merger()

    try:
        result = (
            build_presentation_file_from_plan(
                plan,
                args.output,
                bible_provider=bible_provider,
                pre_service_provider=(
                    pre_service_provider
                ),
                song_asset_provider=(
                    song_provider
                ),
                slide_merger=merger,
                skip_missing_songs=True,
                preserve_blank_after_pre_service=True,
                preserve_trailing_blank=True,
            )
        )
    except PermissionError as error:
        raise RuntimeError(
            f"{args.output} 파일을 저장할 수 없습니다. "
            "PowerPoint에서 결과 파일이 열려 있다면 "
            "닫은 뒤 다시 실행해주세요."
        ) from error

    removed_links = remove_presentation_hyperlinks(result)
    structure = build_in_person_structure(
        plan, args.output.with_name(args.output.stem + "_structure_check.pptx"),
        bible_provider=bible_provider, pre_service_provider=pre_service_provider,
        song_asset_provider=song_provider, skip_missing_songs=True,
        preserve_blank_after_pre_service=True, preserve_trailing_blank=True,
    )
    scripture_passages = {
        block.key: bible_provider.get_passage(block.value.reference)
        for block in plan if block.kind == BlockKind.SCRIPTURE
    }
    final_qa = validate_in_person_structure(
        result, slide_ranges=structure.slide_ranges,
        scripture_passages=scripture_passages,
        standalone_keys=(
            "serving.this_week.second_service.prayer",
            "serving.this_week.second_service.offering_prayer",
            "worship.sermon_title",
        ),
        generated_keys=(
            "serving.this_week.second_service.prayer",
            "serving.this_week.second_service.offering_prayer",
            "worship.sermon_title", "worship.scripture", "worship.additional_scripture",
        ),
    )
    if not final_qa.ok:
        for issue in final_qa.issues:
            print(f"FINAL QA: [{issue.code}] {issue.message}")
        raise RuntimeError("주일 최종 PPT QA 실패")
    print(f"Sunday Final QA: PASS (links removed: {removed_links})")

    reopened = Presentation(result)

    print()
    print("=== Result ===")
    print(
        f"slides : {len(reopened.slides)}"
    )
    print(
        f"output : {result}"
    )

    print()
    print("PPT 생성 완료.")
    print(
        "실제 PowerPoint에서 "
        "악보·본문·담당자·빈 화면·잘림을 "
        "최종 확인해주세요."
    )


if __name__ == "__main__":
    main()