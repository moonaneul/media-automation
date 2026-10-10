import hashlib

import pytest
from pptx import Presentation
from reportlab.pdfgen import canvas

from media_automation.web.existing_inspections import ExistingInspections


def test_existing_pptx_isolated_and_reviewed(tmp_path):
    root = tmp_path / "past-review"
    lib = ExistingInspections(root)
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    original = tmp_path / "source.pptx"
    presentation.save(original)
    original_digest = hashlib.sha256(original.read_bytes()).hexdigest()
    record = lib.create("sunday", "2026-09-27", "source.pptx", original.read_bytes())
    inspection_id = record["inspection_id"]
    assert record["type"] == "existing_file_only"
    assert not record["qa"]["available"]
    assert len(lib.list()) == 1
    result = lib.inspect(inspection_id)
    assert result["available"]
    labels = result["check_labels"]
    assert labels
    checks = {key: "checked" for key in labels}
    changed = lib.confirm(inspection_id, {
        "reviewer": "검수 담당", "note": "별도 화면 검수", "sha256": result["sha256"],
        "checks": checks,
    })
    assert changed["review_record_complete"]
    assert lib.report(inspection_id)
    assert original_digest == hashlib.sha256(original.read_bytes()).hexdigest()
    with pytest.raises(ValueError):
        lib.confirm(inspection_id, {"reviewer": "x", "sha256": "wrong", "checks": checks})


def test_existing_pdf_preview_without_mutating_original(tmp_path):
    lib = ExistingInspections(tmp_path / "review")
    source = tmp_path / "bulletin.pdf"
    c = canvas.Canvas(str(source), pagesize=(841.89, 595.28))
    c.drawString(12, 40, "sample")
    c.showPage()
    c.drawString(12, 40, "sample")
    c.save()
    original = source.read_bytes()
    data = lib.create("bulletin", "2026-09-27", "bulletin.pdf", original)
    inspection_id = data["inspection_id"]
    assert data["preview"]["available"]
    assert lib.preview_bytes(inspection_id) == original
    assert source.read_bytes() == original


def test_rejects_bad_uploads(tmp_path):
    lib = ExistingInspections(tmp_path / "review")
    for filename in ("../escape.pptx", "inside/file.pptx", "song.pdf"):
        with pytest.raises(ValueError):
            lib.create("sunday", "2026-09-27", filename, b"fake")
    with pytest.raises(ValueError):
        lib.create("bulletin", "2026-09-27", "bad.pptx", b"fake")
    with pytest.raises(ValueError):
        lib.create("sunday", "2026-09-27", "bad.pptx", b"")
    with pytest.raises(ValueError):
        lib.get("../test")


def test_changed_copy_cannot_be_used_as_current_review(tmp_path):
    lib = ExistingInspections(tmp_path / "review")
    original = tmp_path / "test.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(original)
    record = lib.create("sunday", "2026-09-27", "test.pptx", original.read_bytes())
    p = lib._dir(record["inspection_id"]) / "source.pptx"
    p.write_bytes(b"changed")
    assert not lib.get(record["inspection_id"])["unchanged"]
    with pytest.raises(ValueError):
        lib.inspect(record["inspection_id"])
