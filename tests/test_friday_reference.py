from pptx import Presentation
from pptx.util import Inches,Pt
from media_automation.ppt.friday_reference import replace_paragraph,replace_prayer

def test_prayer_body_has_one_marker_regardless_of_input_bullets():
    from media_automation.ppt.friday_reference import prayer_body_text
    expected='- 내 구주 예수를 더욱 사랑\n- 두 번째 기도'
    for body in ['내 구주 예수를 더욱 사랑\n두 번째 기도', '- 내 구주 예수를 더욱 사랑\n- 두 번째 기도', '- - 내 구주 예수를 더욱 사랑\n• 두 번째 기도']:
        assert prayer_body_text(body)==expected

def test_in_person_prayer_is_left_aligned_without_changing_other_slides(tmp_path):
    from media_automation.ppt.friday_typography import apply_friday_typography
    from pptx.enum.text import PP_ALIGN
    p=Presentation()
    for text in ['기도 제목\n기도 내용','설교 제목']:
        s=p.slides.add_slide(p.slide_layouts[6]);shape=s.shapes.add_textbox(0,0,Inches(8),Inches(3))
        shape.text=text
        for para in shape.text_frame.paragraphs:para.alignment=PP_ALIGN.CENTER
    path=tmp_path/'in-person.pptx';p.save(path)
    apply_friday_typography(path,[dict(status='rendered',kind='input',key='말씀 관련 기도 제목 1',start=1,end=1),dict(status='rendered',kind='input',key='설교 제목',start=2,end=2)],zoom=False)
    result=Presentation(path)
    assert all(para.alignment==PP_ALIGN.LEFT for para in result.slides[0].shapes[0].text_frame.paragraphs)
    assert result.slides[1].shapes[0].text_frame.paragraphs[0].alignment==PP_ALIGN.CENTER

def test_replacement_keeps_source_geometry_font_and_explicit_soft_breaks():
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);shape=s.shapes.add_textbox(Inches(1),Inches(2),Inches(8),Inches(2))
    para=shape.text_frame.paragraphs[0];para.text='과거 안내';para.runs[0].font.name='나눔스퀘어라운드 Bold';para.runs[0].font.size=Pt(30)
    geometry=(shape.left,shape.top,shape.width,shape.height)
    replace_paragraph(para,'이번 주 안내\v다음 줄')
    assert para.text=='이번 주 안내\v다음 줄'
    assert all(r.font.name=='나눔스퀘어라운드 Bold' and r.font.size==Pt(30) for r in para.runs)
    assert geometry==(shape.left,shape.top,shape.width,shape.height)
    assert '과거' not in shape.text

def test_prayer_uses_source_body_style_and_original_separator():
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);shape=s.shapes.add_textbox(0,0,Inches(10),Inches(3))
    para=shape.text_frame.paragraphs[0];a=para.add_run();a.text='옛 제목';a.font.size=Pt(36);a.font.name='제목 원본'
    b=para.add_run();b.text='옛 내용';b.font.size=Pt(32);b.font.name='내용 원본'
    replace_prayer(shape,'이번 제목','이번 내용')
    assert shape.text=='이번 제목\v\v\v  - 이번 내용'
    assert shape.text_frame.paragraphs[0].runs[-1].font.name=='내용 원본'
    assert '옛' not in shape.text

def test_long_prayer_stays_inside_safe_width_and_preserves_words():
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);shape=s.shapes.add_textbox(0,0,Inches(12),Inches(3))
    para=shape.text_frame.paragraphs[0];a=para.add_run();a.text='제목';a.font.size=Pt(36)
    b=para.add_run();b.text='내용';b.font.size=Pt(32)
    title='2. 성령의 능력으로 변화되는 삶을 살도록'
    body='우리 안에 거하시는 성령의 도우심으로 거룩하게 살아갈 힘을 주소서.'
    replace_prayer(shape,title,body)
    assert shape.left==Inches(1.65)
    assert shape.left+shape.width<=Inches(12.43)
    assert '힘을' in shape.text and '힘\v을' not in shape.text
    assert ' '.join(shape.text.split()).replace(' - ',' ')==' '.join((title+' '+body).split())


def test_blank_retains_inherited_original_background_without_weekly_text():
    from media_automation.ppt.friday_reference import background_only_slide
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
    master=s.slide_layout.slide_master
    master.background.fill.solid();master.background.fill.fore_color.rgb=RGBColor(90,110,120)
    s.shapes.add_textbox(0,0,Inches(2),Inches(1)).text='과거 날짜'
    expected=master._element.find(qn('p:cSld')).find(qn('p:bg')).xml
    background_only_slide(s)
    assert len(s.shapes)==0
    assert s._element.find(qn('p:cSld')).find(qn('p:bg')).xml==expected
    assert s._element.get('showMasterSp')=='0'


def test_blank_keeps_full_slide_background_but_removes_layout_decoration():
    from media_automation.ppt.friday_reference import background_only_slide
    from pptx.enum.shapes import MSO_SHAPE
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);tree=s.slide_layout.shapes._spTree
    bg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,p.slide_width,p.slide_height)
    original=bg._element.xml;tree.append(bg._element)
    decoration=s.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1));decoration.text='지난주 제목';tree.append(decoration._element)
    background_only_slide(s)
    assert len(s.shapes)==1 and len(s.slide_layout.shapes)==0
    assert s.shapes[0]._element.xml==original


def test_reference_parentheses_have_consistent_inner_spacing():
    from media_automation.ppt.friday_reference import normalize_reference
    assert normalize_reference(' ( 요일 2:12~17 ) ')== '(요일 2:12~17)'
    assert normalize_reference('(요일 2:12~17)')== '(요일 2:12~17)'

def test_personal_instruction_group_is_vertically_centered():
    from media_automation.ppt.friday_reference import center_personal
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
    a=s.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1));a.text='개인 기도'
    b=s.shapes.add_textbox(Inches(1),Inches(2.5),Inches(8),Inches(1));b.text='안내'
    footer=s.shapes.add_textbox(0,Inches(7),Inches(1),Inches(.2));footer.text='33'
    center_personal(s,p.slide_height)
    assert abs((a.top+b.top+b.height)/2-p.slide_height/2)<=1
    assert footer.top==Inches(7)
