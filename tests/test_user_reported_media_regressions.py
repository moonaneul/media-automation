"""Regression cases for real worship UI reports (legacy .ppt, mode, download)."""
from pathlib import Path
from pptx import Presentation
from media_automation.web.original import ppt_conversion


UI = Path(__file__).resolve().parents[1] / "src/media_automation/web/original/ui/index.html"


def test_misnamed_real_pptx_converts_without_libreoffice(tmp_path, monkeypatch):
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    original = tmp_path / "song.PPT"
    deck.save(original)
    monkeypatch.setattr(ppt_conversion.shutil, "which",
                        lambda name: (_ for _ in ()).throw(AssertionError("must not run soffice")))
    converted = ppt_conversion.convert_original(original)
    assert converted.suffix == ".pptx"
    assert converted.is_file()
    assert converted.read_bytes() == original.read_bytes()
    assert len(Presentation(converted).slides) == 1
    assert ppt_conversion.conversion_status(original, "song.PPT")["conversion_status"] == "converted"


def test_converter_failure_reports_reason_and_preserves_original(tmp_path, monkeypatch):
    original = tmp_path / "old.ppt"
    original.write_bytes(b"not-a-real-ppt")
    def unavailable(_):
        raise RuntimeError("LibreOffice 미설치")
    monkeypatch.setattr(ppt_conversion, "convert_original", unavailable)
    status = ppt_conversion.conversion_status(original, "old.ppt")
    assert status["conversion_status"] == "failed"
    assert "LibreOffice 미설치" in status["conversion_message"]
    assert original.read_bytes() == b"not-a-real-ppt"


def test_ui_never_imports_entire_onedrive_during_worship_build():
    html = UI.read_text()
    assert "CloudMedia.autoConnect()" not in html
    assert "원본 미삽입" in html
    assert "현재 방식에서 미삽입" in html
    assert "서버 저장 완료 · OneDrive 원본 업로드도 확인됐습니다" in html
    assert "OneDrive 업로드 미확인" in html
    assert "이 서버에서 제작된 파일을 찾지 못했습니다" in html
