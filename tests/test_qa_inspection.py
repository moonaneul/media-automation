from pptx import Presentation
from reportlab.pdfgen import canvas
from media_automation.qa import inspect_file, save_report


def test_empty_pptx_reports_visual_review(tmp_path):
    path = tmp_path / 'test.pptx'
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)
    report = inspect_file(path)
    assert report['stats']['slides'] == 1
    assert report['rendered'] is False
    assert any(x['code'] == 'visual_review_required' for x in report['issues'])
    assert save_report(report, tmp_path / 'report.json').exists()


def test_bad_bulletin_pages_are_errors(tmp_path):
    path = tmp_path / 'bad.pdf'
    pdf = canvas.Canvas(str(path), pagesize=(595, 842))
    pdf.drawString(10, 10, 'Test')
    pdf.save()
    report = inspect_file(path, service='bulletin')
    assert {'page_size', 'folded_page_count'} <= {x['code'] for x in report['issues']}
