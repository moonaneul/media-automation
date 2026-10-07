from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path
from types import SimpleNamespace

import pytest
import yaml
from pptx import Presentation


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "register_sunday_song.py"


@pytest.fixture
def registrar(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.syspath_prepend(str(SCRIPT.parent))
    spec = spec_from_file_location("register_sunday_song", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_pptx(path: Path) -> None:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)


def prepare_week(
    registrar,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    *,
    weekly: bool,
) -> tuple[Path, Path]:
    intake_dir = tmp_path / "output" / "sunday_intake"
    song_dir = tmp_path / "input" / "sunday" / "20260927"
    intake_dir.mkdir(parents=True)
    song_dir.mkdir(parents=True)
    monkeypatch.setattr(registrar, "INTAKE_DIR", intake_dir)
    monkeypatch.setattr(registrar, "SONG_ROOT", song_dir.parent)

    intake = {
        "service": "sunday",
        "date": {"status": "provided", "value": "2026-09-27"},
        "review_required": False,
        "fields": {
            "opening_song_1": {"status": "provided", "value": "이번 주 찬양"},
            "decision_hymn": {"status": "blank", "value": None},
        },
    }
    (intake_dir / "sunday_20260927_intake.yaml").write_text(
        yaml.safe_dump(intake, allow_unicode=True), encoding="utf-8"
    )
    checklist_path = intake_dir / "sunday_20260927_song_checklist.yaml"
    checklist_path.write_text(
        yaml.safe_dump(
            {
                "items": [
                    {
                        "field": "opening_song_1",
                        "title": "수정 전 찬양",
                        "expected_filename": "opening_song_1.pptx",
                        "file": None,
                    }
                ]
            },
            allow_unicode=True,
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    if weekly:
        (intake_dir / "sunday-20260927.yaml").write_text(
            yaml.safe_dump(
                {
                    "service": "sunday",
                    "worship": {
                        "opening_songs": [
                            {"status": "VALUE", "title": "이번 주 찬양"}
                        ]
                    },
                },
                allow_unicode=True,
            ),
            encoding="utf-8",
        )
    return checklist_path, song_dir


def test_register_updates_checklist_and_manifest_when_week_is_ready(
    registrar,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    checklist_path, song_dir = prepare_week(
        registrar, tmp_path, monkeypatch, weekly=True
    )
    source = tmp_path / "사용자 악보.pptx"
    make_pptx(source)
    calls = []

    def fake_run(name, *args, allow=(0,)):
        calls.append(name)
        if name == "refresh_sunday_song_checklist.py":
            checklist = yaml.safe_load(checklist_path.read_text(encoding="utf-8"))
            checklist["items"][0]["file"] = str(song_dir / "opening_song_1.pptx")
            checklist_path.write_text(
                yaml.safe_dump(checklist, allow_unicode=True), encoding="utf-8"
            )
        if name == "resume_sunday_week.py":
            (checklist_path.parent / "sunday-20260927-songs.yaml").write_text(
                "songs: {}\n", encoding="utf-8"
            )
        return 0

    monkeypatch.setattr(registrar, "run_script", fake_run)

    destination, action, item, manifest_ready = registrar.register_song(
        "2026-09-27", "1", source
    )

    assert action == "copied"
    assert destination.read_bytes() == source.read_bytes()
    assert item["title"] == "이번 주 찬양"
    assert item["file"] == str(destination)
    assert manifest_ready
    assert calls == [
        "refresh_sunday_song_checklist.py",
        "resume_sunday_week.py",
    ]


def test_register_without_weekly_keeps_manifest_pending(
    registrar,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    checklist_path, song_dir = prepare_week(
        registrar, tmp_path, monkeypatch, weekly=False
    )
    source = tmp_path / "악보.pptx"
    make_pptx(source)
    calls = []

    def fake_run(name, *args, allow=(0,)):
        calls.append(name)
        if name == "refresh_sunday_song_checklist.py":
            checklist = yaml.safe_load(checklist_path.read_text(encoding="utf-8"))
            checklist["items"][0]["file"] = str(song_dir / "opening_song_1.pptx")
            checklist_path.write_text(
                yaml.safe_dump(checklist, allow_unicode=True), encoding="utf-8"
            )
            return 0
        return 2

    monkeypatch.setattr(registrar, "run_script", fake_run)

    _, _, _, manifest_ready = registrar.register_song("20260927", "1", source)

    assert not manifest_ready
    assert calls == ["refresh_sunday_song_checklist.py", "resume_sunday_week.py"]


def test_register_rejects_slot_absent_from_current_week(
    registrar,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    _, song_dir = prepare_week(registrar, tmp_path, monkeypatch, weekly=True)
    source = tmp_path / "악보.pptx"
    make_pptx(source)

    with pytest.raises(ValueError, match="이번 주 안내"):
        registrar.register_song("2026-09-27", "decision", source)

    assert not (song_dir / "decision_hymn.pptx").exists()


def test_manifest_waits_when_weekly_song_title_is_unset(
    registrar,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    checklist_path, song_dir = prepare_week(
        registrar, tmp_path, monkeypatch, weekly=True
    )
    weekly_path = checklist_path.parent / "sunday-20260927.yaml"
    weekly_path.write_text(
        yaml.safe_dump(
            {
                "worship": {
                    "opening_songs": [{"status": "UNSET", "title": None}]
                }
            }
        ),
        encoding="utf-8",
    )
    source = tmp_path / "악보.pptx"
    make_pptx(source)
    calls = []

    def fake_run(name, *args, allow=(0,)):
        calls.append(name)
        if name == "refresh_sunday_song_checklist.py":
            checklist = yaml.safe_load(checklist_path.read_text(encoding="utf-8"))
            checklist["items"][0]["file"] = str(song_dir / "opening_song_1.pptx")
            checklist_path.write_text(
                yaml.safe_dump(checklist, allow_unicode=True), encoding="utf-8"
            )
            return 0
        return 2

    monkeypatch.setattr(registrar, "run_script", fake_run)

    _, _, _, manifest_ready = registrar.register_song("2026-09-27", "1", source)

    assert not manifest_ready
    assert calls == ["refresh_sunday_song_checklist.py", "resume_sunday_week.py"]


def test_status_does_not_report_ready_from_stale_manifest(
    registrar,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
):
    checklist_path, song_dir = prepare_week(
        registrar, tmp_path, monkeypatch, weekly=True
    )
    (checklist_path.parent / "sunday-20260927-songs.yaml").write_text(
        "songs: {}\n", encoding="utf-8"
    )
    (checklist_path.parent / "sunday-20260927-bible.yaml").write_text(
        "passages: {}\n", encoding="utf-8"
    )

    sunday_path = SCRIPT.parent.parent / "sunday.py"
    spec = spec_from_file_location("sunday_cli_song_status", sunday_path)
    assert spec is not None and spec.loader is not None
    sunday = module_from_spec(spec)
    spec.loader.exec_module(sunday)
    monkeypatch.setattr(sunday, "INTAKE_DIR", checklist_path.parent)
    monkeypatch.setattr(sunday, "INPUT_ROOT", song_dir.parent)

    sunday.status(SimpleNamespace(date="2026-09-27"))

    output = capsys.readouterr().out
    assert "songs    : NOT READY" in output
    assert "score    : 0/1 READY" in output
    assert "SUNDAY INPUTS READY" not in output
