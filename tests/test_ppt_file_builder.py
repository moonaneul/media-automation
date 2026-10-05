from pathlib import Path

import pytest
from pptx import Presentation

from media_automation.bible import (
    BiblePassage,
    BibleVerse,
    InMemoryBibleProvider,
)
from media_automation.planning import (
    BlockKind,
    WorshipBlock,
)
from media_automation.ppt import (
    FilePreServiceSlideProvider,
    FileSongAssetProvider,
    MissingRenderDependencyError,
    MissingSongAssetError,
    build_presentation_file_from_plan,
)
from media_automation.weekly_data.models import (
    PersonField,
    ScriptureField,
    Song,
)


def make_ppt(
    path: Path,
    text: str | None = None,
):
    prs = Presentation()
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)

    if text is not None:
        box = slide.shapes.add_textbox(
            0,
            0,
            3000000,
            1000000,
        )
        box.text = text

    prs.save(path)


class FakeSlideMerger:
    def insert_all(
        self,
        destination,
        source,
        *,
        after_slide,
    ):
        source_prs = Presentation(source)

        self.insert_range(
            destination,
            source,
            after_slide=after_slide,
            start_slide=1,
            end_slide=len(source_prs.slides),
        )

    def insert_range(
        self,
        destination,
        source,
        *,
        after_slide,
        start_slide,
        end_slide,
    ):
        destination_prs = Presentation(
            destination
        )
        source_prs = Presentation(
            source
        )

        for slide_number in range(
            start_slide,
            end_slide + 1,
        ):
            source_slide = (
                source_prs.slides[
                    slide_number - 1
                ]
            )

            layout = (
                destination_prs
                .slide_layouts[6]
            )

            new_slide = (
                destination_prs
                .slides
                .add_slide(layout)
            )

            for shape in source_slide.shapes:
                if not hasattr(
                    shape,
                    "text",
                ):
                    continue

                box = (
                    new_slide
                    .shapes
                    .add_textbox(
                        0,
                        0,
                        3000000,
                        1000000,
                    )
                )
                box.text = shape.text

        destination_prs.save(
            destination
        )


def test_file_builder_requires_bible_provider(
    tmp_path,
):
    output = tmp_path / "result.pptx"

    plan = [
        WorshipBlock(
            kind=BlockKind.SCRIPTURE,
            key="scripture",
            value=ScriptureField(
                status="VALUE",
                reference="삼상 16:6-7",
            ),
        )
    ]

    with pytest.raises(
        MissingRenderDependencyError
    ):
        build_presentation_file_from_plan(
            plan,
            output,
        )


def test_file_builder_requires_song_provider(
    tmp_path,
):
    output = tmp_path / "result.pptx"

    plan = [
        WorshipBlock(
            kind=BlockKind.SONG,
            key="song",
            value=Song(
                status="VALUE",
                title="찬양",
            ),
        )
    ]

    with pytest.raises(
        MissingRenderDependencyError
    ):
        build_presentation_file_from_plan(
            plan,
            output,
            slide_merger=FakeSlideMerger(),
        )


def test_file_builder_rejects_missing_song(
    tmp_path,
):
    output = tmp_path / "result.pptx"

    plan = [
        WorshipBlock(
            kind=BlockKind.SONG,
            key="song",
            value=Song(
                status="VALUE",
                title="없는 찬양",
            ),
        )
    ]

    with pytest.raises(
        MissingSongAssetError
    ):
        build_presentation_file_from_plan(
            plan,
            output,
            song_asset_provider=(
                FileSongAssetProvider({})
            ),
            slide_merger=FakeSlideMerger(),
        )


def test_file_builder_can_skip_missing_song(
    tmp_path,
):
    output = tmp_path / "result.pptx"

    plan = [
        WorshipBlock(
            kind=BlockKind.SONG,
            key="song",
            value=Song(
                status="VALUE",
                title="없는 찬양",
            ),
        ),
        WorshipBlock(
            kind=BlockKind.PRAYER,
            key="prayer",
            value=PersonField(
                status="VALUE",
                person="홍길동",
            ),
        ),
    ]

    result = build_presentation_file_from_plan(
        plan,
        output,
        song_asset_provider=(
            FileSongAssetProvider({})
        ),
        slide_merger=FakeSlideMerger(),
        skip_missing_songs=True,
    )

    prs = Presentation(result)

    assert len(prs.slides) == 1
    assert (
        prs.slides[0]
        .shapes[0]
        .text
        == "기 도 : 홍길동"
    )


def test_missing_middle_song_does_not_create_double_blank(
    tmp_path,
):
    song_a = tmp_path / "a.pptx"
    song_c = tmp_path / "c.pptx"
    output = tmp_path / "result.pptx"

    make_ppt(
        song_a,
        "SONG A",
    )
    make_ppt(
        song_c,
        "SONG C",
    )

    provider = FileSongAssetProvider(
        {
            "찬양 A": song_a,
            "찬양 C": song_c,
        }
    )

    plan = [
        WorshipBlock(
            kind=BlockKind.SONG,
            key="a",
            value=Song(
                status="VALUE",
                title="찬양 A",
            ),
        ),
        WorshipBlock(
            kind=BlockKind.BLANK,
            key="a->b",
        ),
        WorshipBlock(
            kind=BlockKind.SONG,
            key="b",
            value=Song(
                status="VALUE",
                title="없는 찬양 B",
            ),
        ),
        WorshipBlock(
            kind=BlockKind.BLANK,
            key="b->c",
        ),
        WorshipBlock(
            kind=BlockKind.SONG,
            key="c",
            value=Song(
                status="VALUE",
                title="찬양 C",
            ),
        ),
    ]

    result = build_presentation_file_from_plan(
        plan,
        output,
        song_asset_provider=provider,
        slide_merger=FakeSlideMerger(),
        skip_missing_songs=True,
    )

    prs = Presentation(result)

    assert len(prs.slides) == 3


def test_file_builder_uses_song_slide_range(
    tmp_path,
):
    source = tmp_path / "songs.pptx"
    output = tmp_path / "result.pptx"

    prs = Presentation()

    layout = prs.slide_layouts[6]

    for number in range(1, 6):
        slide = prs.slides.add_slide(
            layout
        )
        box = slide.shapes.add_textbox(
            0,
            0,
            3000000,
            1000000,
        )
        box.text = f"SLIDE {number}"

    prs.save(source)

    from media_automation.ppt import (
        SongAsset,
    )

    provider = FileSongAssetProvider(
        {
            "찬양": SongAsset(
                title="찬양",
                path=source,
                start_slide=2,
                end_slide=4,
            )
        }
    )

    plan = [
        WorshipBlock(
            kind=BlockKind.SONG,
            key="song",
            value=Song(
                status="VALUE",
                title="찬양",
            ),
        )
    ]

    result = build_presentation_file_from_plan(
        plan,
        output,
        song_asset_provider=provider,
        slide_merger=FakeSlideMerger(),
    )

    reopened = Presentation(result)

    assert len(reopened.slides) == 3


def test_file_builder_renders_scripture(
    tmp_path,
):
    output = tmp_path / "result.pptx"

    provider = InMemoryBibleProvider(
        {
            "삼상 16:6-7": BiblePassage(
                reference="삼상 16:6~7",
                verses=(
                    BibleVerse(
                        number=6,
                        text="6절",
                    ),
                    BibleVerse(
                        number=7,
                        text="7절",
                    ),
                ),
            )
        }
    )

    plan = [
        WorshipBlock(
            kind=BlockKind.SCRIPTURE,
            key="scripture",
            value=ScriptureField(
                status="VALUE",
                reference="삼상 16:6-7",
            ),
        )
    ]

    result = build_presentation_file_from_plan(
        plan,
        output,
        bible_provider=provider,
    )

    prs = Presentation(result)

    assert len(prs.slides) == 2

def test_file_builder_can_preserve_sunday_edge_blanks(
    tmp_path,
):
    pre_service = (
        tmp_path
        / "pre_service.pptx"
    )
    output = (
        tmp_path
        / "sunday-result.pptx"
    )

    make_ppt(
        pre_service,
        "PRE SERVICE",
    )

    provider = (
        FilePreServiceSlideProvider(
            {
                "pre_service": (
                    pre_service
                ),
            },
            slide_counts={
                "pre_service": 1,
            },
        )
    )

    plan = [
        WorshipBlock(
            kind=BlockKind.PRE_SERVICE,
            key="pre_service",
        ),
        WorshipBlock(
            kind=BlockKind.BLANK,
            key=(
                "transition:"
                "pre_service->prayer"
            ),
        ),
        WorshipBlock(
            kind=BlockKind.PRAYER,
            key="prayer",
            value=PersonField(
                status="VALUE",
                person="홍길동",
            ),
        ),
        WorshipBlock(
            kind=BlockKind.BLANK,
            key=(
                "transition:"
                "prayer->end"
            ),
        ),
    ]

    result = (
        build_presentation_file_from_plan(
            plan,
            output,
            pre_service_provider=provider,
            preserve_blank_after_pre_service=True,
            preserve_trailing_blank=True,
        )
    )

    prs = Presentation(result)

    assert len(prs.slides) == 4

    assert (
        "PRE SERVICE"
        in prs.slides[0].shapes[0].text
    )

    assert not any(
        getattr(shape, "text", "").strip()
        for shape in prs.slides[1].shapes
    )

    assert any(
        "기 도 : 홍길동"
        in getattr(shape, "text", "")
        for shape in prs.slides[2].shapes
    )

    assert not any(
        getattr(shape, "text", "").strip()
        for shape in prs.slides[3].shapes
    )

