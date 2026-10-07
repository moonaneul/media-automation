from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path

import pytest
import yaml
from pptx import Presentation


SCRIPT = (
    Path(__file__).resolve().parents[1]
    / "scripts"
    / "register_wednesday_song.py"
)

spec = spec_from_file_location("register_wednesday_song", SCRIPT)
module = module_from_spec(spec)
assert spec is not None
assert spec.loader is not None
spec.loader.exec_module(module)


def make_pptx(path: Path, text: str = "악보") -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 1_000_000, 1_000_000)
    box.text = text
    presentation.save(path)


def prepare_week(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Path, Path]:
    intake_dir = tmp_path / "output" / "wednesday_intake"
    song_root = tmp_path / "input" / "wednesday"
    intake_dir.mkdir(parents=True)
    monkeypatch.setattr(module, "ROOT", tmp_path)
    monkeypatch.setattr(module, "INTAKE_DIR", intake_dir)
    monkeypatch.setattr(module, "SONG_ROOT", song_root)

    checklist = intake_dir / "wednesday_20261007_song_checklist.yaml"
    checklist.write_text(
        yaml.safe_dump(
            {
                "items": [
                    {
                        "field": "opening_song_1",
                        "title": "새 찬양",
                        "expected_filename": "opening_song_1.pptx",
                        "file": None,
                    }
                ]
            },
            allow_unicode=True,
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    weekly = intake_dir / "wednesday-20261007.yaml"
    weekly.write_text("service: wednesday\n", encoding="utf-8")
    return checklist, song_root


def test_slot_aliases_are_explicit():
    assert module.canonical_slot("1") == "opening_song_1"
    assert module.canonical_slot("additional") == "additional_song"
    assert module.canonical_slot("decision-hymn") == "decision_hymn"

    with pytest.raises(ValueError, match="알 수 없는 슬롯"):
        module.canonical_slot("지난주 찬양")


def test_register_pptx_copies_to_expected_slot_and_refreshes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    checklist, song_root = prepare_week(tmp_path, monkeypatch)
    source = tmp_path / "사용자 제공 악보.pptx"
    make_pptx(source)
    calls = []

    def fake_run(name, *args, allow=(0,)):
        calls.append((name, args, allow))
        if name == "refresh_wednesday_song_checklist.py":
            data = yaml.safe_load(checklist.read_text(encoding="utf-8"))
            data["items"][0]["file"] = str(
                song_root / "20261007" / "opening_song_1.pptx"
            )
            checklist.write_text(
                yaml.safe_dump(data, allow_unicode=True, sort_keys=False),
                encoding="utf-8",
            )
        return 0

    monkeypatch.setattr(module, "run_script", fake_run)

    destination, action, item = module.register_song(
        "2026-10-07", "1", source
    )

    assert action == "copied"
    assert destination == song_root / "20261007" / "opening_song_1.pptx"
    assert destination.read_bytes() == source.read_bytes()
    assert item["file"] == str(destination)
    assert [call[0] for call in calls] == [
        "refresh_wednesday_song_checklist.py",
        "build_wednesday_song_manifest.py",
    ]


def test_register_ppt_uses_conversion_without_modifying_source(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    checklist, song_root = prepare_week(tmp_path, monkeypatch)
    source = tmp_path / "사용자 제공 악보.ppt"
    original = b"legacy-ppt-content"
    source.write_bytes(original)

    def fake_convert(given_source: Path, destination: Path):
        assert given_source == source.resolve()
        make_pptx(destination, "변환본")

    def fake_run(name, *args, allow=(0,)):
        if name == "refresh_wednesday_song_checklist.py":
            data = yaml.safe_load(checklist.read_text(encoding="utf-8"))
            data["items"][0]["file"] = str(
                song_root / "20261007" / "opening_song_1.pptx"
            )
            checklist.write_text(
                yaml.safe_dump(data, allow_unicode=True, sort_keys=False),
                encoding="utf-8",
            )
        return 0

    monkeypatch.setattr(module, "convert_legacy_ppt", fake_convert)
    monkeypatch.setattr(module, "run_script", fake_run)

    destination, action, _ = module.register_song("20261007", "opening1", source)

    assert action == "converted"
    assert source.read_bytes() == original
    assert destination.is_file()
    assert Presentation(destination).slides[0].shapes[0].text == "변환본"


def test_register_refuses_slot_not_present_this_week(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    prepare_week(tmp_path, monkeypatch)
    source = tmp_path / "악보.pptx"
    make_pptx(source)

    with pytest.raises(ValueError, match="이번 주 체크리스트"):
        module.register_song("2026-10-07", "decision", source)


def test_stage_song_requires_replace_for_existing_destination(tmp_path: Path):
    source = tmp_path / "source.pptx"
    destination = tmp_path / "slot.pptx"
    make_pptx(source, "새 파일")
    make_pptx(destination, "기존 파일")

    with pytest.raises(FileExistsError, match="--replace"):
        module.stage_song(source, destination)

    action = module.stage_song(source, destination, replace=True)

    assert action == "copied"
    assert destination.read_bytes() == source.read_bytes()


def test_date_token_rejects_invalid_date():
    with pytest.raises(ValueError, match="날짜"):
        module.date_token("2026-02-30")
