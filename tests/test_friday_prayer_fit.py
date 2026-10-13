from pptx.util import Inches, Pt
from media_automation.ppt.base import create_16x9_presentation
from media_automation.ppt.friday_reference import replace_prayer
import pytest


def original_shape():
    deck=create_16x9_presentation()
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    shape=slide.shapes.add_textbox(Inches(1.65),Inches(2),Inches(10.78),Inches(4.8))
    title=shape.text_frame.paragraphs[0].add_run();title.text='원본 제목';title.font.size=Pt(36)
    body=shape.text_frame.add_paragraph().add_run();body.text='원본 보조 본문';body.font.size=Pt(32)
    return shape


def test_short_prayer_keeps_original_sizes():
    shape=original_shape()
    replace_prayer(shape,'1. 짧은 기도','짧은 보조 말씀')
    sizes={r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs}
    assert sizes=={36,32}
    assert shape.top==Inches(2)


def test_long_prayer_only_adjusts_its_own_shape_and_preserves_text():
    short=original_shape();long=original_shape()
    title='4. 하늘빛기쁨교회가 영혼 구원하여 제자 삼는 교회로 쓰임 받을 수 있도록'
    body='지상 사명 말씀 연결 검증을 위한 긴 보조 문장입니다. '*8
    replace_prayer(short,'1. 짧은 기도','짧은 보조 말씀')
    replace_prayer(long,title,body)
    sizes={r.font.size.pt for p in long.text_frame.paragraphs for r in p.runs}
    assert min(sizes)>=24
    assert long.top+long.height<=Inches(7)
    assert min(sizes)<32
    assert {r.font.size.pt for p in short.text_frame.paragraphs for r in p.runs}=={36,32}
    normalized=''.join(long.text.split())
    assert ''.join(title.split()) in normalized
    assert ''.join(body.split()) in normalized


def test_unreadable_amount_is_rejected_instead_of_tiny_text():
    with pytest.raises(ValueError,match='두 화면'):
        replace_prayer(original_shape(),'기도 제목','너무 긴 검증용 내용입니다. '*100)
