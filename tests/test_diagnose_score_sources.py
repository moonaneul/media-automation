"""The score diagnostic never alters a church DB or source assets."""
import hashlib
import sqlite3
from pathlib import Path

from pptx import Presentation

from media_automation.web.original.work_store import WorkStore
from scripts.diagnose_score_sources import read_rows, probe_file, report_generated


def test_score_probe_converts_disposable_copy_without_touching_original(tmp_path, capsys):
    store=WorkStore(tmp_path/'work.db')
    deck=Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    original=tmp_path/'misnamed.ppt'
    deck.save(original)
    attachment=store.attach('2026-10-14','wednesday','시작 찬양 1','1.ppt',original.read_bytes())
    digest=attachment['digest']
    stored=tmp_path/'work.files'/digest
    before=hashlib.sha256(stored.read_bytes()).hexdigest()
    files_before={str(path.relative_to(tmp_path/'work.files')):path.stat().st_mtime_ns for path in (tmp_path/'work.files').rglob('*') if path.is_file()}
    listed,reports=read_rows(tmp_path/'work.db',['1.ppt'])
    assert listed==[('1.ppt',digest,'2026-10-14/wednesday/시작 찬양 1')]
    assert reports==[]
    probe_file(tmp_path/'work.db','1.ppt',digest)
    text=capsys.readouterr().out
    assert '변환/읽기: 성공' in text
    assert '1장' in text
    assert hashlib.sha256(stored.read_bytes()).hexdigest()==before
    files_after={str(path.relative_to(tmp_path/'work.files')):path.stat().st_mtime_ns for path in (tmp_path/'work.files').rglob('*') if path.is_file()}
    assert files_after==files_before


def test_report_shows_stale_generated_records(tmp_path,capsys):
    report_generated(tmp_path/'work.db', [('2026-10-14','wednesday','a'*32)])
    text=capsys.readouterr().out
    assert 'PPT 파일 없음' in text
