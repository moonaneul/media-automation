import platform

import pytest
from pptx import Presentation

from media_automation.ppt import (
    PowerPointComSlideMerger,
)


def make_ppt(
    path,
    slide_count=1,
):
    prs = Presentation()

    layout = prs.slide_layouts[6]

    for _ in range(
        slide_count
    ):
        prs.slides.add_slide(
            layout
        )

    prs.save(
        path
    )


def test_insert_all_empty_source_does_nothing(
    tmp_path,
):
    destination = (
        tmp_path / "destination.pptx"
    )
    source = (
        tmp_path / "source.pptx"
    )

    Presentation().save(
        destination
    )
    Presentation().save(
        source
    )

    merger = (
        PowerPointComSlideMerger()
    )

    # source에 슬라이드가 없으므로
    # COM을 호출하지 않아도 된다.
    merger.insert_all(
        destination,
        source,
        after_slide=0,
    )


def test_insert_range_rejects_invalid_start(
    tmp_path,
):
    if platform.system() != "Windows":
        pytest.skip(
            "COM 실행 전 Windows 환경 검사가 우선됨"
        )

    merger = (
        PowerPointComSlideMerger()
    )

    with pytest.raises(
        ValueError
    ):
        merger.insert_range(
            tmp_path / "a.pptx",
            tmp_path / "b.pptx",
            after_slide=0,
            start_slide=0,
            end_slide=1,
        )


def test_com_merger_is_windows_only():
    if platform.system() == "Windows":
        pytest.skip(
            "Windows에서는 COM 사용 가능"
        )

    merger = (
        PowerPointComSlideMerger()
    )

    with pytest.raises(
        RuntimeError,
        match="Windows",
    ):
        merger.insert_range(
            "destination.pptx",
            "source.pptx",
            after_slide=0,
            start_slide=1,
            end_slide=1,
        )
def test_insert_all_delegates_full_slide_range(
    tmp_path,
):
    source = (
        tmp_path
        / "source.pptx"
    )

    destination = (
        tmp_path
        / "destination.pptx"
    )

    make_ppt(
        source,
        slide_count=3,
    )

    make_ppt(
        destination,
        slide_count=1,
    )

    calls = []

    class RecordingMerger(
        PowerPointComSlideMerger
    ):
        def insert_range(
            self,
            destination,
            source,
            *,
            after_slide,
            start_slide,
            end_slide,
        ):
            calls.append(
                {
                    "after_slide": after_slide,
                    "start_slide": start_slide,
                    "end_slide": end_slide,
                }
            )

    merger = RecordingMerger()

    merger.insert_all(
        destination,
        source,
        after_slide=1,
    )

    assert calls == [
        {
            "after_slide": 1,
            "start_slide": 1,
            "end_slide": 3,
        }
    ]
from media_automation.ppt import (
    MacPowerPointAppleScriptSlideMerger,
    create_platform_slide_merger,
)


def test_platform_merger_rejects_macos_for_final_build(
    monkeypatch,
):
    monkeypatch.setattr(
        "media_automation.ppt.slide_merge.platform.system",
        lambda: "Darwin",
    )

    with pytest.raises(
        RuntimeError,
        match="Windows 제작 환경",
    ):
        create_platform_slide_merger()


def test_mac_merger_rejects_insert_before_first_slide():
    if platform.system() != "Darwin":
        pytest.skip(
            "macOS 전용 검증"
        )

    merger = (
        MacPowerPointAppleScriptSlideMerger()
    )

    with pytest.raises(
        ValueError,
        match="첫 슬라이드 앞",
    ):
        merger.insert_range(
            "destination.pptx",
            "source.pptx",
            after_slide=0,
            start_slide=1,
            end_slide=1,
        )