"""Office treats OPC image part names case-insensitively. Preserve mixed-case scores."""
from collections import defaultdict
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from media_automation.ppt.slide_merge import OpenXmlSlideMerger


def test_uppercase_score_images_do_not_collide_with_lowercase_welcome(tmp_path):
    first = tmp_path / 'welcome.png'
    score = tmp_path / 'score.png'
    Image.new('RGB', (24, 24), (0, 0, 250)).save(first)
    Image.new('RGB', (24, 24), (250, 0, 0)).save(score)

    target_path, source_path = tmp_path / 'target.pptx', tmp_path / 'source.pptx'
    target = Presentation()
    target.slide_width, target.slide_height = Inches(10), Inches(7.5)
    target.slides.add_slide(target.slide_layouts[6]).shapes.add_picture(str(first), 0, 0)
    target.save(target_path)

    source = Presentation()
    source.slide_width, source.slide_height = Inches(10), Inches(7.5)
    source.slides.add_slide(source.slide_layouts[6]).shapes.add_picture(str(score), 0, 0)
    source.save(source_path)

    # Simulate a real downloaded score whose media extension is uppercase PNG,
    # while the worship template uses lowercase png for a *different* image.
    temporary = tmp_path / 'source-uppercase.pptx'
    with ZipFile(source_path) as existing, ZipFile(temporary, 'w', ZIP_DEFLATED) as changed:
        for entry in existing.infolist():
            name = entry.filename
            data = existing.read(name)
            if name == 'ppt/media/image1.png':
                name = 'ppt/media/image1.PNG'
            elif name == 'ppt/slides/_rels/slide1.xml.rels':
                data = data.replace(b'../media/image1.png', b'../media/image1.PNG')
            elif name == '[Content_Types].xml':
                root = ET.fromstring(data)
                types = 'http://schemas.openxmlformats.org/package/2006/content-types'
                ET.SubElement(root, '{%s}Default' % types,
                              {'Extension': 'PNG', 'ContentType': 'image/png'})
                data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            changed.writestr(name, data)

    OpenXmlSlideMerger().insert_all(target_path, temporary, after_slide=1)
    made = Presentation(target_path)
    assert len(made.slides) == 2
    assert made.slides[0].shapes[0].image.blob == first.read_bytes()
    assert made.slides[1].shapes[0].image.blob == score.read_bytes()

    with ZipFile(target_path) as archive:
        names = defaultdict(list)
        for name in archive.namelist():
            names[name.casefold()].append(name)
        assert not {k: v for k, v in names.items() if len(v) > 1}
