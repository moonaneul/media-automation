from media_automation.ppt.base import (
    create_16x9_presentation,
)
from media_automation.ppt.friday_zoom import (
    add_prayer_topics_slide,
)
from pptx.util import Pt


def test_supporting_scripture_is_smaller_in_same_topic_shape():
    prs = create_16x9_presentation()
    title = "염려가 나를 찾아올 때마다 말씀과 기도로 승리할 수 있도록"
    scripture = "[6] 아무것도 염려하지 말고 [7] 너희 마음과 생각을 지키시리라 (빌 4:6~7)"
    slide = add_prayer_topics_slide(prs, [title + "\n" + scripture, "두 번째 제목"])
    shape = next(s for s in slide.shapes if s.name == "FridayZoomPrayerTopic1")
    paragraphs = shape.text_frame.paragraphs
    assert shape.text == "1. " + title + "\n" + scripture
    assert paragraphs[0].runs[0].font.size == Pt(38)
    assert paragraphs[1].runs[0].font.size == Pt(28)
    assert len([s for s in slide.shapes if s.name.startswith("FridayZoomPrayerTopic")]) == 2


def test_prayer_topics_use_separate_shapes():
    prs = create_16x9_presentation()

    slide = add_prayer_topics_slide(
        prs,
        [
            "first topic",
            "second topic",
            "third topic",
        ],
    )

    shapes = [
        shape
        for shape in slide.shapes
        if shape.name.startswith(
            "FridayZoomPrayerTopic"
        )
    ]

    assert len(shapes) == 3

    assert [
        shape.name
        for shape in shapes
    ] == [
        "FridayZoomPrayerTopic1",
        "FridayZoomPrayerTopic2",
        "FridayZoomPrayerTopic3",
    ]

    # ?? Zoom?? ?? ??? ?? ???
    # animation?? ?? ??? ? ??.
    assert len(
        {
            (
                shape.left,
                shape.top,
                shape.width,
                shape.height,
            )
            for shape in shapes
        }
    ) == 1
