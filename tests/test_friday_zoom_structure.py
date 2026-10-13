from pathlib import Path

import yaml
from pptx import Presentation

from media_automation.bible import (
    BiblePassage,
    BibleVerse,
    InMemoryBibleProvider,
)
from media_automation.planning import (
    build_friday_zoom_plan,
)
from media_automation.ppt import (
    build_friday_zoom_structure,
)
from media_automation.weekly_data.models import (
    FridayZoomData,
    parse_weekly_data,
)


ROOT = Path(__file__).resolve().parents[1]


def load_weekly(
    filename: str,
) -> FridayZoomData:
    path = (
        ROOT
        / "samples"
        / "weekly"
        / filename
    )

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:
        raw = yaml.safe_load(file)

    weekly = parse_weekly_data(raw)

    assert isinstance(
        weekly,
        FridayZoomData,
    )

    return weekly


def load_bible(
    filename: str,
) -> InMemoryBibleProvider:
    path = (
        ROOT
        / "samples"
        / "weekly"
        / filename
    )

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:
        raw = yaml.safe_load(file)

    passages = {}

    for lookup_reference, data in (
        raw["passages"].items()
    ):
        verses = tuple(
            BibleVerse(
                number=int(number),
                text=text,
            )
            for number, text
            in data["verses"].items()
        )

        passages[
            lookup_reference
        ] = BiblePassage(
            reference=data["reference"],
            verses=verses,
        )

    return InMemoryBibleProvider(
        passages
    )



def assert_complete_structure(prs, result, bible, additional):
    ranges = result.slide_ranges
    keys = ['pre_service', 'opening_songs[0]', 'opening_songs[1]', 'first_prayer',
            'song_after_prayer', 'scripture', 'sermon_title']
    if additional: keys.append('additional_scripture')
    else: assert 'additional_scripture' not in ranges
    keys += ['response_song', 'word_prayer', 'intercession_song', 'community_prayer', 'personal_prayer']
    cursor = 1
    for key in keys:
        assert ranges[key].start == cursor
        assert ranges[key].end >= ranges[key].start
        cursor = ranges[key].end + 1
    assert cursor - 1 == len(prs.slides)
    # Pagination may change with typography; all supplied verse text must remain.
    text = ''.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    compact = lambda value: ''.join(value.split()).replace('세례','침례')
    for passage in bible._passages.values():
        for verse in passage.verses:
            assert compact(verse.text) in compact(text)

def test_20260911_structure_preserves_order_and_verses(
    tmp_path,
):
    weekly = load_weekly(
        "friday-zoom-20260911.yaml"
    )

    bible = load_bible(
        "friday-zoom-20260911-bible.yaml"
    )

    plan = build_friday_zoom_plan(
        weekly
    )

    output = (
        tmp_path
        / "20260911.pptx"
    )

    result = build_friday_zoom_structure(
        plan,
        output,
        bible_provider=bible,
    )

    prs = Presentation(output)

    assert_complete_structure(prs, result, bible, True)


def test_20260918_structure_preserves_order_and_verses(
    tmp_path,
):
    weekly = load_weekly(
        "friday-zoom-20260918.yaml"
    )

    bible = load_bible(
        "friday-zoom-20260918-bible.yaml"
    )

    plan = build_friday_zoom_plan(
        weekly
    )

    output = (
        tmp_path
        / "20260918.pptx"
    )

    result = build_friday_zoom_structure(
        plan,
        output,
        bible_provider=bible,
    )

    prs = Presentation(output)

    assert_complete_structure(prs, result, bible, False)


def test_20260904_structure_preserves_order_and_verses(
    tmp_path,
):
    weekly = load_weekly(
        "friday-zoom-20260904.yaml"
    )

    bible = load_bible(
        "friday-zoom-20260904-bible.yaml"
    )

    plan = build_friday_zoom_plan(
        weekly
    )

    output = (
        tmp_path
        / "20260904.pptx"
    )

    result = build_friday_zoom_structure(
        plan,
        output,
        bible_provider=bible,
    )

    prs = Presentation(output)

    assert_complete_structure(prs, result, bible, False)
