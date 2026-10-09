import hashlib
import json

import pytest
from pptx import Presentation

from media_automation.web.application import Application


def setup_generated(tmp_path):
    repo = tmp_path / "repo"
    repo.mkdir()
    for service in ("sunday", "wednesday", "friday", "bulletin"):
        (repo / (service + ".py")).write_text("print('placeholder')\n")
    app = Application(repo, tmp_path / "jobs")
    job_id = app.create("sunday", "2025-01-08")["job_id"]
    folder = app.jobs._directory(job_id)
    artifact = folder / "workspace/output/sunday/result.pptx"
    artifact.parent.mkdir(parents=True)
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(artifact)
    record = json.loads((folder / "job.json").read_text())
    record.update(state="generated", artifacts=["output/sunday/result.pptx"])
    app.jobs._save(folder, record)
    return app, job_id, folder, artifact


def test_review_cannot_access_non_generated_artifacts(tmp_path):
    app, job, folder, artifact = setup_generated(tmp_path)
    with pytest.raises(ValueError):
        app.qa_inspect(job, "input/notice.txt")
    record = json.loads((folder / "job.json").read_text())
    record["state"] = "failed"
    app.jobs._save(folder, record)
    with pytest.raises(ValueError):
        app.qa_inspect(job, "output/sunday/result.pptx")


def test_review_requires_explicit_human_checks_and_hash(tmp_path):
    app, job, folder, artifact = setup_generated(tmp_path)
    path = "output/sunday/result.pptx"
    assert not app.qa_status(job, path)["available"]
    inspected = app.qa_inspect(job, path)
    assert inspected["available"]
    assert inspected["automatic"]["rendered"] is False
    assert not inspected["human_review_complete"]
    labels = inspected["check_labels"]
    with pytest.raises(ValueError):
        app.qa_confirm(job, path, {"sha256": "bad", "reviewer": "tester", "checks": {}})
    values = {key: "checked" for key in labels}
    saved = app.qa_confirm(job, path, {
        "sha256": inspected["sha256"], "reviewer": "QA 담당",
        "note": "실제 화면 별도 확인", "checks": values,
    })
    assert saved["review_record_complete"]
    assert app.qa_report(job, path)
    changed = Presentation()
    changed.slides.add_slide(changed.slide_layouts[6])
    changed.slides.add_slide(changed.slide_layouts[6])
    changed.save(artifact)
    assert app.qa_status(job, path)["stale"]
    with pytest.raises(ValueError):
        app.qa_report(job, path)
    again = app.qa_inspect(job, path)
    assert not again["human_review_complete"]
