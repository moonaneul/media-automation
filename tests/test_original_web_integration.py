import json
from pathlib import Path
import shutil
import subprocess
import threading
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from zipfile import ZipFile
import pytest
from pptx import Presentation
from pypdf import PdfReader
from media_automation.web.original.server import make_server
from media_automation.web.original.work_store import WorkStore, RevisionConflict
from media_automation.web.original.generation import generate, ROOT
from media_automation.web.original.bulletin_generation import generate_bulletin
from media_automation.ppt.base import create_4x3_presentation,add_centered_text_slide
from media_automation.web.original.notice_files import extract_notice


def value(text): return {'state':'VALUE','value':text}

def setup(tmp_path, service='sunday', mode=None):
    store=WorkStore(tmp_path/'work.db')
    data={'fields':{'성경 본문':value('단 1:8~9'),'설교 제목':value('연결 시험')},'outputs':{'ppt':{},'pdf':{}}}
    if mode:data['fridayMode']=mode
    store.save('2026-10-11',service,0,data)
    body=dict(date='2026-10-11',service=service,revision=1,groups=[['말씀',['설교 제목','성경 본문']]])
    return store,body


def test_primary_root_is_original_ui_and_revision_storage(tmp_path):
    store=WorkStore(tmp_path/'work.db');server=make_server(store,0)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    base=f'http://127.0.0.1:{server.server_port}'
    def request(path,data=None):
        return urlopen(Request(base+path,data=json.dumps(data).encode() if data else None,headers={'Content-Type':'application/json'}),timeout=3)
    try:
        html=request('/').read()
        assert b'notice-parser.js' in html
        from media_automation.web.original.generation import BUILD_VERSION, FRIDAY_TYPOGRAPHY_VERSION
        assert BUILD_VERSION.encode() in html and FRIDAY_TYPOGRAPHY_VERSION.encode() in html
        assert b'__ENGINE_VERSION__' not in html
        data={'fields':{'기도 담당자':value('검증 담당')}}
        assert json.load(request('/api/work?date=2026-10-11&service=sunday',{'revision':0,'data':data}))['revision']==1
        assert json.load(request('/api/work?date=2026-10-11&service=sunday'))['data']==data
        with pytest.raises(HTTPError) as e:request('/api/work?date=2026-10-11&service=sunday',{'revision':0,'data':data})
        assert e.value.code==409
        with pytest.raises(HTTPError) as e:urlopen(Request(base+'/api/work',headers={'Origin':'https://outside.invalid'}))
        assert e.value.code==403
        with pytest.raises(HTTPError):request('/blue')
    finally:server.shutdown();server.server_close();thread.join()


@pytest.mark.parametrize('service,mode,ratio',[('sunday',None,4/3),('wednesday',None,4/3),('friday','in_person',4/3),('friday','zoom',16/9)])
def test_all_service_paths_make_real_ppt_using_official_renderer(tmp_path,service,mode,ratio):
    store,body=setup(tmp_path,service,mode)
    record=store.get(body['date'],service)
    record['data']['fields']['설교 제목']=value('  . - 연결 시험')
    record['data']['fields']['성경 본문']=value('  . - 단 1:8~9')
    record['data']['fields']['추가 말씀']=value('- 사 43:21, . 왕상 8:27, • 히 10:19~20')
    store.save(body['date'],service,record['revision'],record['data'])
    body['revision']=record['revision']+1
    body['groups'][0][1].append('추가 말씀')
    result=generate(store,body);path=Path(store.path).with_suffix('.generated')/(result['file_id']+'.pptx')
    deck=Presentation(path)
    assert abs(deck.slide_width/deck.slide_height-ratio)<.01
    text='\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert '연결 시험' in text and '다니엘' in text
    assert '. - 연결 시험' not in text
    assert not any(item.startswith('추가 말씀:') for item in result['unresolved'])
    assert '세례' not in text
    assert result['complete'] is False
    with ZipFile(path) as z:
        assert z.testzip() is None
        if service in ('sunday', 'wednesday'):
            from media_automation.ppt.sunday_typography import VERSION
            assert result['typography_version'] == VERSION
            assert '경기천년제목V Bold' in ''.join(z.read(n).decode('utf-8') for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml'))
        else:
            from media_automation.ppt.friday_typography import VERSION as FRIDAY_VERSION
            assert result['typography_version'] == FRIDAY_VERSION


def test_stale_revision_rejected_before_creating_output(tmp_path):
    store,body=setup(tmp_path);body['revision']=0
    with pytest.raises(ValueError,match='입력이 변경'):generate(store,body)
    assert not Path(store.path).with_suffix('.generated').exists()


def test_song_original_entire_sequence_and_missing_blank(tmp_path):
    store,body=setup(tmp_path)
    song=tmp_path/'song.pptx';deck=create_4x3_presentation()
    for text in ('원본 1절','원본 후렴','원본 아멘'):add_centered_text_slide(deck,text)
    deck.save(song);store.attach(body['date'],'sunday','시작 찬양 1','song.pptx',song.read_bytes())
    data=store.get(body['date'],'sunday')['data'];data['fields']['시작 찬양 1']=value('검증곡');data['fields']['시작 찬양 2']=value('미확보곡')
    body['revision']=store.save(body['date'],'sunday',1,data)['revision'];body['groups'].insert(0,['시작 찬양',['시작 찬양 1','시작 찬양 2']])
    result=generate(store,body);directory=Path(store.path).with_suffix('.generated');report=json.loads((directory/(result['file_id']+'.json')).read_text());deck=Presentation(directory/(result['file_id']+'.pptx'))
    entry=next(e for e in report['mapping'] if e['key']=='시작 찬양 1');assert entry['end']-entry['start']==2
    texts=['\n'.join(s.text for s in slide.shapes if s.has_text_frame) for slide in deck.slides]
    assert [texts[n-1] for n in range(entry['start'],entry['end']+1)]==['원본 1절','원본 후렴','원본 아멘']
    blank=next(e for e in report['mapping'] if e['key']=='시작 찬양 2');assert texts[blank['start']-1]==''
    assert '시작 찬양 2:score_asset' in result['unresolved']


def test_bulletin_requires_explicit_issue_and_does_not_promote_cell_group(tmp_path):
    store,body=setup(tmp_path)
    initial=generate_bulletin(store,body)
    assert initial['issue_number']=='13-41'
    initial_report=json.loads((Path(store.path).with_suffix('.generated')/(initial['file_id']+'.json')).read_text())
    assert initial_report['effective_fields']['이번 주 1부 기도']['value']=='인도자'
    assert initial_report['effective_fields']['다음 주 1부 기도']['value']=='인도자'
    assert '이번 주 1부 기도' not in initial['unresolved']
    assert '다음 주 1부 기도' not in initial['unresolved']
    assert '확정 주보 호수' not in store.get(body['date'],'sunday')['data']['fields']
    current=store.get(body['date'],'sunday');data=current['data'];data['fields']['확정 주보 호수']=value('검증용')
    data['fields']['목장 본문']=value('단 1:8~9');data['fields']['목장 제목']=value('연결 시험')
    for i in range(1,5):data['fields']['질문 '+str(i)]=value('확인 질문 '+str(i))
    body['revision']=store.save(body['date'],'sunday',1,data)['revision']
    result=generate_bulletin(store,body);directory=Path(store.path).with_suffix('.generated');pdf=PdfReader(directory/(result['file_id']+'.pdf'))
    assert len(pdf.pages)==2
    assert float(pdf.pages[0].mediabox.width)>float(pdf.pages[0].mediabox.height)
    assert '다음 주 1부 기도' not in result['unresolved']
    assert '13-41' in pdf.pages[0].extract_text() and '연결 시험' in pdf.pages[1].extract_text().replace('\xa0',' ')
    assert store.get(body['date'],'sunday')['data']['fields']['설교 제목']['value']=='연결 시험'


def test_hwpx_preserves_coordinates_and_merged_spans(tmp_path):
    file=tmp_path/'notice.hwpx'
    xml='''<section><p><run><t>10/11</t></run></p><tbl rowCnt="2" colCnt="2"><tr><tc><cellAddr rowAddr="0" colAddr="0"/><cellSpan rowSpan="2" colSpan="1"/><subList><p><run><t>이번 주</t></run></p></subList></tc></tr></tbl></section>'''
    with ZipFile(file,'w') as z:z.writestr('Contents/section0.xml',xml)
    doc=extract_notice(file.name,file.read_bytes());assert doc['tables'][0]['cells'][0]['rowspan']==2


def test_original_parser_with_structured_table_never_assigns_role_as_person():
    node=shutil.which('node')
    if not node:pytest.skip('Node.js required for browser parser regression')
    parser=ROOT/'src/media_automation/web/original/ui/notice-parser.js'
    code='''const p=require(process.argv[1]);const raw='10/11\\n찬양 : 79, 337, 382장\\n<목장 말씀 나누기>\\n<단 1:8~9 / 목장 제목>';const table={rows:3,columns:4,cells:[{row:0,column:2,rowspan:1,colspan:1,text:'기 도'},{row:0,column:3,rowspan:1,colspan:1,text:'봉헌기도'},{row:1,column:0,rowspan:2,colspan:1,text:'이번 주'},{row:1,column:1,rowspan:1,colspan:1,text:'1부'},{row:2,column:1,rowspan:1,colspan:1,text:'2부'},{row:1,column:2,rowspan:1,colspan:1,text:'인도자'},{row:2,column:2,rowspan:1,colspan:1,text:'검증 기도자'},{row:2,column:3,rowspan:1,colspan:1,text:'검증 봉헌자'}]};console.log(JSON.stringify(p.parseBulletin(raw,[table])));'''
    result=json.loads(subprocess.check_output([node,'-e',code,str(parser)],text=True))
    assert result['fields']['이번 주 1부 기도']['value']=='인도자'
    assert result['fields']['2부 기도 담당자']['value']=='검증 기도자'
    assert result['fields']['설교 제목']['value']=='목장 제목'
    assert result['fields']['성경 본문']['value']=='단 1:8~9'
    assert [result['fields'][k]['value'] for k in ['별도 찬송','봉헌 찬송','결단 찬송']]==['79장','337장','382장']
    assert '시작 찬양 1' not in result['fields']


def test_original_ui_guards_dates_existing_fields_and_issue_inference():
    html=(ROOT/'src/media_automation/web/original/ui/index.html').read_text()
    assert 'apply.disabled=dateMismatch' in html
    assert "c.checked=!old||old.state==='UNSET'" in html
    assert '날짜로 계산한 주보 호수:' in html


def test_zoom_embeds_five_exact_videos_and_four_explicit_audio_files(tmp_path, monkeypatch):
    from PIL import Image
    poster=tmp_path/"video-frame.png"
    Image.new("RGB",(160,90),"navy").save(poster)
    monkeypatch.setattr("media_automation.ppt.video_poster.video_poster",lambda source:poster)
    template_dir=tmp_path/'templates';template_dir.mkdir()
    start=Presentation();start.slide_width=12192000;start.slide_height=6858000
    start.slides.add_slide(start.slide_layouts[6]);start.save(template_dir/'friday-zoom-start.pptx')
    monkeypatch.setenv('MEDIA_WORSHIP_TEMPLATES',str(template_dir))
    store,body=setup(tmp_path,'friday','zoom')
    fields=store.get(body['date'],'friday')['data']['fields']
    video_bytes=b'\x00\x00\x00\x18ftypmp42'+b'fixture-video-payload'
    audio_bytes=b'ID3'+b'fixture-audio-payload'
    labels=[f'찬양 {i}' for i in range(1,6)]
    for label in labels:
        fields[label]=value('시험 곡');store.attach(body['date'],'friday',label,label+'.mp4',video_bytes)
    prayers=['첫 기도 제목 1','말씀 관련 기도 제목 1','공동체·중보 기도 제목 1']
    for label in prayers:fields[label]=value('시험 기도')
    for label in ('배경음악 예배 준비','배경음악 첫 기도','배경음악 말씀 기도','배경음악 공동체 기도'):
        store.attach(body['date'],'friday',label,label+'.mp3',audio_bytes)
    data=store.get(body['date'],'friday')['data'];data['fields']=fields
    body['revision']=store.save(body['date'],'friday',1,data)['revision']
    body['groups']=[['찬양',[labels[0],labels[1]]],['첫 기도',[prayers[0]]],['찬양',[labels[2]]],
        ['말씀',['성경 본문','설교 제목']],['응답',[labels[3]]],['말씀 기도',[prayers[1]]],
        ['찬양',[labels[4]]],['공동체 기도',[prayers[2]]]]
    result=generate(store,body);directory=Path(store.path).with_suffix('.generated')
    report=json.loads((directory/(result['file_id']+'.json')).read_text())
    with ZipFile(directory/(result['file_id']+'.pptx')) as z:
        # OPC may deduplicate identical bytes. Validate five movie objects and four audio objects.
        xml=[z.read(n) for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')]
        assert sum(b'<a:videoFile' in blob for blob in xml)==5
        assert sum(b'<a:audioFile' in blob for blob in xml)==4
        assert any(z.read(n)==video_bytes for n in z.namelist() if n.startswith('ppt/media/'))
        assert any(z.read(n)==audio_bytes for n in z.namelist() if n.startswith('ppt/media/'))
    assert all(Path(n).suffix in ('.mp4', '.mp3') for n in z.namelist() if n.startswith('ppt/media/media'))
    assert report['background_music']['actual_playback_verified'] is False
    assert len(report['background_music']['inserted'])==4
    for audio in report['background_music']['inserted']:
        first=next(m for m in report['mapping'] if m['start']==audio['page'])
        assert first['kind']!='blank'
    assert any(m['key']=='개인 기도' for m in report['mapping'])


def test_http_build_download_and_report_use_saved_original_fields(tmp_path):
    store,body=setup(tmp_path);server=make_server(store,0)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    base=f'http://127.0.0.1:{server.server_port}'
    try:
        result=json.load(urlopen(Request(base+'/api/build',data=json.dumps(body).encode(),headers={'Content-Type':'application/json'}),timeout=15))
        blob=urlopen(base+'/api/generated?id='+result['file_id']).read()
        file=tmp_path/'download.pptx';file.write_bytes(blob);assert len(Presentation(file).slides)==result['slide_count']
        report=json.load(urlopen(base+'/api/generated-report?id='+result['file_id']))
        assert report['fields']==store.get(body['date'],'sunday')['data']['fields']
        assert report['complete'] is False
    finally:server.shutdown();server.server_close();thread.join()


def test_store_migration_preserves_records_original_bytes_and_rejects_overwrite(tmp_path):
    from media_automation.web.original.migration import snapshot_store
    source=tmp_path/'old.db';old=WorkStore(source)
    old.save('2026-09-27','sunday',0,{'fields':{'설교 제목':value('보존 데이터')}})
    old.attach('2026-09-27','sunday','검증','a.mp3',b'ID3-original')
    before=source.read_bytes();target=tmp_path/'new.db';snapshot_store(source,target)
    new=WorkStore(target)
    assert new.get('2026-09-27','sunday')==old.get('2026-09-27','sunday')
    assert new.get('2026-10-11','sunday')['data']=={'fields': {'폐회 기도 담당자': value('이은철 목사')}}
    assert source.read_bytes()==before
    with pytest.raises(ValueError):snapshot_store(source,target)

@pytest.mark.parametrize('change',['fields','choices','attachments','unresolved','review','file','typography'])
def test_completion_rejects_changed_or_unreviewed_output(tmp_path,change):
    from media_automation.web.original.generation import render_configuration, SUNDAY_TYPOGRAPHY_VERSION
    store,body=setup(tmp_path)
    data=store.get(body['date'],'sunday')['data'];identifier='a'*32
    data['outputs']['ppt']={'fileId':identifier,'changed':False,'complete':False}
    directory=Path(store.path).with_suffix('.generated');directory.mkdir()
    report={'date':body['date'],'service':'sunday','format':'pptx','typography_version':SUNDAY_TYPOGRAPHY_VERSION,'fields':data['fields'],
            'attachments':[],'unresolved':[],'render_configuration':render_configuration(data)}
    (directory/(identifier+'.pptx')).write_bytes(b'fixture')
    # A fixture report isolates completion checks from the independently tested renderer.
    (directory/(identifier+'.json')).write_text(json.dumps(report))
    if change=='fields':data['fields']['설교 제목']=value('수정 후')
    if change=='choices':data['assetChoices']={'찬양 1':'changed'}
    if change=='attachments':store.attach(body['date'],'sunday','찬양 1','fixture.mp3',b'ID3')
    if change=='unresolved':
        report['unresolved']=['미해결'];(directory/(identifier+'.json')).write_text(json.dumps(report))
    if change=='typography':
        report['typography_version']='old';(directory/(identifier+'.json')).write_text(json.dumps(report))
    if change=='file':(directory/(identifier+'.pptx')).unlink()
    revision=store.save(body['date'],'sunday',1,data)['revision']
    server=make_server(store,0);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:
        request={'date':body['date'],'service':'sunday','revision':revision,'file_id':identifier,'reviewed':change!='review'}
        with pytest.raises(HTTPError) as error:
            urlopen(Request(f'http://127.0.0.1:{server.server_port}/api/complete',data=json.dumps(request).encode(),headers={'Content-Type':'application/json'}))
        assert error.value.code==400
        assert not store.get(body['date'],'sunday')['data']['outputs']['ppt']['complete']
    finally:server.shutdown();server.server_close();thread.join()


def test_completion_accepts_matching_reviewed_output(tmp_path):
    from media_automation.web.original.generation import render_configuration, SUNDAY_TYPOGRAPHY_VERSION
    store,body=setup(tmp_path);data=store.get(body['date'],'sunday')['data'];identifier='b'*32
    data['outputs']['ppt']={'fileId':identifier,'changed':False,'complete':False}
    directory=Path(store.path).with_suffix('.generated');directory.mkdir()
    (directory/(identifier+'.pptx')).write_bytes(b'fixture')
    (directory/(identifier+'.json')).write_text(json.dumps({'date':body['date'],'service':'sunday','format':'pptx',
        'typography_version':SUNDAY_TYPOGRAPHY_VERSION,'fields':data['fields'],'attachments':[],'unresolved':[],'render_configuration':render_configuration(data)}))
    revision=store.save(body['date'],'sunday',1,data)['revision']
    server=make_server(store,0);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:
        request=dict(date=body['date'],service='sunday',revision=revision,file_id=identifier,reviewed=True)
        result=json.load(urlopen(Request(f'http://127.0.0.1:{server.server_port}/api/complete',data=json.dumps(request).encode(),headers={'Content-Type':'application/json'})))
        assert result['data']['outputs']['ppt']['complete'] is True
    finally:server.shutdown();server.server_close();thread.join()


def test_library_asset_upload_never_assigns_a_weekly_slot(tmp_path):
    store=WorkStore(tmp_path/'work.db');song=store.register_song({'title':'검증곡'})
    result=store.attach_song_asset(song['id'],'original.mp3',b'ID3fixture')
    assert store.songs()[0]['assets'][0]['digest']==result['digest']
    assert store.attachments('2026-10-11','sunday')==[]
    assert not store.songs()[0]['assets'][0]['eligible']
    assert store.original(result['digest'])[0].read_bytes()==b'ID3fixture'


def test_pdf_preview_renders_both_imposed_pages(tmp_path):
    from media_automation.web.original.ppt_preview import preview
    bundled=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/pdftoppm'
    if not shutil.which('pdftoppm') and not bundled.exists():pytest.skip('PDF page renderer unavailable')
    store,body=setup(tmp_path);data=store.get(body['date'],'sunday')['data'];data['fields']['확정 주보 호수']=value('검증용')
    body['revision']=store.save(body['date'],'sunday',1,data)['revision'];result=generate_bulletin(store,body)
    for page in (1,2):assert preview(store,result['file_id'],page).startswith(b'\x89PNG\r\n\x1a\n')
    with pytest.raises(ValueError):preview(store,result['file_id'],3)


def test_original_score_preview_exposes_all_pages_as_images(tmp_path):
    from reportlab.pdfgen import canvas
    from pptx import Presentation
    store=WorkStore(tmp_path/'work.db')
    song=store.register_song({'title':'검증곡'})
    deck=tmp_path/'score.pptx';Presentation().save(deck)
    asset=store.attach_song_asset(song['id'],'score.pptx',deck.read_bytes())
    cache=Path(store.path).with_suffix('.previews');cache.mkdir()
    pdf=cache/(asset['digest']+'.pdf');c=canvas.Canvas(str(pdf))
    for number in range(3):c.drawString(72,72,str(number));c.showPage()
    c.save()
    server=make_server(store,0);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    endpoint=f'http://127.0.0.1:{server.server_port}/api/asset-preview?digest='+asset['digest']
    try:
        assert json.load(urlopen(endpoint+'&info=1'))['page_count']==3
        for page in (1,3):assert urlopen(endpoint+'&page='+str(page)).read().startswith(b'\x89PNG\r\n\x1a\n')
        with pytest.raises(HTTPError):urlopen(endpoint+'&page=4')
    finally:server.shutdown();server.server_close();thread.join()


def test_bulletin_restores_original_cover_font_and_artwork(tmp_path):
    import pymupdf
    store,body=setup(tmp_path);data=store.get(body['date'],'sunday')['data'];data['fields']['확정 주보 호수']=value('검증용')
    body['revision']=store.save(body['date'],'sunday',1,data)['revision'];result=generate_bulletin(store,body)
    doc=pymupdf.open(Path(store.path).with_suffix('.generated')/(result['file_id']+'.pdf'))
    spans=[s for b in doc[0].get_text('dict')['blocks'] for line in b.get('lines',[]) for s in line['spans']]
    slogan=next(s for s in spans if ''.join(s['text'].split())=='영혼구원하여제자삼는교회')
    assert slogan['font']=='AppleSDGothicNeo-Bold' and slogan['size']==pytest.approx(23)
    assert doc[0].get_images()
    # A missing current question or serving name must clear the reference week.
    text=doc[1].get_text()
    assert '다니엘이뜻을정한것처럼' not in ''.join(text.split())
    assert '장덕수' not in text and '최진용' not in text


def test_latest_bulletin_aligns_shared_serving_and_four_column_meetings(tmp_path):
    store,body=setup(tmp_path);data=store.get(body['date'],'sunday')['data']
    for label,text in [('확정 주보 호수','검증용'),('설교 제목','믿음과 상황이 충돌할 때'),('목장 제목','믿음과 상황이 충돌할 때'),('목장 본문','단 1:8~9'),
                       ('이번 주 1부 기도','첫기도'),('2부 기도 담당자','둘기도'),('이번 주 설거지','공통설거지')]:
        data['fields'][label]=value(text)
    body['revision']=store.save(body['date'],'sunday',1,data)['revision'];result=generate_bulletin(store,body)
    pdf=PdfReader(Path(store.path).with_suffix('.generated')/(result['file_id']+'.pdf'))
    import pymupdf
    path=Path(store.path).with_suffix('.generated')/(result['file_id']+'.pdf')
    doc=pymupdf.open(path)
    assert len(doc)==2
    inner=doc[1]
    title=inner.search_for('믿음')[0]
    reference=inner.search_for('단1:8~9')[0]
    assert title.y0>reference.y0 # original puts scripture above sermon
    first=next(r for r in inner.search_for('인도자') if r.y0>380);second=next(r for r in inner.search_for('둘기도') if r.y0>380)
    dishes=inner.search_for('공통설거지')[0]
    assert dishes.y0==pytest.approx((first.y0+second.y0)/2,abs=1)
    assert doc[0].get_images() # original cover artwork preserved
    assert '오전예배순서' in inner.get_text().replace(' ','')
    assert 'AppleSDGothicNeo' in str(inner.get_fonts())
    assert '장덕수' not in inner.get_text() # no previous offering prayer remains


def test_setting_change_flags_pending_bulletin_and_preserves_completed_history(tmp_path):
    store,body=setup(tmp_path);data=store.get(body['date'],'sunday')['data'];data['fields']['확정 주보 호수']=value('검증용')
    body['revision']=store.save(body['date'],'sunday',1,data)['revision'];result=generate_bulletin(store,body)
    data['outputs']['pdf']={'fileId':result['file_id'],'generated':True,'complete':False,'changed':False}
    pending=store.save(body['date'],'sunday',body['revision'],data)
    history=json.loads(json.dumps(data));history['outputs']['pdf']['complete']=True
    store.save('2026-10-04','sunday',0,history)
    pdf=Path(store.path).with_suffix('.generated')/(result['file_id']+'.pdf');original=pdf.read_bytes()
    server=make_server(store,0);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    base=f'http://127.0.0.1:{server.server_port}'
    try:
        settings=json.load(urlopen(base+'/api/bulletin-settings'))
        settings['data']['church']['전화']='검증 연락처'
        from PIL import Image
        from io import BytesIO
        import base64
        stream=BytesIO();Image.effect_noise((512,512),100).convert('RGB').save(stream,format='PNG')
        settings['data']['cover_photo']={'filename':'qa.png','data':'data:image/png;base64,'+base64.b64encode(stream.getvalue()).decode()}
        assert len(json.dumps(settings))>65536
        urlopen(Request(base+'/api/bulletin-settings',data=json.dumps(settings).encode(),headers={'Content-Type':'application/json'})).read()
        assert json.load(urlopen(base+'/api/bulletin-settings'))['data']['cover_photo']==settings['data']['cover_photo']
        current=store.get(body['date'],'sunday')
        assert current['revision']==pending['revision']+1 and current['data']['outputs']['pdf']['changed']
        assert current['data']['fields']==data['fields']
        assert store.get('2026-10-04','sunday')['data']['outputs']['pdf']==history['outputs']['pdf']
        assert pdf.read_bytes()==original
        with pytest.raises(RevisionConflict):store.save(body['date'],'sunday',pending['revision'],data)
    finally:server.shutdown();server.server_close();thread.join()


@pytest.mark.parametrize('problem',['missing_file','old_layout','settings_change'])
def test_bulletin_completion_rejects_invalid_current_artifact(tmp_path,problem):
    from media_automation.bulletin.original_design import LAYOUT_VERSION
    store,body=setup(tmp_path);data=store.get(body['date'],'sunday')['data'];identifier='c'*32
    data['outputs']['pdf']={'fileId':identifier,'generated':True,'complete':False,'changed':False}
    directory=Path(store.path).with_suffix('.generated');directory.mkdir()
    file=directory/(identifier+'.pdf');file.write_bytes(b'fixture')
    settings=json.loads(Path('src/media_automation/web/original/bulletin-base-reference.json').read_text())
    report={'layout_version':LAYOUT_VERSION,'format':'pdf','date':body['date'],'service':'sunday','fields':data['fields'],'unresolved':[],'settings':settings}
    if problem=='missing_file':file.unlink()
    elif problem=='old_layout':report['layout_version']='earlier-design'
    else:report['settings']=dict(settings,status='different')
    (directory/(identifier+'.json')).write_text(json.dumps(report))
    revision=store.save(body['date'],'sunday',1,data)['revision']
    server=make_server(store,0);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:
        req=Request(f'http://127.0.0.1:{server.server_port}/api/complete-bulletin',data=json.dumps(dict(date=body['date'],service='sunday',revision=revision,file_id=identifier,reviewed=True)).encode(),headers={'Content-Type':'application/json'})
        with pytest.raises(HTTPError):urlopen(req)
        assert not store.get(body['date'],'sunday')['data']['outputs']['pdf']['complete']
    finally:server.shutdown();server.server_close();thread.join()


def test_sunday_people_are_split_in_bulletin_and_second_service_only_in_ppt():
    from media_automation.bulletin.document import BulletinWorshipPage
    from media_automation.bulletin.worship_page import build_worship_order_rows
    from media_automation.web.original.production_order import build_order
    page=BulletinWorshipPage(praise_first='첫찬양',praise_second='둘찬양',first_service_prayer='인도자',second_service_prayer='둘기도',
        first_service_offering_prayer='첫봉헌',second_service_offering_prayer='둘봉헌')
    rows=dict(build_worship_order_rows(page))
    assert rows['경배와 찬양']=='1부 첫찬양 / 2부 둘찬양'
    assert rows['기 도']=='1부 인도자 / 2부 둘기도'
    assert rows['봉 헌 기 도']=='1부 첫봉헌 / 2부 둘봉헌'
    fields={key:value(text) for key,text in [('주일 인도자','인도자'),('1부 찬양 인도자','첫찬양'),('2부 찬양 인도자','둘찬양'),
        ('이번 주 1부 기도','인도자'),('2부 기도 담당자','둘기도'),('이번 주 1부 봉헌기도','첫봉헌'),('2부 봉헌기도 담당자','둘봉헌')]}
    groups=[['주보 예배·섬김',['주일 인도자','1부 찬양 인도자','2부 찬양 인도자']],
        ['이번 주 섬김 · 주보',['이번 주 1부 기도','이번 주 1부 봉헌기도']],['기도와 봉헌',['2부 기도 담당자','2부 봉헌기도 담당자']]]
    blocks=[b for b in build_order(groups,fields,'sunday') if b['kind']!='blank']
    assert [(b['key'],b['content']) for b in blocks]==[('2부 기도 담당자','기도\n둘기도'),('2부 봉헌기도 담당자','봉헌기도\n둘봉헌')]


def test_sunday_fixed_closing_prayer_default_preserves_saved_override(tmp_path):
    from media_automation.web.original.work_store import WorkStore
    store = WorkStore(tmp_path / 'work.db')
    fresh = store.get('2026-10-18', 'sunday')
    assert fresh['data']['fields']['폐회 기도 담당자'] == {'state': 'VALUE', 'value': '이은철 목사'}
    override = {'fields': {'폐회 기도 담당자': {'state': 'VALUE', 'value': '별도 지정'}}}
    store.save('2026-10-18', 'sunday', 0, override)
    assert store.get('2026-10-18', 'sunday')['data'] == override
    assert store.get('2026-10-14', 'wednesday')['data'] == {}


def test_unset_starting_song_stays_blank_and_reports_input_pending(tmp_path, monkeypatch):
    monkeypatch.setenv("MEDIA_WORSHIP_TEMPLATES",str(tmp_path/"no-templates"))
    store, body = setup(tmp_path)
    body['groups'] = [['시작 찬양', ['시작 찬양 1']]]
    result = generate(store, body)
    assert result['unresolved'] == ['예배 전 안내:template','시작 찬양 1:input']
    deck = Presentation(Path(store.path).with_suffix('.generated') / (result['file_id'] + '.pptx'))
    assert not any(shape.has_text_frame and shape.text.strip() for shape in deck.slides[-1].shapes)


def test_https_proxy_requires_auth_and_exact_origin(tmp_path):
    password = 'isolated-test-password-16plus'
    store=WorkStore(tmp_path/'proxy.db')
    server=make_server(store,0,password,public_origin='https://media.example.test')
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    base=f'http://127.0.0.1:{server.server_port}'
    headers={'Host':'media.example.test','Origin':'https://media.example.test','X-Forwarded-Proto':'https','Content-Type':'application/json'}
    try:
        def request(path, extra=None, data=None):
            return urlopen(Request(base+path,headers={**headers,**(extra or {})},data=json.dumps(data).encode() if data else None))
        with pytest.raises(HTTPError) as failure:request('/api/work?date=2026-10-09&service=friday')
        assert failure.value.code==401
        for extra in [{'Origin':'https://other.example.test'}, {'Host':'other.example.test'}, {'X-Forwarded-Proto':'http'}]:
            with pytest.raises(HTTPError) as failure:request('/api/mode',extra)
            assert failure.value.code==403
        login=request('/api/login',data={'password':password});cookie=login.headers['Set-Cookie']
        assert all(flag in cookie for flag in ['Secure','HttpOnly','SameSite=Strict'])
        assert json.load(request('/api/work?date=2026-10-09&service=friday',{'Cookie':cookie.split(';')[0]}))['revision']==0
    finally:server.shutdown();server.server_close();thread.join()

def test_https_proxy_rejects_unsafe_startup_configuration(tmp_path):
    store=WorkStore(tmp_path/'proxy.db')
    for origin,password in [('http://media.example.test','a'*20), ('https://user@media.example.test','a'*20), ('https://media.example.test/path','a'*20), ('https://media.example.test',None), ('https://media.example.test','short')]:
        with pytest.raises(ValueError):make_server(store,0,password,public_origin=origin)

def test_backup_command_preserves_fields_and_original_files(tmp_path):
    from media_automation.web.backup import backup_store
    store=WorkStore(tmp_path/'work.db');data={'fields':{'설교 제목':value('보존 검증')}}
    store.save('2026-10-09','friday',0,data);store.attach('2026-10-09','friday','시작 찬양 1','original.mp4',b'original media')
    first=backup_store(store.path,tmp_path/'backups');second=backup_store(store.path,tmp_path/'backups')
    assert first!=second
    assert WorkStore(first).get('2026-10-09','friday')['data']==data
    assert WorkStore(first).attachments('2026-10-09','friday')==store.attachments('2026-10-09','friday')
    digest=store.attachments('2026-10-09','friday')[0]['digest']
    assert first.with_suffix('.files').joinpath(digest).read_bytes()==b'original media'


def test_concurrent_build_is_rejected_then_released(tmp_path, monkeypatch):
    import media_automation.web.original.server as module
    from concurrent.futures import ThreadPoolExecutor
    entered=threading.Event();release=threading.Event()
    def slow_generate(store, body):
        entered.set();assert release.wait(3);return {'test_generated':True}
    monkeypatch.setattr(module,'generate',slow_generate)
    server=make_server(WorkStore(tmp_path/'work.db'),0)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    url=f'http://127.0.0.1:{server.server_port}/api/build'
    def build():return json.load(urlopen(Request(url,data=b'{}',headers={'Content-Type':'application/json'}),timeout=4))
    try:
        with ThreadPoolExecutor(max_workers=2) as pool:
            pending=pool.submit(build);assert entered.wait(2)
            with pytest.raises(HTTPError) as failure:build()
            assert failure.value.code==409
            release.set();assert pending.result()['test_generated']
            assert build()['test_generated']
    finally:release.set();server.shutdown();server.server_close();thread.join()


def test_bulletin_hymn_number_with_title_fits_original_number_slot():
    from media_automation.web.original.bulletin_generation import bulletin_hymn_label
    assert bulletin_hymn_label('430장 주와 같이 길 가는 것')=='430장'
    assert bulletin_hymn_label('새찬송가 337장 내 모든 시험 무거운 짐을')=='337장'
    assert bulletin_hymn_label('주님 말씀하시면')=='주님 말씀하시면'
    assert bulletin_hymn_label(None) is None


def test_structured_key_value_notice_table_preserves_blank_and_multiline():
    node=shutil.which('node')
    if not node:pytest.skip('Node.js required for browser parser regression')
    parser=ROOT/'src/media_automation/web/original/ui/notice-parser.js'
    code='''const p=require(process.argv[1]);const table={cells:[{row:0,column:0,colspan:2,text:'본 문'},{row:0,column:2,text:'시 24:3~5'},{row:1,column:0,colspan:2,text:'설교 제목'},{row:1,column:2,text:'새 제목'},{row:2,column:0,colspan:2,text:'추가 말씀'},{row:2,column:2,text:''},{row:3,column:0,colspan:2,text:'교회 소식'},{row:3,column:2,text:'첫 소식\\n둘째 소식'}]};console.log(JSON.stringify(p.parseTableFields([table],['성경 본문','설교 제목','추가 말씀','교회 소식'],{'본문':'성경 본문'},'sunday')));'''
    result=json.loads(subprocess.check_output([node,'-e',code,str(parser)],text=True))
    assert result['fields']['성경 본문']==value('시 24:3~5')
    assert result['fields']['추가 말씀']=={'state':'NONE','value':''}
    assert result['fields']['교회 소식']['value']=='첫 소식\n둘째 소식'


def test_ui_new_week_normalizes_partial_server_record_without_erasing_values():
    node=shutil.which('node')
    if not node:pytest.skip('Node.js required for browser regression')
    html=ROOT/'src/media_automation/web/original/ui/index.html'
    code='''const fs=require('fs');const src=fs.readFileSync(process.argv[1],'utf8');const recordFunction='function record()'+src.split('function record()')[1].split('function msg')[0];const key=()=>'new';const service='sunday';let records={new:{fields:{'폐회 기도 담당자':{state:'VALUE',value:'이은철 목사'}}}};eval(recordFunction);let r=record();if(!r.outputs.ppt||!r.outputs.pdf||!r.sourceValues)throw Error('partial record');if(r.fields['1부 찬양 인도자'].value!=='이광민'||r.fields['2부 찬양 인도자'].value!=='장기훈')throw Error('missing defaults');r.fields['2부 찬양 인도자'].value='직접 수정';if(record().fields['2부 찬양 인도자'].value!=='직접 수정')throw Error('override lost');r.outputs.pdf.fileId='preserved';if(record().outputs.pdf.fileId!=='preserved'||record().fields['폐회 기도 담당자'].value!=='이은철 목사')throw Error('lost stored data');console.log('ok');'''
    assert subprocess.check_output([node,'-e',code,str(html)],text=True).strip()=='ok'


def test_sunday_news_is_bulletin_only_even_for_older_clients():
    from media_automation.web.original.production_order import build_order
    blocks=build_order([['기도',['2부 기도 담당자','교회 소식']]],{'2부 기도 담당자':value('검증 담당자'),'교회 소식':value('주보 소식')},'sunday')
    assert not any(b['key']=='교회 소식' for b in blocks)


def test_sunday_shared_passage_rejects_saved_conflicting_title(tmp_path):
    store,body=setup(tmp_path)
    r=store.get(body['date'],'sunday')
    r['data']['fields']['확정 주보 호수']=value('검증용')
    r['data']['fields']['목장 제목']=value('서로 다른 제목')
    body['revision']=store.save(body['date'],'sunday',r['revision'],r['data'])['revision']
    with pytest.raises(ValueError,match='값이 다릅니다'):generate_bulletin(store,body)
    assert store.get(body['date'],'sunday')['data']['fields']['목장 제목']['value']=='서로 다른 제목'


def test_serving_table_merged_cells_resolve_both_weeks(tmp_path):
    node=shutil.which('node')
    if not node:pytest.skip('Node.js required')
    cells=[]
    def cell(r,c,text,rs=1,cs=1):cells.append(dict(row=r,column=c,text=text,rowspan=rs,colspan=cs))
    cell(0,0,'구 분',cs=2)
    for c,t in enumerate(['기 도','봉 헌 기 도','설 거 지','수요예배 기도'],2):cell(0,c,t)
    for row,week in [(1,'이 번 주'),(3,'다 음 주')]:
        cell(row,0,week,rs=2)
        for r,part in [(row,'1 부'),(row+1,'2 부')]:
            cell(r,1,part);cell(r,2,'인도자' if part=='1 부' else f'기도{r}');cell(r,3,f'봉헌{r}')
        cell(row,4,f'설거지{row}',rs=2);cell(row,5,f'수요{row}',rs=2)
    code="const p=require(process.argv[1]);const t=JSON.parse(process.argv[2]);console.log(JSON.stringify(p.parseBulletin('10/11\\n<목장 말씀 나누기>\\n<단 1:8~9 / 제목>',[t])));"
    r=json.loads(subprocess.check_output([node,'-e',code,str(ROOT/'src/media_automation/web/original/ui/notice-parser.js'),json.dumps(dict(rows=5,columns=6,cells=cells))],text=True))
    for prefix,row in [('이번 주',1),('다음 주',3)]:
        assert r['fields'][prefix+' 1부 기도']['value']=='인도자'
        assert r['fields'][prefix+' 설거지']['value']==f'설거지{row}'
        assert r['fields'][prefix+' 수요예배 기도']['value']==f'수요{row}'
    assert r['fields']['2부 기도 담당자']['value']=='기도2'
    assert r['fields']['다음 주 2부 기도']['value']=='기도4'


def test_pending_refresh_never_replaces_a_new_local_input():
    node=shutil.which('node')
    if not node:pytest.skip('Node.js required')
    code="""const fs=require('fs');const src=fs.readFileSync(process.argv[1],'utf8');const fn='async function loadShared'+src.split('async function loadShared')[1].split('async function flushShared')[0];const st={loaded:true,revision:1,dirty:false};const syncState=()=>st,endpoint=()=>'',key=()=>'k',syncLabel=()=>{};let saveOk=false;let records={k:{fields:{title:'local'}}};let release;const fetch=()=>new Promise(r=>release=r);const document={activeElement:{matches:()=>true}};const render=()=>{throw Error('input replaced');};eval(fn);(async()=>{let pending=loadShared('k');st.dirty=true;release({ok:true,json:async()=>({revision:2,data:{fields:{title:'remote'}}})});await pending;if(records.k.fields.title!=='local'||st.revision!==1)throw Error('lost local edit');console.log('ok');})();"""
    assert subprocess.check_output([node,'-e',code,str(ROOT/'src/media_automation/web/original/ui/index.html')],text=True).strip()=='ok'
