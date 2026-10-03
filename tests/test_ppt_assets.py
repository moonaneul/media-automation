from pptx import Presentation

from media_automation.ppt import (
    FilePreServiceSlideProvider,
)


def make_ppt(
    path,
    slide_count,
):
    prs = Presentation()

    layout = prs.slide_layouts[6]

    for _ in range(
        slide_count
    ):
        prs.slides.add_slide(
            layout
        )

    prs.save(
        path
    )


def test_pre_service_provider_keeps_requested_slides(
    tmp_path,
):
    source = (
        tmp_path
        / "source.pptx"
    )

    make_ppt(
        source,
        6,
    )

    provider = (
        FilePreServiceSlideProvider(
            {
                "wednesday": source,
            },
            slide_counts={
                "wednesday": 4,
            },
        )
    )

    prs = (
        provider
        .create_presentation(
            "wednesday"
        )
    )

    assert len(
        prs.slides
    ) == 4


def test_pre_service_provider_rejects_unknown_key(
    tmp_path,
):
    provider = (
        FilePreServiceSlideProvider(
            {}
        )
    )

    try:
        provider.create_presentation(
            "wednesday"
        )
    except KeyError:
        pass
    else:
        raise AssertionError(
            "KeyError expected"
        )


def test_pre_service_provider_rejects_missing_file(
    tmp_path,
):
    provider = (
        FilePreServiceSlideProvider(
            {
                "wednesday": (
                    tmp_path
                    / "missing.pptx"
                )
            }
        )
    )

    try:
        provider.create_presentation(
            "wednesday"
        )
    except FileNotFoundError:
        pass
    else:
        raise AssertionError(
            "FileNotFoundError expected"
        )