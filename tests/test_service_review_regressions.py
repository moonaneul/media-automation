import json
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation

from media_automation.web.original.production_order import build_order, combine_friday_prayers
from media_automation.web.original.work_store import WorkStore
from media_automation.web.original.generation import generate


def value(text):
    return {'state': 'VALUE', 'value': text}


@pytest.mark.parametrize('service,prayer,notice', [
    ('wednesday', '기도 담당자', '광고'),
    ('sunday', '2부 기도 담당자', '교회 소식'),
])
def test_fixed_notice_does_not_require_weekly_input(service, prayer, notice):
    groups=[['기도', [prayer]], ['말씀', ['성경 본문']]]
    fields={prayer:value('검증 담당자'), '성경 본문':value('요 3:16')}
    blocks=[b for b in build_order(groups,fields,service) if b['kind']!='blank']
    assert [b['key'] for b in blocks]==[prayer,'광고 순서','성경 본문']
    assert blocks[1]['content']==notice


def test_table_heading_cannot_be_offering_person():
    with pytest.raises(ValueError,match='표의 제목'):
        build_order([['봉헌',['2부 봉헌기도 담당자']]],
                    {'2부 봉헌기도 담당자':value('설 거 지')},'sunday')


def test_numbered_topics_in_old_single_field_become_separate_screens():
    blocks=combine_friday_prayers([
        dict(kind='input',key='말씀 관련 기도',state='VALUE',
             content='1. 첫째 제목\n첫째 보조 내용\n2. 둘째 제목\n둘째 보조 내용')
    ],{},include_personal=False)
    assert [b['content'] for b in blocks]==[
        '1. 첫째 제목\n첫째 보조 내용','2. 둘째 제목\n둘째 보조 내용']
    assert len({b['key'] for b in blocks})==2


def test_flattened_heading_is_not_imported_as_offering_person():
    parser=Path('src/media_automation/web/original/ui/notice-parser.js').resolve()
    raw='오전 예배 순서\n봉헌기도\n2부\n설 거 지'
    driver='const p=require(process.argv[1]);console.log(JSON.stringify(p.parseBulletin(process.argv[2])));'
    result=json.loads(subprocess.check_output(['node','-e',driver,str(parser),raw],text=True))
    assert '2부 봉헌기도 담당자' not in result['fields']
    assert any('표 제목' in warning for warning in result['unknown'])


@pytest.mark.parametrize('service,mode',[('wednesday',None),('sunday',None),('friday','in_person'),('friday','zoom')])
def test_explicit_external_bible_path_reaches_generated_service(tmp_path,monkeypatch,service,mode):
    # Synthetic text proves wiring only; this is not a verified Bible quotation.
    bible=tmp_path/'external-bible.json'
    bible.write_text(json.dumps({'요3:16':'본문 연결 검증용 문자열'},ensure_ascii=False))
    monkeypatch.setenv('MEDIA_BIBLE_FILE',str(bible))
    monkeypatch.setenv('MEDIA_BIBLE_SOURCES',str(tmp_path/'sources'))
    monkeypatch.setenv('MEDIA_WORSHIP_TEMPLATES',str(tmp_path/'templates'))
    store=WorkStore(tmp_path/'work.db')
    data={'fields':{'성경 본문':value('요 3:16'),'설교 제목':value('연결 검증')}}
    if mode:data['fridayMode']=mode
    revision=store.save('2026-10-11',service,0,data)['revision']
    result=generate(store,dict(date='2026-10-11',service=service,revision=revision,
                             groups=[['말씀',['성경 본문','설교 제목']]]))
    deck=Presentation(Path(store.path).with_suffix('.generated')/(result['file_id']+'.pptx'))
    text=''.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert '본문 연결 검증용 문자열' in text
    assert '성경 본문:bible_text' not in result['unresolved']
    assert '성경 본문:bible_review' in result['unresolved']
