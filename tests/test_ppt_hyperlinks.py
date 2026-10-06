from pptx import Presentation

from media_automation.ppt.base import create_4x3_presentation
from media_automation.ppt.hyperlinks import remove_presentation_hyperlinks


def test_remove_presentation_hyperlinks_removes_shape_click_link(tmp_path):
    path = tmp_path / "hyperlinks.pptx"

    prs = create_4x3_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, 2000000, 1000000)
    shape.text = "클릭"
    shape.click_action.hyperlink.address = "https://example.com/shop"
    prs.save(path)

    removed = remove_presentation_hyperlinks(path)

    assert removed >= 1

    reopened = Presentation(path)
    address = reopened.slides[0].shapes[0].click_action.hyperlink.address
    assert address is None


def test_remove_presentation_hyperlinks_keeps_visible_text_for_qa(tmp_path):
    path = tmp_path / "visible-url.pptx"

    prs = create_4x3_presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, 3000000, 1000000)
    shape.text = "https://example.com/shop"
    prs.save(path)

    removed = remove_presentation_hyperlinks(path)

    assert removed == 0
    reopened = Presentation(path)
    assert "https://example.com/shop" in reopened.slides[0].shapes[0].text
