"""Whole-chapter John 1–3 input must resolve complete existing verses.

The fixtures deliberately contain synthetic, NON-Bible text. This tests
reference expansion and UI production without distributing copyrighted text.
"""
import json
from pathlib import Path

import pytest
from pptx import Presentation

from media_automation.web.original.bible_connection import (
    CHAPTER_ENDS, expand_chapter_reference, resolve,
)
from media_automation.web.original.generation import generate
from media_automation.web.original.work_store import WorkStore


EXPECTED={'요 1:1~51':51, '요 2:1~25':25, '요 3:1~36':36}


def synthetic_json(missing=None):
    rows={}
    for chapter,end in [(1,51),(2,25),(3,36)]:
        for verse in range(1,end+1):
            reference=f'요{chapter}:{verse}'
            if reference!=missing:
                rows[reference]=f'{chapter}장 {verse}절 검증용 모의 문장 (성경 원문 아님)'
    return rows


def bible_block(reference='요 1~3장'):
    return [dict(kind='input',key='성경 본문',
                 state='VALUE',content=reference)]


def lookup_json(tmp_path,ref='요 1~3장',missing=None):
    path=tmp_path/'provided.json'
    path.write_text(json.dumps(synthetic_json(missing),ensure_ascii=False),encoding='utf8')
    return resolve(bible_block(ref),source_dir=tmp_path/'no-official',
                   provided_file=path)


def test_john_one_to_three_chapter_ends_are_explicit():
    assert [CHAPTER_ENDS['요',ch] for ch in (1,2,3)]==[51,25,36]
    assert expand_chapter_reference('요 1~3장')==list(EXPECTED)
    assert expand_chapter_reference('요1~3장')==list(EXPECTED)
    assert expand_chapter_reference('요 2장')==['요 2:1~25']


def test_chapter_expansion_produces_all_112_existing_verses(tmp_path):
    passage=lookup_json(tmp_path)['성경 본문']
    assert passage['source_validated'] is False
    assert passage['text_review_pending'] is True
    assert len(passage['units'])==112
    assert [ref for ref in dict.fromkeys(u['passage_reference'] for u in passage['units'])]==list(EXPECTED)
    actual=[(u['chapter'],u['start_verse']) for u in passage['units']]
    wanted=[(c,n) for c,end in [(1,51),(2,25),(3,36)] for n in range(1,end+1)]
    assert actual==wanted
    assert not any('세례' in u['text'] for u in passage['units'])


def test_missing_actual_verse_does_not_shortcut_range_or_insert_partial(tmp_path):
    resolved=lookup_json(tmp_path,missing='요2:25')['성경 본문']
    assert resolved['units']==[]
    assert '요2:25' in resolved['error']


def test_unsupported_chapter_cannot_guess_json_last_verse(tmp_path):
    result=lookup_json(tmp_path,ref='요 1~4장')['성경 본문']
    assert result['units']==[]
    assert '마지막 절이 등록되지 않았습니다' in result['error']
    assert '요 1~4장' in result['error']


def test_actual_wednesday_generation_has_one_slide_per_verse(tmp_path,monkeypatch):
    file=tmp_path/'provided.json'
    file.write_text(json.dumps(synthetic_json(),ensure_ascii=False),encoding='utf-8')
    monkeypatch.setenv('MEDIA_BIBLE_FILE',str(file))
    monkeypatch.setenv('MEDIA_BIBLE_SOURCES',str(tmp_path/'empty'))
    monkeypatch.setenv('MEDIA_WORSHIP_TEMPLATES',str(tmp_path/'templates'))
    store=WorkStore(tmp_path/'work.db')
    fields={'성경 본문':{'state':'VALUE','value':'요 1~3장'},
            '설교 제목':{'state':'VALUE','value':'테스트 설교'}}
    data={'fields':fields}
    rev=store.save('2026-10-21','wednesday',0,data)['revision']
    result=generate(store,dict(date='2026-10-21',service='wednesday',
                               revision=rev,groups=[['말씀',['성경 본문','설교 제목']]]))
    report=json.loads((Path(store.path).with_suffix('.generated')/(result['file_id']+'.json')).read_text())
    portions=[row for row in report['mapping'] if row['key']=='성경 본문']
    assert len(portions)==1 and portions[0]['status']=='bible_inserted'
    assert portions[0]['end']-portions[0]['start']+1==112
    assert '성경 본문:bible_text' not in result['unresolved']
    assert '성경 본문:bible_review' in result['unresolved']
    assert result['bible_errors']=={}
    path=Path(store.path).with_suffix('.generated')/(result['file_id']+'.pptx')
    presentation=Presentation(path)
    scripture_slide=presentation.slides[portions[0]['start']-1]
    assert any('요 1:1~51' in shape.text for shape in scripture_slide.shapes
               if shape.has_text_frame)


def test_missing_verse_yields_useful_production_report_and_no_partial_text(tmp_path,monkeypatch):
    file=tmp_path/'provided.json'
    file.write_text(json.dumps(synthetic_json(missing='요2:25'),ensure_ascii=False),encoding='utf-8')
    monkeypatch.setenv('MEDIA_BIBLE_FILE',str(file))
    monkeypatch.setenv('MEDIA_BIBLE_SOURCES',str(tmp_path/'empty'))
    monkeypatch.setenv('MEDIA_WORSHIP_TEMPLATES',str(tmp_path/'templates'))
    store=WorkStore(tmp_path/'work.db')
    rev=store.save('2026-10-21','wednesday',0,
                   {'fields':{'성경 본문':{'state':'VALUE','value':'요 1~3장'},
                              '설교 제목':{'state':'VALUE','value':'테스트 설교'}}})['revision']
    result=generate(store,dict(date='2026-10-21',service='wednesday',revision=rev,
                               groups=[['말씀',['성경 본문','설교 제목']]]))
    assert result['bible_errors'] and '요2:25' in result['bible_errors']['성경 본문']
    assert '성경 본문:bible_text' in result['unresolved']
    report=json.loads((Path(store.path).with_suffix('.generated')/
                       (result['file_id']+'.json')).read_text())
    assert report['bible_errors']==result['bible_errors']
    assert next(row for row in report['mapping'] if row['key']=='성경 본문')['status']=='missing_verified_bible_text'


def test_web_shows_real_source_error_not_just_input_saved():
    ui=Path('src/media_automation/web/original/ui/index.html').read_text('utf8')
    server=Path('src/media_automation/web/original/server.py').read_text('utf8')
    assert 'o.bibleErrors=result.bible_errors||{}' in ui
    assert '성경 본문 미삽입 원인:' in ui
    assert "report.bible_errors?.[item.key]" in ui
    assert "'bible_errors':data.get('bible_errors',{})" in server
