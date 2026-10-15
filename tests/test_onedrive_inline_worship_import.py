"""OneDrive inline source selection imports only the user-selected file/slot."""
from io import BytesIO
import json
from pathlib import Path

import pytest
from pptx import Presentation

from media_automation.web.original.work_store import WorkStore
from media_automation.web.original.onedrive import CloudError, OneDrive
from media_automation.web.original.generation import generate


DAY = "2026-10-14"


def make_store(tmp_path, service="wednesday", mode="in_person"):
    store = WorkStore(tmp_path / "review.db")
    fields = {
        "성경 본문": {"state": "VALUE", "value": "단 1:8"},
        "설교 제목": {"state": "VALUE", "value": "원본 연결 검수"},
        "시작 찬양 1": {"state": "VALUE", "value": "380장"},
    }
    store.save(DAY, service, 0, {"fields": fields, "fridayMode": mode})
    return store


def score_file():
    deck = Presentation()
    deck.slide_width = 9144000
    deck.slide_height = 6858000
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, 3000000, 1000000)
    shape.text = "원드라이브로 가져온 악보"
    buff = BytesIO()
    deck.save(buff)
    return buff.getvalue()


def fake_cloud(store, monkeypatch, name="새찬송가 380장 악보.pptx", content=None,
               inside=True, file_id="ONE_SELECTED_ITEM"):
    cloud = OneDrive(store)
    monkeypatch.setattr(cloud, "selected", lambda purpose: {"id": "source-folder"})
    monkeypatch.setattr(cloud, "item", lambda ident:
                        {"id": ident, "name": name, "size": len(content or b"none"),
                         "file": {}, "parentReference": {"id": "source-folder"}})
    monkeypatch.setattr(cloud, "within", lambda item, folder: inside)
    monkeypatch.setattr(cloud, "_download", lambda item: content if content is not None else score_file())
    return cloud


def test_selected_wednesday_score_imported_and_inserted_without_changing_song_text(tmp_path, monkeypatch):
    store = make_store(tmp_path)
    cloud = fake_cloud(store, monkeypatch)
    first = cloud.import_for_service("ONE_SELECTED_ITEM", DAY, "wednesday", "시작 찬양 1")
    assert first["imported"] is True
    assert first["source"] == "onedrive"
    assert first["ready_for_ppt"] is True
    assert len(store.attachments(DAY, "wednesday")) == 1
    assert store.get(DAY, "wednesday")["data"]["fields"]["시작 찬양 1"]["value"] == "380장"
    work_revision = store.get(DAY, "wednesday")["revision"]
    result = generate(store, {"date": DAY, "service": "wednesday", "revision": work_revision,
                              "groups": [["찬양", ["시작 찬양 1"]], ["말씀", ["성경 본문", "설교 제목"]]]})
    report = json.loads((Path(store.path).with_suffix(".generated") /
                         (result["file_id"] + ".json")).read_text())
    mapping = next(item for item in report["mapping"] if item["key"] == "시작 찬양 1")
    assert mapping["status"] == "original_inserted"
    assert report["connections"]["시작 찬양 1"]["status"] == "connected"
    deck = Presentation(Path(store.path).with_suffix(".generated") / (result["file_id"] + ".pptx"))
    texts = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert "원드라이브로 가져온 악보" in texts


@pytest.mark.parametrize("slot,name,mode", [
    ("기도 담당자", "score.pptx", "in_person"),
    ("시작 찬양 1", "video.mp4", "in_person"),
    ("시작 찬양 1", "cover.pptx", "in_person"),
    ("시작 찬양 1", "score.pptx", "zoom"),
])
def test_wrong_slot_or_format_refused_before_download(tmp_path, monkeypatch, slot, name, mode):
    store = make_store(tmp_path, "friday" if mode == "zoom" else "wednesday", mode)
    cloud = fake_cloud(store, monkeypatch, name=name)
    def disallow(_): raise AssertionError("incompatible file must not download")
    monkeypatch.setattr(cloud, "_download", disallow)
    with pytest.raises(CloudError):
        cloud.import_for_service("ONE_SELECTED_ITEM", DAY,
                                 "friday" if mode == "zoom" else "wednesday", slot)
    assert not store.attachments(DAY, "friday" if mode == "zoom" else "wednesday")


def test_out_of_selected_folder_refused_before_download(tmp_path, monkeypatch):
    store = make_store(tmp_path)
    cloud = fake_cloud(store, monkeypatch, inside=False)
    monkeypatch.setattr(cloud, "_download",
                        lambda item: (_ for _ in ()).throw(AssertionError("not downloaded")))
    with pytest.raises(CloudError, match="폴더"):
        cloud.import_for_service("BAD_ITEM", DAY, "wednesday", "시작 찬양 1")
    assert not store.attachments(DAY, "wednesday")


def test_missing_song_field_is_not_auto_filled_from_filename(tmp_path, monkeypatch):
    store = make_store(tmp_path)
    record = store.get(DAY, "wednesday")
    record["data"]["fields"]["시작 찬양 1"] = {"state": "UNSET", "value": ""}
    store.save(DAY, "wednesday", record["revision"], record["data"])
    cloud = fake_cloud(store, monkeypatch)
    monkeypatch.setattr(cloud, "_download",
                        lambda item: (_ for _ in ()).throw(AssertionError("no download")))
    with pytest.raises(CloudError, match="곡명"):
        cloud.import_for_service("ONE_SELECTED_ITEM", DAY, "wednesday", "시작 찬양 1")
    assert store.get(DAY, "wednesday")["data"]["fields"]["시작 찬양 1"]["state"] == "UNSET"


def test_weekly_inline_picker_and_library_navigation():
    html = (Path(__file__).resolve().parents[1] /
            "src/media_automation/web/original/ui/index.html").read_text()
    js = (Path(__file__).resolve().parents[1] /
          "src/media_automation/web/original/ui/onedrive.js").read_text()
    assert 'data-page="library"' not in html
    assert '<option value="library">' not in html
    assert "function inlineSourcePicker(" in html
    assert "CloudMedia.searchWorshipSource(query,kind)" in html
    assert "CloudMedia.importWorshipSource(item.id,date,svc,label)" in html
    assert 'id="archive-library-manager"' in html
    assert "function searchWorshipSource(" in js
    assert "function importWorshipSource(" in js
    assert "CloudMedia.autoConnect()" not in html


def test_friday_zoom_one_selected_video_can_be_imported_without_other_sources(tmp_path, monkeypatch):
    store = make_store(tmp_path, "friday", mode="zoom")
    cloud = fake_cloud(store, monkeypatch, name="찬양영상.mp4", content=b"test-video")
    result = cloud.import_for_service("ONE_SELECTED_ITEM", DAY, "friday", "시작 찬양 1")
    assert result["source"] == "onedrive" and result["imported"]
    assets = store.attachments(DAY, "friday")
    assert len(assets) == 1
    assert assets[0]["slot"] == "시작 찬양 1"
    assert assets[0]["filename"] == "찬양영상.mp4"


def test_other_service_song_slot_is_not_modified(tmp_path, monkeypatch):
    store = make_store(tmp_path)
    store.save(DAY, "sunday", 0, {"fields": {
        "시작 찬양 1": {"state": "VALUE", "value": "다른 주일곡"}
    }})
    cloud = fake_cloud(store, monkeypatch)
    cloud.import_for_service("ONE_SELECTED_ITEM", DAY, "wednesday", "시작 찬양 1")
    assert not store.attachments(DAY, "sunday")
    assert store.get(DAY, "sunday")["data"]["fields"]["시작 찬양 1"]["value"] == "다른 주일곡"


def test_hosted_one_selected_one_drive_file_mutates_one_attachment_only(tmp_path):
    """Cloud mode must not fall back to the legacy song-only import API."""
    from types import SimpleNamespace
    from unittest.mock import Mock
    from media_automation.cloud.http import Hosted
    hosted = object.__new__(Hosted)
    data = score_file()
    digest = __import__("hashlib").sha256(data).hexdigest()
    cloud = Mock()
    cloud.item.return_value = {
        "id": "chosen-item", "name": "380장-악보.pptx",
        "size": len(data), "file": {},
        "parentReference": {"driveId": "drive-1", "id": "source"}
    }
    cloud.selected.return_value = {"id": "source"}
    cloud.within.return_value = True
    cloud._download.return_value = data
    hosted.cloud = cloud
    hosted.backend = Mock()
    hosted.catalog = lambda client, day, service: {
        "work": {"data": {"fridayMode": "in_person", "fields": {
            "시작 찬양 1": {"state": "VALUE", "value": "380장"},
        }}},
        "tables": {"files": []},
    }
    hosted.edit = Mock()
    result = hosted.import_for_slot(object(), {
        "id": "chosen-item", "date": DAY,
        "service": "wednesday", "slot": "시작 찬양 1"
    })
    assert result["digest"] == digest
    assert result["imported"] is True
    assert hosted.edit.call_args.args[2] == "attachment"
    assert hosted.edit.call_args.args[3]["slot"] == "시작 찬양 1"
    hosted.backend.call.assert_called_once()
    cloud._download.assert_called_once()
