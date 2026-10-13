import json
import subprocess
from pathlib import Path

from pptx import Presentation

from media_automation.ppt.base import create_4x3_presentation, add_centered_text_slide
from media_automation.web.original.generation import _asset, generate
from media_automation.web.original.work_store import WorkStore


def test_none_roundtrip_preserves_latest_manual_value_across_reload():
    html=Path('src/media_automation/web/original/ui/index.html').read_text()
    source=html[html.index('function setFieldAbsent('):html.index('async function selectOriginal(')]
    driver=source+'''
    let r={fields:{'기도 담당자':{state:'VALUE',value:'수정 담당자'}},sourceValues:{'기도 담당자':{state:'VALUE',value:'안내 담당자'}}};
    setFieldAbsent(r,'기도 담당자',true);
    if(r.fields['기도 담당자'].state!=='NONE'||r.fields['기도 담당자'].value)throw Error('NONE');
    r=JSON.parse(JSON.stringify(r));setFieldAbsent(r,'기도 담당자',false);
    if(r.fields['기도 담당자'].value!=='수정 담당자')throw Error('manual value lost');
    r.fields['기도 담당자']={state:'NONE',value:''};delete r.suspendedFields['기도 담당자'];
    setFieldAbsent(r,'기도 담당자',false);
    if(r.fields['기도 담당자'].value!=='안내 담당자')throw Error('source restoration lost');
    r.fields['기도 담당자']={state:'UNSET',value:''};setFieldAbsent(r,'기도 담당자',true);setFieldAbsent(r,'기도 담당자',false);
    if(r.fields['기도 담당자'].state!=='UNSET')throw Error('explicit empty changed');
    console.log('ok');
    '''
    assert subprocess.check_output(['node','-e',driver],text=True).strip()=='ok'


def test_original_cancel_targets_captured_week_even_after_navigation():
    html=Path('src/media_automation/web/original/ui/index.html').read_text()
    source=next(line for line in html.splitlines() if line.startswith('async function selectOriginal('))
    driver='''
    const old='wednesday:2026-10-14',current='sunday:2026-10-18';
    const records={[old]:{outputs:{ppt:{generated:true}}},[current]:{outputs:{}}};
    const savedAttachments=new Map([['wednesday|2026-10-14|시작 찬양 1',{digest:'old-original'}]]);
    let removed,saved;const removeOriginal=async target=>{removed=target;};
    const save=k=>saved=k,key=()=>current,msg=()=>{},assets=()=>{throw Error('wrong week rendered');},renderState=()=>{};
    '''+source+'''
    selectOriginal('wednesday|2026-10-14|시작 찬양 1','시작 찬양 1','NONE').then(()=>{
      if(removed.service!=='wednesday'||removed.date!=='2026-10-14'||removed.digest!=='old-original')throw Error('wrong attachment removed');
      if(saved!==old||records[old].assetChoices['시작 찬양 1']!=='NONE'||!records[old].outputs.ppt.changed)throw Error('cancel not saved');
      if(records[current].assetChoices||savedAttachments.size)throw Error('navigation affected cancel');
      console.log('ok');
    });
    '''
    assert subprocess.check_output(['node','-e',driver],text=True).strip()=='ok'


def setup_library(tmp_path):
    store=WorkStore(tmp_path/'work.db')
    song=store.register_song({'title':'검증 찬양'})
    assets=[]
    for name in ('first','second'):
        deck=create_4x3_presentation()
        for text in (name+' 원본 첫 장',name+' 원본 끝 장'):
            add_centered_text_slide(deck,text)
        path=tmp_path/(name+'.pptx');deck.save(path)
        assets.append(store.attach_song_asset(song['id'],path.name,path.read_bytes()))
    data={'fields':{'시작 찬양 1':{'state':'VALUE','value':'검증 찬양'},
                    '성경 본문':{'state':'VALUE','value':'요 3:16'},
                    '설교 제목':{'state':'VALUE','value':'검증 제목'}},
          'songChoices':{'시작 찬양 1':song['id']}}
    store.save('2026-10-14','wednesday',0,data)
    body=dict(date='2026-10-14',service='wednesday',revision=1,
              groups=[['찬양',['시작 찬양 1']],['말씀',['성경 본문','설교 제목']]])
    block=dict(key='시작 찬양 1',content='검증 찬양')
    return store,song,assets,data,body,block


def test_explicit_choice_inserts_only_chosen_library_original(tmp_path,monkeypatch):
    monkeypatch.setenv('MEDIA_WORSHIP_TEMPLATES',str(tmp_path/'templates'))
    store,song,assets,data,body,block=setup_library(tmp_path)
    assert _asset(store,body,block) is None
    data['assetChoices']={'시작 찬양 1':assets[1]['digest']}
    body['revision']=store.save(body['date'],body['service'],1,data)['revision']
    result=generate(store,body)
    directory=Path(store.path).with_suffix('.generated')
    report=json.loads((directory/(result['file_id']+'.json')).read_text())
    entry=next(x for x in report['mapping'] if x['key']=='시작 찬양 1')
    assert entry['status']=='original_inserted' and entry['end']-entry['start']==1
    deck=Presentation(directory/(result['file_id']+'.pptx'))
    text='\n'.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    assert 'second 원본 첫 장' in text and 'second 원본 끝 장' in text
    assert 'first 원본' not in text
    assert '시작 찬양 1:original_review' in result['unresolved']
    assert store.songs()[0]['assets'][1]['review']['permission']=='unknown'


def test_cancel_attachment_prevents_library_auto_reconnect_without_deleting_bytes(tmp_path):
    store,song,assets,data,body,block=setup_library(tmp_path)
    for asset in assets:
        store.review_asset(song['id'],asset['digest'],dict(technical=True,church=True,permission='allowed'))
    data['assetChoices']={'시작 찬양 1':assets[0]['digest']}
    store.save(body['date'],body['service'],1,data)
    latest=store.attach(body['date'],body['service'],block['key'],'second.pptx',
                       (tmp_path/'second.pptx').read_bytes(),song['id'])
    assert _asset(store,body,block)[2]['digest']==assets[1]['digest']
    store.remove_original(dict(date=body['date'],service=body['service'],slot=block['key'],digest=latest['digest']))
    data['assetChoices'][block['key']]='NONE'
    store.save(body['date'],body['service'],2,data)
    assert _asset(store,body,block) is None
    assert len(store.songs()[0]['assets'])==2
    assert store.original(latest['digest'])[0].is_file()


def test_cancelled_single_eligible_library_asset_is_not_inserted(tmp_path):
    store,song,assets,data,body,block=setup_library(tmp_path)
    store.remove_original(dict(song_id=song['id'],digest=assets[1]['digest']))
    store.review_asset(song['id'],assets[0]['digest'],dict(technical=True,church=True,permission='allowed'))
    assert _asset(store,body,block)[2]['digest']==assets[0]['digest']
    data['assetChoices']={block['key']:'NONE'}
    store.save(body['date'],body['service'],1,data)
    assert _asset(store,body,block) is None


def test_legacy_conversion_failure_reports_filename_without_aborting_generation(tmp_path,monkeypatch):
    monkeypatch.setenv('MEDIA_WORSHIP_TEMPLATES',str(tmp_path/'templates'))
    store,song,assets,data,body,block=setup_library(tmp_path)
    legacy=store.attach_song_asset(song['id'],'legacy-original.ppt',b'legacy test original')
    data['assetChoices']={block['key']:legacy['digest']}
    body['revision']=store.save(body['date'],body['service'],1,data)['revision']
    def fail(path):
        raise RuntimeError('converter unavailable')
    monkeypatch.setattr('media_automation.web.original.generation.convert_original',fail)
    result=generate(store,body)
    directory=Path(store.path).with_suffix('.generated')
    report=json.loads((directory/(result['file_id']+'.json')).read_text())
    assert report['connections'][block['key']]['status']=='conversion_failed'
    assert report['connections'][block['key']]['filename']=='legacy-original.ppt'
    assert block['key']+':score_asset' in result['unresolved']
    assert (directory/(result['file_id']+'.pptx')).is_file()
    assert store.original(legacy['digest'])[0].read_bytes()==b'legacy test original'
