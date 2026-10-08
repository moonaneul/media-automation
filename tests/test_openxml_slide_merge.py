from io import BytesIO
import hashlib
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.xmlchemy import OxmlElement
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from media_automation.ppt.slide_merge import OpenXmlSlideMerger


def picture_bytes(color):
    stream = BytesIO()
    Image.new('RGB', (40, 30), color).save(stream, format='PNG')
    return stream.getvalue()


def deck(path, labels, color):
    p = Presentation()
    for text in labels:
        s = p.slides.add_slide(p.slide_layouts[6])
        s.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = text
        s.shapes.add_picture(BytesIO(picture_bytes(color)), 0, Inches(1))
    p.save(path)
    return p


def test_merge_preserves_source_xml_images_notes_and_theme_with_name_collisions(tmp_path):
    target, source = tmp_path/'target.pptx', tmp_path/'source.pptx'
    deck(target, ['before', 'after'], 'red')
    p = deck(source, ['one', 'two', 'three'], 'blue')
    slide = p.slides[1]
    slide.notes_slide.notes_text_frame.text = 'original speaker notes'
    slide.shapes[0].text_frame.paragraphs[0].runs[0].hyperlink.address = 'https://example.org/score'
    transition = OxmlElement('p:transition')
    transition.append(OxmlElement('p:fade'))
    slide._element.append(transition)
    p.save(source)
    p = Presentation(source)
    source_xml = p.slides[1].part.blob
    source_theme = p.slides[1].slide_layout.slide_master.part.part_related_by(RT.THEME).blob
    original_sha = hashlib.sha256(source.read_bytes()).digest()

    merger = OpenXmlSlideMerger()
    merger.insert_range(target, source, after_slide=1, start_slide=2, end_slide=3)
    merged = Presentation(target)
    assert [s.shapes[0].text for s in merged.slides] == ['before', 'two', 'three', 'after']
    assert merged.slides[1].part.blob == source_xml
    assert merged.slides[1].notes_slide.notes_text_frame.text == 'original speaker notes'
    assert merged.slides[1].shapes[0].text_frame.paragraphs[0].runs[0].hyperlink.address == 'https://example.org/score'
    assert merged.slides[1].slide_layout.slide_master.part.part_related_by(RT.THEME).blob == source_theme
    assert merged.slides[0].shapes[1].image.blob == picture_bytes('red')
    assert merged.slides[1].shapes[1].image.blob == picture_bytes('blue')
    assert hashlib.sha256(source.read_bytes()).digest() == original_sha
    # Repeated imports must not overwrite parts imported by the previous call.
    merger.insert_range(target, source, after_slide=0, start_slide=1, end_slide=1)
    merged = Presentation(target)
    assert [s.shapes[0].text for s in merged.slides] == ['one', 'before', 'two', 'three', 'after']
    with zipfile.ZipFile(target) as z:
        assert len(z.namelist()) == len(set(z.namelist()))
    for part in merged.part.package.iter_parts():
        for rel in part.rels.values():
            if not rel.is_external:
                assert rel.target_part.package is merged.part.package


@pytest.mark.parametrize('start,end,after', [(0,1,0), (1,4,0), (2,1,0), (1,1,3)])
def test_invalid_range_does_not_change_destination(tmp_path, start, end, after):
    target, source = tmp_path/'target.pptx', tmp_path/'source.pptx'
    deck(target, ['target'], 'red')
    deck(source, ['source'], 'blue')
    before = target.read_bytes()
    with pytest.raises(ValueError):
        OpenXmlSlideMerger().insert_range(target, source, after_slide=after, start_slide=start, end_slide=end)
    assert target.read_bytes() == before


def test_different_canvas_rejected_without_scaling_original(tmp_path):
    target, source = tmp_path/'target.pptx', tmp_path/'source.pptx'
    deck(target, ['target'], 'red')
    p = deck(source, ['source'], 'blue')
    p.slide_width = Inches(13.333)
    p.save(source)
    with pytest.raises(ValueError, match='화면 크기'):
        OpenXmlSlideMerger().insert_all(target, source, after_slide=1)


def test_legacy_half_point_canvas_rounding_expands_margin_without_scaling(tmp_path):
    target, source = tmp_path/'target.pptx', tmp_path/'source.pptx'
    dest = deck(target, ['target'], 'red')
    dest.slide_width -= 6350
    dest.slide_height -= 6350
    dest.save(target)
    original = deck(source, ['source'], 'blue')
    blob = Presentation(source).slides[0].part.blob
    OpenXmlSlideMerger().insert_all(target, source, after_slide=1)
    merged = Presentation(target)
    assert (merged.slide_width, merged.slide_height) == (original.slide_width, original.slide_height)
    assert merged.slides[1].part.blob == blob
