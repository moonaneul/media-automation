import hashlib
import json
from pathlib import Path

import pytest
from pptx import Presentation

from media_automation.ppt.base import create_4x3_presentation, add_centered_text_slide
from media_automation.web.original.generation import generate
from media_automation.web.original.work_store import WorkStore


@pytest.mark.parametrize('invalid_bytes',[
    b'not a pptx package',
    b'version https://git-lfs.github.com/spec/v1\noid sha256:placeholder\nsize 999\n',
])
def test_invalid_score_is_reported_without_aborting_other_originals(tmp_path,monkeypatch,invalid_bytes):
    monkeypatch.setenv('MEDIA_WORSHIP_TEMPLATES',str(tmp_path/'templates'))
    monkeypatch.setenv('MEDIA_BIBLE_FILE',str(tmp_path/'no-bible.json'))
    monkeypatch.setenv('MEDIA_BIBLE_SOURCES',str(tmp_path/'no-sources'))
    store=WorkStore(tmp_path/'work.db')
    fields={label:{'state':'VALUE','value':text} for label,text in [
        ('성경 본문','요 3:16'),('설교 제목','연결 검증'),
        ('시작 찬양 1','확인 필요한 원본'),('시작 찬양 2','정상 원본')
    ]}
    revision=store.save('2026-10-07','wednesday',0,{'fields':fields})['revision']
    bad=store.attach('2026-10-07','wednesday','시작 찬양 1','unreadable.pptx',invalid_bytes)
    good=tmp_path/'good.pptx';deck=create_4x3_presentation()
    add_centered_text_slide(deck,'정상 악보 원본 검증');deck.save(good)
    store.attach('2026-10-07','wednesday','시작 찬양 2','good.pptx',good.read_bytes())
    result=generate(store,dict(date='2026-10-07',service='wednesday',revision=revision,
        groups=[['찬양',['시작 찬양 1','시작 찬양 2']],['말씀',['성경 본문','설교 제목']]]))
    directory=Path(store.path).with_suffix('.generated')
    report=json.loads((directory/(result['file_id']+'.json')).read_text())
    assert report['connections']['시작 찬양 1']['status']=='invalid_ppt_package'
    assert report['connections']['시작 찬양 1']['filename']=='unreadable.pptx'
    assert report['connections']['시작 찬양 2']['status']=='connected'
    assert '시작 찬양 1:score_asset' in result['unresolved']
    deck=Presentation(directory/(result['file_id']+'.pptx'))
    blank=next(x for x in report['mapping'] if x['key']=='시작 찬양 1')
    assert not deck.slides[blank['start']-1].shapes
    assert any(s.text=='정상 악보 원본 검증' for slide in deck.slides for s in slide.shapes if s.has_text_frame)
    preserved=Path(store.path).with_suffix('.files')/bad['digest']
    assert preserved.read_bytes()==invalid_bytes
    assert hashlib.sha256(preserved.read_bytes()).hexdigest()==bad['digest']
    assert result['complete'] is False
