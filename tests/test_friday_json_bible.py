from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path

import pytest
from pptx import Presentation
from media_automation.bible import BiblePassage, BibleVerse, InMemoryBibleProvider
from media_automation.bible.json_source import build_index, lookup, CombinedVerseRangeError
from media_automation.planning import BlockKind, WorshipBlock
from media_automation.ppt.friday_zoom import render_friday_zoom_block

ROOT = Path(__file__).resolve().parents[1]


def script(name):
    spec = spec_from_file_location(name, ROOT / 'scripts' / f'{name}.py')
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_combined_units_are_not_duplicated_or_silently_split():
    index = build_index({'신6:17': 'before', '신6:18-19': 'combined', '신6:20': 'after'})
    result = lookup(index, '신 6:17~20')
    assert [u['text'] for u in result['units']] == ['before', 'combined', 'after']
    assert result['units'][1]['reference'] == '신 6:18~19'
    with pytest.raises(CombinedVerseRangeError):
        lookup(index, '신 6:19')
    with pytest.raises(KeyError):
        lookup(index, '신 6:21')


def test_json_index_rejects_overlap_and_malformed_key():
    with pytest.raises(ValueError):
        build_index({'신6:18-19': 'combined', '신6:19': 'duplicate'})
    with pytest.raises(ValueError):
        build_index({'요18:이': 'broken'})


def test_resolver_splits_books_and_preserves_combined_numbers():
    resolver = script('resolve_friday_zoom_bible')
    index = build_index({'사43:21': 'isaiah', '왕상8:27': 'kings', '히10:19': 'nineteen', '히10:20': 'twenty', '신6:18-19': 'combined'})
    result, missing = resolver.resolve_requests({'requests': [
        {'reference': '사 43:21, 왕상 8:27, 히 10:19~20'},
        {'reference': '신 6:18~19'},
    ]}, {}, {}, index)
    assert missing == []
    assert list(result) == ['사 43:21', '왕상 8:27', '히 10:19~20', '신 6:18~19']
    assert result['신 6:18~19']['verses'] == {18: 'combined'}
    assert result['신 6:18~19']['verse_ends'] == {18: 19}
    _, missing = resolver.resolve_requests({'requests': [{'reference': '신 6:18'}]}, {}, {}, index)
    assert missing


def test_legacy_range_library_still_works():
    resolver = script('resolve_friday_zoom_bible')
    value = {'reference': '왕상 9:1~2', 'verses': {1: 'a', 2: 'b'}, '_source': 'test'}
    result, missing = resolver.resolve_requests({'requests': [{'reference': '왕상 9:1~2'}]}, {'passages': {'왕상 9:1-2': value}}, {})
    assert not missing
    assert '_source' not in result['왕상 9:1~2']
    assert value['_source'] == 'test'


NOTICE = '''[금요 기도회 순서] 10/9
<두 곡 찬양 후 기도 제목>
1) 곡 하나
2) 곡 둘
1. 첫 기도 제목
 - 빌 4:6~7
2. 둘째 기도 제목
 - 설명 원문
3) 곡 셋
<메시지 후 기도 제목>
고전 3:16 / 설교 제목
읽을 말씀 – 사 43:21, 왕상 8:27, 히 10:19~20
4) 곡 넷
1. 말씀 기도 하나
 - 설명 하나
2. 말씀 기도 둘
 - 설명 둘
<마지막 찬양 후 기도제목>
5) 곡 다섯
1. 공동체 기도 하나
 - 설명 하나
2. 공동체 기도 둘
 - 설명 둘
3. 공동체 기도 셋
 - 설명 셋
'''


def test_pastor_notice_and_reference_only_prayer():
    parser = script('parse_friday_zoom_notice')
    data, unknown = parser.parse_notice(NOTICE, reference_year=2026)
    assert unknown == []
    assert data['date']['value'] == '2026-10-09'
    assert data['personal_prayer']['value'] == ['개인 기도']
    assert [len(data[k]['value']) for k in ['first_prayer', 'word_prayer', 'community_prayer']] == [2, 2, 3]
    requests = parser.build_bible_requests(data)
    assert {'source': 'first_prayer', 'reference': '빌 4:6~7', 'translation': '개역개정', 'display_rule': '세례→침례'} in requests['requests']
    intake = {'fields': data}
    filler = script('fill_friday_zoom_prayer_scripture')
    count = filler.fill_prayer_scripture(intake, {'passages': {'빌 4:6~7': {'verses': {6: 'verse six', 7: 'verse seven'}}}})
    assert count == 1
    assert data['first_prayer']['value'][0] == '첫 기도 제목\n[6] verse six [7] verse seven (빌 4:6~7)'
    assert data['first_prayer']['value'][1] == '둘째 기도 제목\n설명 원문'
    assert filler.fill_prayer_scripture(intake, {'passages': {}}) == 0
    blank, _ = parser.parse_notice(NOTICE + '\n개인 기도:', reference_year=2026)
    assert blank['personal_prayer']['status'] == 'blank'


def test_friday_renderer_multiple_books_and_combined_verse_label():
    from types import SimpleNamespace
    provider = InMemoryBibleProvider({
        '사 43:21': BiblePassage('사 43:21', (BibleVerse(21, 'isaiah'),)),
        '신 6:18~19': BiblePassage('신 6:18~19', (BibleVerse(18, 'combined text', 19),)),
    })
    block = WorshipBlock(kind=BlockKind.SCRIPTURE, key='additional_scripture', value=SimpleNamespace(reference='사 43:21, 신 6:18~19'))
    prs = Presentation()
    slides = render_friday_zoom_block(prs, block, bible_provider=provider)
    assert len(slides) == 2
    text = '\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    assert '사 43:21' in text and '신 6:18~19' in text and '18~19' in text
    assert text.count('combined text') == 1


def test_new_weekly_fields_have_valid_status_without_inheriting_media():
    converter = script('intake_to_friday_zoom_weekly')
    assert converter.status_from_shape({'status': 'NONE'}) == 'VALUE'
    song = converter.song_value({'status': 'VALUE', 'title': 'old', 'media': {'status': 'VALUE'}}, 'new')
    assert song == {'status': 'VALUE', 'title': 'new', 'media': {'status': 'UNSET'}}
