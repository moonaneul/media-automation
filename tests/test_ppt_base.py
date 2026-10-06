from pptx import Presentation
from pptx.enum.text import (
    MSO_ANCHOR,
    PP_ALIGN,
)

from media_automation.ppt import (
    add_blank_slide,
    add_centered_text_slide,
    add_sermon_title_slide,
    create_4x3_presentation,
    create_16x9_presentation,
    get_blank_layout,
)


def test_create_4x3_presentation():
    prs = create_4x3_presentation()

    ratio = (
        prs.slide_width
        / prs.slide_height
    )

    assert round(ratio, 2) == 1.33


def test_create_16x9_presentation():
    prs = create_16x9_presentation()

    ratio = (
        prs.slide_width
        / prs.slide_height
    )

    assert round(ratio, 2) == 1.78


def test_get_blank_layout_returns_layout():
    prs = Presentation()

    layout = get_blank_layout(prs)

    assert layout is not None


def test_add_blank_slide_adds_one_slide():
    prs = create_4x3_presentation()

    add_blank_slide(prs)

    assert len(prs.slides) == 1


def test_centered_text_slide():
    prs = create_4x3_presentation()

    slide = add_centered_text_slide(
        prs,
        "기 도 : 홍길동",
    )

    assert len(prs.slides) == 1

    textbox = slide.shapes[0]
    frame = textbox.text_frame

    assert frame.vertical_anchor == MSO_ANCHOR.MIDDLE

    assert (
        frame.paragraphs[0].alignment
        == PP_ALIGN.CENTER
    )

    assert (
        frame.paragraphs[0].text
        == "기 도 : 홍길동"
    )

    assert all(
        run.font.bold is True
        for run in frame.paragraphs[0].runs
        if run.text.strip()
    )


def test_sermon_title_slide_contains_reference_and_title():
    prs = create_4x3_presentation()

    slide = add_sermon_title_slide(
        prs,
        "외모와 중심",
        "삼상 16:6-7",
    )

    paragraphs = (
        slide.shapes[0]
        .text_frame.paragraphs
    )

    text = "\n".join(
        paragraph.text
        for paragraph in paragraphs
    )

    assert "삼상 16:6-7" in text
    assert "외모와 중심" in text

    assert all(
        run.font.bold is True
        for paragraph in paragraphs
        for run in paragraph.runs
        if run.text.strip()
    )
