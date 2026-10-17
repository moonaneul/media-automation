"""Non-private smoke test for the read-only score/assembly comparison."""
from pptx import Presentation
from pptx.util import Inches
from scripts.diagnose_ppt_slide_mapping import compare


def test_read_only_slide_diagnostic_detects_same_and_changed(tmp_path, capsys):
    original = tmp_path / "song.pptx"
    built = tmp_path / "worship.pptx"

    source = Presentation()
    source.slide_width, source.slide_height = Inches(10), Inches(7.5)
    for text in ["악보 첫 화면", "악보 둘째 화면"]:
        slide = source.slides.add_slide(source.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = text
    source.save(original)

    made = Presentation(original)
    assert len(compare(original, original, 1)) == 2
    changed = made.slides[0].shapes[0]
    changed.text = "다른 내용"
    made.save(built)
    rows = compare(original, built, 1)
    assert rows[0]["same"]["text"] is False
    assert rows[1]["same"]["text"] is True
    assert "CORE-DIFF" in capsys.readouterr().out
