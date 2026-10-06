import pytest
from pptx.util import Pt

from media_automation.bible import (
    BiblePassage,
    BibleVerse,
)
from media_automation.ppt import (
    add_scripture_passage_slides,
    create_4x3_presentation,
)


def make_passage(
    *verses: BibleVerse,
) -> BiblePassage:
    return BiblePassage(
        reference="삼상 16:6~7",
        verses=verses,
    )


def get_body_paragraph(slide):
    return (
        slide.shapes[1]
        .text_frame
        .paragraphs[1]
    )


def test_scripture_creates_one_slide_per_verse():
    prs = create_4x3_presentation()

    passage = make_passage(
        BibleVerse(
            number=6,
            text="6절 본문",
        ),
        BibleVerse(
            number=7,
            text="7절 본문",
        ),
    )

    slides = add_scripture_passage_slides(
        prs,
        passage,
    )

    assert len(slides) == 2
    assert len(prs.slides) == 2


def test_scripture_slide_contains_reference():
    prs = create_4x3_presentation()

    passage = make_passage(
        BibleVerse(
            number=6,
            text="본문",
        ),
    )

    slide = add_scripture_passage_slides(
        prs,
        passage,
    )[0]

    assert (
        slide.shapes[0].text
        == "삼상 16:6~7"
    )


def test_scripture_keeps_number_separate_from_body():
    prs = create_4x3_presentation()

    passage = make_passage(
        BibleVerse(
            number=6,
            text="본문 내용",
        ),
    )

    slide = add_scripture_passage_slides(
        prs,
        passage,
    )[0]

    paragraphs = (
        slide.shapes[1]
        .text_frame
        .paragraphs
    )

    assert paragraphs[0].text == "6."
    assert paragraphs[1].text == "본문 내용"


def test_scripture_generated_text_is_bold():
    prs = create_4x3_presentation()

    passage = make_passage(
        BibleVerse(
            number=6,
            text="본문 내용",
        ),
    )

    slide = add_scripture_passage_slides(
        prs,
        passage,
    )[0]

    paragraphs = [
        *slide.shapes[0].text_frame.paragraphs,
        *slide.shapes[1].text_frame.paragraphs,
    ]

    assert all(
        run.font.bold is True
        for paragraph in paragraphs
        for run in paragraph.runs
        if run.text.strip()
    )


def test_short_verse_uses_32pt():
    prs = create_4x3_presentation()

    passage = make_passage(
        BibleVerse(
            number=6,
            text="짧은 본문",
        ),
    )

    slide = add_scripture_passage_slides(
        prs,
        passage,
    )[0]

    paragraph = get_body_paragraph(
        slide
    )

    assert (
        paragraph.runs[0].font.size
        == Pt(32)
    )


def test_medium_verse_uses_30pt():
    prs = create_4x3_presentation()

    passage = make_passage(
        BibleVerse(
            number=6,
            text="가" * 100,
        ),
    )

    slide = add_scripture_passage_slides(
        prs,
        passage,
    )[0]

    paragraph = get_body_paragraph(
        slide
    )

    assert (
        paragraph.runs[0].font.size
        == Pt(30)
    )


def test_long_verse_uses_28pt():
    prs = create_4x3_presentation()

    passage = make_passage(
        BibleVerse(
            number=6,
            text="가" * 150,
        ),
    )

    slide = add_scripture_passage_slides(
        prs,
        passage,
    )[0]

    paragraph = get_body_paragraph(
        slide
    )

    assert (
        paragraph.runs[0].font.size
        == Pt(28)
    )


def test_baptism_word_is_changed_for_display():
    prs = create_4x3_presentation()

    passage = make_passage(
        BibleVerse(
            number=6,
            text="그가 세례를 받으니",
        ),
    )

    slide = add_scripture_passage_slides(
        prs,
        passage,
    )[0]

    paragraph = get_body_paragraph(
        slide
    )

    assert (
        paragraph.text
        == "그가 침례를 받으니"
    )


def test_missing_verse_is_rejected():
    prs = create_4x3_presentation()

    passage = BiblePassage(
        reference="삼상 16:6~8",
        verses=(
            BibleVerse(
                number=6,
                text="6절",
            ),
            BibleVerse(
                number=8,
                text="8절",
            ),
        ),
    )

    with pytest.raises(
        ValueError,
        match="누락",
    ):
        add_scripture_passage_slides(
            prs,
            passage,
        )
