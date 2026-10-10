"""PPTX-to-PDF preview contract tests; no Office application needed."""
from pathlib import Path
import hashlib
import subprocess

import pytest
from pptx import Presentation
from reportlab.pdfgen import canvas

from media_automation.web import pptx_preview, job_previews
from media_automation.web.existing_inspections import ExistingInspections


def _pptx(path: Path, n: int = 2):
    prs = Presentation()
    for _ in range(n):
        prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(path)


def _mock_office(monkeypatch, count: int):
    monkeypatch.setattr(pptx_preview, "availability", lambda: ("powerpoint", "PowerPoint ready"))
    def fake_run(command, **kwargs):
        assert command[0] == "osascript"
        target = Path(command[-1])
        doc = canvas.Canvas(str(target))
        for _ in range(count):
            doc.drawString(20, 20, "rendered")
            doc.showPage()
        doc.save()
        return subprocess.CompletedProcess(command, 0, b"", b"")
    monkeypatch.setattr(pptx_preview.subprocess, "run", fake_run)


def test_powerpoint_preview_writes_matching_pdf(monkeypatch, tmp_path):
    _mock_office(monkeypatch, 2)
    source = tmp_path / "original.pptx"
    _pptx(source)
    original = hashlib.sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "render" / "slides.pdf"
    assert pptx_preview.convert_pptx(source, output) == "powerpoint"
    assert output.is_file()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original


def test_page_mismatch_rejects_preview(monkeypatch, tmp_path):
    _mock_office(monkeypatch, 1)
    source = tmp_path / "original.pptx"
    _pptx(source, n=2)
    output = tmp_path / "preview.pdf"
    with pytest.raises(ValueError, match="페이지 수"):
        pptx_preview.convert_pptx(source, output)
    assert not output.exists()


def test_existing_file_preview_can_use_powerpoint(monkeypatch, tmp_path):
    _mock_office(monkeypatch, 2)
    src = tmp_path / "source.pptx"
    _pptx(src)
    lib = ExistingInspections(tmp_path / "archive")
    meta = lib.create("sunday", "2026-09-27", "source.pptx", src.read_bytes())
    result = lib.build_preview(meta["inspection_id"])
    assert result["available"]
    assert result["method"] == "powerpoint"
    assert lib.preview_bytes(meta["inspection_id"]).startswith(b"%PDF")
    assert lib._dir(meta["inspection_id"]).joinpath("source.pptx").read_bytes() == src.read_bytes()


def test_generated_job_preview_uses_existing_artifact(monkeypatch, tmp_path):
    _mock_office(monkeypatch, 2)
    src = tmp_path / "slides.pptx"
    _pptx(src)
    folder = tmp_path / "job"
    preview = job_previews.build(folder, "output/sunday/slides.pptx", src)
    assert preview["available"]
    assert job_previews.read(folder, "output/sunday/slides.pptx", src).startswith(b"%PDF")


def test_pdf_previews_need_no_converter(monkeypatch, tmp_path):
    monkeypatch.setattr(pptx_preview, "availability", lambda: (None, "unavailable"))
    p = tmp_path / "bulletin.pdf"
    p.write_bytes(b"%PDF-irrelevant-for-direct-preview")
    assert job_previews.status(tmp_path, "output/test.pdf", p)["available"]
    assert job_previews.read(tmp_path, "output/test.pdf", p) == p.read_bytes()
