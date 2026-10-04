from pptx import Presentation

from media_automation.bible import (
    BiblePassage,
    BibleVerse,
    InMemoryBibleProvider,
)
from media_automation.planning import (
    BlockKind,
    SermonTitleContent,
    WorshipBlock,
)
from media_automation.ppt import (
    build_friday_zoom_preview,
)
from media_automation.weekly_data.models import (
    ScriptureField,
    TopicsField,
    WeeklyStatus,
    ZoomSong,
    MediaField,
    FileStatus,
    TechnicalCheck,
    ChurchReview,
    UsagePermission,
)


def make_zoom_song(
    title: str,
) -> ZoomSong:
    return ZoomSong(
        status=WeeklyStatus.VALUE,
        title=title,
        media=MediaField(
            status=WeeklyStatus.VALUE,
            file_status=FileStatus.AVAILABLE,
            technical_check=(
                TechnicalCheck.UNVERIFIED
            ),
            church_review=(
                ChurchReview.UNVERIFIED
            ),
            usage_permission=(
                UsagePermission.UNKNOWN
            ),
        ),
    )


def test_zoom_preview_skips_song_without_placeholder(
    tmp_path,
):
    output = (
        tmp_path / "zoom-preview.pptx"
    )

    plan = [
        WorshipBlock(
            kind=BlockKind.PRE_SERVICE,
            key="pre_service",
        ),
        WorshipBlock(
            kind=BlockKind.ZOOM_SONG,
            key="opening_song",
            value=make_zoom_song(
                "예시 찬양"
            ),
        ),
        WorshipBlock(
            kind=BlockKind.BLANK,
            key="transition",
        ),
        WorshipBlock(
            kind=BlockKind.PRAYER_TOPICS,
            key="first_prayer",
            value=TopicsField(
                status=WeeklyStatus.VALUE,
                topics=[
                    "첫 번째 기도 제목",
                ],
            ),
        ),
    ]

    provider = InMemoryBibleProvider(
        {}
    )

    build_friday_zoom_preview(
        plan,
        output,
        bible_provider=provider,
    )

    prs = Presentation(output)

    # 예배 준비 + 기도 제목
    # 영상 대체 화면은 생성하지 않음
    assert len(prs.slides) == 2

    texts = []

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                texts.append(
                    shape.text
                )

    joined = "\n".join(texts)

    assert "예배 준비" in joined
    assert "첫 번째 기도 제목" in joined
    assert "예시 찬양" not in joined


def test_zoom_preview_renders_scripture_and_title(
    tmp_path,
):
    output = (
        tmp_path / "zoom-scripture.pptx"
    )

    plan = [
        WorshipBlock(
            kind=BlockKind.SCRIPTURE,
            key="scripture",
            value=ScriptureField(
                status=WeeklyStatus.VALUE,
                reference="요 3:16-18",
            ),
        ),
        WorshipBlock(
            kind=BlockKind.SERMON_TITLE,
            key="sermon_title",
            value=SermonTitleContent(
                title="예시 설교 제목",
                scripture_reference=(
                    "요 3:16~18"
                ),
            ),
        ),
    ]

    provider = InMemoryBibleProvider(
        {
            "요 3:16-18": BiblePassage(
                reference="요 3:16~18",
                verses=(
                    BibleVerse(
                        number=16,
                        text="샘플 16절",
                    ),
                    BibleVerse(
                        number=17,
                        text="샘플 17절",
                    ),
                    BibleVerse(
                        number=18,
                        text="샘플 18절",
                    ),
                ),
            ),
        }
    )

    build_friday_zoom_preview(
        plan,
        output,
        bible_provider=provider,
    )

    prs = Presentation(output)

    # 본문 범위 안내 1장
    # + 본문 2장
    # + 설교 제목 1장
    assert len(prs.slides) == 4

    all_text = "\n".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )

    assert "16. 샘플 16절" in all_text
    assert "예시 설교 제목" in all_text