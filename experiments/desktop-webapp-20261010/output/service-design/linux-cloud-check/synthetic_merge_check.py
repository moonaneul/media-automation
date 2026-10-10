"""Exercise the existing local merger with LibreOffice-created neutral fixtures."""
from pathlib import Path
import hashlib
import subprocess
import zipfile
from lxml import etree as E
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from slide_merge_candidate import OpenXmlSlideMerger, create_platform_slide_merger
from PIL import Image, ImageDraw

ROOT=Path('/tmp/merge-check')
ROOT.mkdir()
def run(*args): subprocess.run(args,check=True,timeout=90)
score=Image.new('RGB',(800,200),'white')
d=ImageDraw.Draw(score)
for y in (50,70,90,110,130): d.line((30,y,770,y),fill='black',width=3)
for x,y in ((160,110),(330,90),(500,70)):
    d.ellipse((x-13,y-9,x+13,y+9),fill='black')
    d.line((x+12,y,x+12,y-55),fill='black',width=3)
score.save(ROOT/'score.png')
import base64
encoded=base64.b64encode((ROOT/'score.png').read_bytes()).decode()

def deck(name, labels, with_score=False):
    pages=[]
    for label in labels:
        picture=f'<draw:frame svg:x="2cm" svg:y="5cm" svg:width="20cm" svg:height="5cm"><draw:image draw:mime-type="image/png"><office:binary-data>{encoded}</office:binary-data></draw:image></draw:frame>' if with_score else ''
        pages.append(f'<draw:page draw:name="{label}"><draw:frame svg:x="2cm" svg:y="1cm" svg:width="24cm" svg:height="3cm"><draw:text-box><text:p>{label}</text:p></draw:text-box></draw:frame>{picture}</draw:page>')
    xml=f'''<?xml version="1.0" encoding="UTF-8"?><office:document xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" office:version="1.2" office:mimetype="application/vnd.oasis.opendocument.presentation"><office:body><office:presentation>{''.join(pages)}</office:presentation></office:body></office:document>'''
    source=ROOT/f'{name}.fodp'
    source.write_text(xml)
    run('libreoffice',f'-env:UserInstallation=file:///tmp/lo-{name}','--headless','--convert-to','pptx','--outdir',str(ROOT),str(source))
    path=ROOT/f'{name}.pptx'
    assert path.is_file()
    return path

target=deck('target',['BEFORE','AFTER'])
source=deck('source',['SCORE_ONE','SCORE_TWO','SCORE_THREE'],True)
original_sha=hashlib.sha256(source.read_bytes()).hexdigest()
original=Presentation(source)
original_xml=[s.part.blob for s in original.slides]
original_theme=original.slides[0].slide_layout.slide_master.part.part_related_by(RT.THEME).blob
original_images=[s.shapes[1].image.blob for s in original.slides]
assert isinstance(create_platform_slide_merger(),OpenXmlSlideMerger)
merger=create_platform_slide_merger()
merger.insert_range(target,source,after_slide=1,start_slide=2,end_slide=3)
merged=Presentation(target)
assert [s.shapes[0].text for s in merged.slides]==['BEFORE','SCORE_TWO','SCORE_THREE','AFTER']
assert [s.part.blob for s in list(merged.slides)[1:3]]==original_xml[1:3]
assert merged.slides[1].slide_layout.slide_master.part.part_related_by(RT.THEME).blob==original_theme
assert [s.shapes[1].image.blob for s in list(merged.slides)[1:3]]==original_images[1:3]
merger.insert_range(target,source,after_slide=0,start_slide=1,end_slide=1)
merged=Presentation(target)
assert [s.shapes[0].text for s in merged.slides]==['SCORE_ONE','BEFORE','SCORE_TWO','SCORE_THREE','AFTER']
assert hashlib.sha256(source.read_bytes()).hexdigest()==original_sha
with zipfile.ZipFile(target) as z:
    assert z.testzip() is None
    assert len(z.namelist())==len(set(z.namelist()))
for part in merged.part.package.iter_parts():
    for rel in part.rels.values():
        if not rel.is_external: assert rel.target_part.package is merged.part.package
before=target.read_bytes()
for start,end,after in ((0,1,0),(1,4,0),(2,1,0),(1,1,6)):
    try: merger.insert_range(target,source,after_slide=after,start_slide=start,end_slide=end)
    except ValueError: pass
    else: raise AssertionError('Invalid insertion accepted')
    assert target.read_bytes()==before
print('PASS: Linux merger; order, source XML/theme/images, repeated imports, source unchanged, invalid range rejection')
run('libreoffice','-env:UserInstallation=file:///tmp/lo-merged','--headless','--convert-to','pdf','--outdir',str(ROOT),str(target))
run('pdftotext','-layout',str(ROOT/'target.pdf'),str(ROOT/'target.txt'))
pages=(ROOT/'target.txt').read_text().split('\f')
expected=['SCORE_ONE','BEFORE','SCORE_TWO','SCORE_THREE','AFTER']
assert all(label in pages[i] for i,label in enumerate(expected))
run('pdftoppm','-scale-to','1000','-png',str(ROOT/'target.pdf'),str(ROOT/'page'))
for n in (1,3,4):
    im=Image.open(ROOT/f'page-{n}.png').convert('RGB')
    w,h=im.size
    region=im.crop((int(w*.05),int(h*.3),int(w*.9),int(h*.8)))
    dark=sum(1 for p in region.getdata() if max(p)<100)
    assert dark>500,(n,dark)
    print(f'PAGE {n}: staff pixels={dark}')
print('PASS: five merged slides rendered; PDF page order and visible staff pixels')
print('NOT CHECKED: church template fidelity, PowerPoint repair warnings, full worship assembly')
