from __future__ import annotations

import platform
import subprocess
import os
import tempfile
import re
from pathlib import Path
from typing import Protocol

from pptx import Presentation
from pptx.opc.package import _Relationship
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT, RELATIONSHIP_TARGET_MODE as RTM
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn


class OpenXmlSlideMerger:
    """Copy original PPTX parts and relationships without a PowerPoint process.

    Slide XML, pictures, layouts, masters, themes and other dependent parts are
    retained; this does not reconstruct shapes or flatten the source slides.
    """

    def insert_all(self, destination, source, *, after_slide):
        count = len(Presentation(source).slides)
        if count:
            self.insert_range(destination, source, after_slide=after_slide,
                              start_slide=1, end_slide=count)

    def insert_range(self, destination, source, *, after_slide, start_slide, end_slide):
        destination, source = Path(destination).resolve(), Path(source).resolve()
        if destination == source:
            raise ValueError("원본과 대상 PPT는 서로 다른 파일이어야 합니다.")
        target, original = Presentation(destination), Presentation(source)
        if not 1 <= start_slide <= end_slide <= len(original.slides):
            raise ValueError("삽입할 슬라이드 범위가 원본 장수를 벗어났습니다.")
        if not 0 <= after_slide <= len(target.slides):
            raise ValueError("삽입 위치가 대상 장수를 벗어났습니다.")
        if (target.slide_width, target.slide_height) != (original.slide_width, original.slide_height):
            # Legacy 4:3 templates can differ by a half point due to rounding.
            # Expand that small margin; never scale or crop the original shapes.
            if (abs(target.slide_width - original.slide_width) > 12700 or
                    abs(target.slide_height - original.slide_height) > 12700):
                raise ValueError("원본 화면 보존을 위해 대상과 악보 PPT의 화면 크기가 같아야 합니다.")
            target.slide_width = max(target.slide_width, original.slide_width)
            target.slide_height = max(target.slide_height, original.slide_height)

        package = target.part.package
        names = {str(part.partname) for part in package.iter_parts()}
        copied = {}

        def clone(part):
            if part in copied:
                return copied[part]
            name = str(part.partname)
            stem, extension = name.rsplit('.', 1)
            numbered = re.fullmatch(r'(.*?)(\d+)', stem)
            index = int(numbered.group(2)) if numbered else 1
            while name in names:
                index += 1
                name = (f"{numbered.group(1)}{index}.{extension}" if numbered
                        else f"{stem}_import{index}.{extension}")
            names.add(name)
            result = type(part).load(PackURI(name), part.content_type, package, part.blob)
            copied[part] = result  # layouts/masters and notes/slides contain cycles
            for rel in part.rels.values():
                dependency = rel.target_ref if rel.is_external else clone(rel.target_part)
                # Preserve rIds used inside the unchanged source XML.
                result.rels._rels[rel.rId] = _Relationship(
                    result.partname.baseURI, rel.rId, rel.reltype,
                    RTM.EXTERNAL if rel.is_external else RTM.INTERNAL, dependency,
                )
            return result

        ids = target.slides._sldIdLst
        for offset, slide in enumerate(list(original.slides)[start_slide-1:end_slide]):
            part = clone(slide.part)
            rid = target.part.relate_to(part, RT.SLIDE)
            element = ids.add_sldId(rid)
            ids.remove(element)
            ids.insert(after_slide + offset, element)

        masters = target.part._element.find(qn('p:sldMasterIdLst'))
        if masters is None:
            masters = OxmlElement('p:sldMasterIdLst')
            target.part._element.insert(0, masters)
        master_id = max([2147483647] + [int(e.get('id')) for e in masters])
        for part in copied.values():
            if part.content_type != 'application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml':
                continue
            master_id += 1
            element = OxmlElement('p:sldMasterId')
            element.set('id', str(master_id))
            element.set(qn('r:id'), target.part.relate_to(part, RT.SLIDE_MASTER))
            masters.append(element)

        # A failed save must not leave a partially written destination.
        with tempfile.NamedTemporaryFile(dir=destination.parent, suffix='.pptx', delete=False) as f:
            temporary = Path(f.name)
        try:
            target.save(temporary)
            os.replace(temporary, destination)
        finally:
            temporary.unlink(missing_ok=True)


class SlideMerger(Protocol):
    def insert_all(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
    ) -> None:
        ...

    def insert_range(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
        start_slide: int,
        end_slide: int,
    ) -> None:
        ...


class PowerPointComSlideMerger:
    """
    Windows + Microsoft PowerPoint용 슬라이드 병합기.
    """

    def _require_windows(self) -> None:
        if platform.system() != "Windows":
            raise RuntimeError(
                "PowerPoint COM 슬라이드 병합은 "
                "Windows에서만 지원합니다."
            )

    def insert_all(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
    ) -> None:
        source = Path(source)

        source_prs = Presentation(source)
        slide_count = len(source_prs.slides)

        if slide_count == 0:
            return

        self.insert_range(
            destination,
            source,
            after_slide=after_slide,
            start_slide=1,
            end_slide=slide_count,
        )

    def insert_range(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
        start_slide: int,
        end_slide: int,
    ) -> None:
        self._require_windows()

        destination = Path(destination).resolve()
        source = Path(source).resolve()

        if start_slide < 1:
            raise ValueError(
                "start_slide은 1 이상이어야 합니다."
            )

        if end_slide < start_slide:
            raise ValueError(
                "end_slide은 start_slide보다 "
                "작을 수 없습니다."
            )

        if not destination.exists():
            raise FileNotFoundError(
                f"대상 PPT가 없습니다: {destination}"
            )

        if not source.exists():
            raise FileNotFoundError(
                f"삽입할 PPT가 없습니다: {source}"
            )

        try:
            import win32com.client
        except ImportError as error:
            raise RuntimeError(
                "Windows PowerPoint 병합에는 "
                "pywin32가 필요합니다."
            ) from error

        powerpoint = None
        presentation = None

        try:
            powerpoint = (
                win32com.client.DispatchEx(
                    "PowerPoint.Application"
                )
            )

            presentation = (
                powerpoint.Presentations.Open(
                    str(destination),
                    False,
                    False,
                    False,
                )
            )

            presentation.Slides.InsertFromFile(
                str(source),
                after_slide,
                start_slide,
                end_slide,
            )

            presentation.Save()

        finally:
            if presentation is not None:
                presentation.Close()

            if powerpoint is not None:
                powerpoint.Quit()


class MacPowerPointAppleScriptSlideMerger:
    """
    macOS + Microsoft PowerPoint용 슬라이드 병합기.

    PowerPoint 자체의 copy/paste 기능을 AppleScript로 호출해서
    python-pptx가 원본 악보 슬라이드를 재구성하지 않도록 한다.
    """

    def _require_macos(self) -> None:
        if platform.system() != "Darwin":
            raise RuntimeError(
                "Mac PowerPoint 슬라이드 병합은 "
                "macOS에서만 지원합니다."
            )

    def insert_all(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
    ) -> None:
        source = Path(source)

        source_prs = Presentation(source)
        slide_count = len(source_prs.slides)

        if slide_count == 0:
            return

        self.insert_range(
            destination,
            source,
            after_slide=after_slide,
            start_slide=1,
            end_slide=slide_count,
        )

    def insert_range(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
        start_slide: int,
        end_slide: int,
    ) -> None:
        self._require_macos()

        destination = Path(destination).resolve()
        source = Path(source).resolve()

        if start_slide < 1:
            raise ValueError(
                "start_slide은 1 이상이어야 합니다."
            )

        if end_slide < start_slide:
            raise ValueError(
                "end_slide은 start_slide보다 "
                "작을 수 없습니다."
            )

        # 현재 실제 예배 PPT는 PRE_SERVICE 뒤에 찬양이 들어가므로
        # after_slide >= 1이다.
        # 0번 위치 삽입은 별도 구현 전까지 명시적으로 막는다.
        if after_slide < 1:
            raise ValueError(
                "Mac PowerPoint 병합은 현재 "
                "첫 슬라이드 앞 삽입을 지원하지 않습니다."
            )

        if not destination.exists():
            raise FileNotFoundError(
                f"대상 PPT가 없습니다: {destination}"
            )

        if not source.exists():
            raise FileNotFoundError(
                f"삽입할 PPT가 없습니다: {source}"
            )

        script = r'''
on run argv
    set destinationPath to item 1 of argv
    set sourcePath to item 2 of argv
    set afterSlideNumber to (item 3 of argv) as integer
    set startSlideNumber to (item 4 of argv) as integer
    set endSlideNumber to (item 5 of argv) as integer

    set destinationFile to POSIX file destinationPath as alias
    set sourceFile to POSIX file sourcePath as alias

    tell application "Microsoft PowerPoint"
        activate

        open sourceFile
        set sourcePresentation to active presentation

        open destinationFile
        set destinationPresentation to active presentation

        set insertionPoint to afterSlideNumber

        repeat with sourceSlideNumber from startSlideNumber to endSlideNumber
            copy object slide sourceSlideNumber of sourcePresentation

            select slide insertionPoint of destinationPresentation

            tell active window
                set view type to slide sorter view
                paste object its view
                set view type to normal view
            end tell

            set insertionPoint to insertionPoint + 1
        end repeat

        close sourcePresentation saving no
        close destinationPresentation saving yes
    end tell
end run
'''

        result = subprocess.run(
            [
                "osascript",
                "-e",
                script,
                str(destination),
                str(source),
                str(after_slide),
                str(start_slide),
                str(end_slide),
            ],
            text=True,
            capture_output=True,
        )

        if result.returncode != 0:
            message = (
                result.stderr.strip()
                or result.stdout.strip()
                or "알 수 없는 AppleScript 오류"
            )

            raise RuntimeError(
                "Mac PowerPoint 슬라이드 병합에 "
                f"실패했습니다: {message}"
            )


def create_platform_slide_merger() -> SlideMerger:
    system = platform.system()

    if system == "Windows":
        return PowerPointComSlideMerger()

    if system == "Darwin":
        return OpenXmlSlideMerger()

    if system == "Linux":
        return OpenXmlSlideMerger()

    raise RuntimeError(
        f"지원하지 않는 운영체제입니다: {system}"
    )
