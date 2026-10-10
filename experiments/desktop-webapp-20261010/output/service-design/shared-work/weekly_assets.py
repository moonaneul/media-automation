"""Use explicit attachments for this week's slot before reviewed library matches."""
import hashlib
from pathlib import Path
from pptx import Presentation

def resolve_attachment(attachment, directory):
    digest=attachment['digest']; file=Path(directory)/digest
    if not file.is_file() or hashlib.sha256(file.read_bytes()).hexdigest()!=digest:
        return {'status':'file_missing_or_changed'}
    if file.read_bytes()[:8]==bytes.fromhex('d0cf11e0a1b11ae1'):
        from ppt_conversion import convert_original
        try:file=convert_original(file)
        except Exception:return {'status':'conversion_failed','filename':attachment['filename']}
        digest=hashlib.sha256(file.read_bytes()).hexdigest()
    if Path(attachment['filename']).suffix.lower() not in ('.ppt','.pptx'):return {'status':'unsupported_score_format'}
    try:
        deck=Presentation(file)
        if not deck.slides:return {'status':'empty_original'}
        if abs(deck.slide_width-9144000)>12700 or abs(deck.slide_height-6858000)>12700:return {'status':'canvas_mismatch'}
    except Exception:return {'status':'invalid_pptx'}
    return {'status':'connected','source':'weekly_attachment','permission_review':'not_confirmed','path':str(file.resolve()),'digest':digest,'slide_count':len(deck.slides)}

def prepend_start(file, report, service, engine):
    sources={name:Path(__file__).with_name('static')/(name+'-start.pptx') for name in ('wednesday','sunday')}
    source=sources.get(service)
    if not source:return report
    if not source.is_file():raise ValueError('기존 시작 안내 원본을 찾지 못했습니다.')
    engine.OpenXmlSlideMerger().insert_range(file,source,after_slide=0,start_slide=1,end_slide=4)
    for entry in report['mapping']:entry['start']+=4;entry['end']+=4
    report['mapping'].insert(0,{'key':'예배 전 안내','kind':'pre_service','start':1,'end':4,'status':'original_inserted'})
    report['slide_count']+=4
    return report
