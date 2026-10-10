import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from attach_originals import attach_originals

ROOT=Path(__file__).parent

class AttachOriginalTests(unittest.TestCase):
    def setUp(self):
        self.preview=ROOT/'output/friday-order-with-bible-preview.pptx'
        self.report=json.loads((ROOT/'bible-assembly-report.json').read_text())
        self.source=ROOT/'output/friday-prayer-topics.pptx'
        self.connection={'status':'connected','path':str(self.source),'digest':hashlib.sha256(self.source.read_bytes()).hexdigest(),'slide_count':3}

    def test_variable_source_count_updates_later_blocks_without_changing_originals(self):
        before=self.preview.read_bytes(),self.source.read_bytes()
        with tempfile.TemporaryDirectory() as d:
            out=Path(d)/'assembled.pptx'
            result=attach_originals(self.preview,self.report,{'song_1':self.connection,'song_6':self.connection},out)
            self.assertEqual(result['slide_count'],29)
            mapping={e['key']:e for e in result['mapping']}
            self.assertEqual((mapping['song_1']['start'],mapping['song_1']['end']),(1,3))
            self.assertEqual(mapping['word_prayer']['start'],23)
            self.assertFalse(result['complete'])
            deck=Presentation(out)
            original=Presentation(self.source)
            self.assertEqual([s.part.blob for s in list(deck.slides)[:3]],[s.part.blob for s in original.slides])
            self.assertEqual((self.preview.read_bytes(),self.source.read_bytes()),before)

    def test_changed_original_fails_without_creating_output(self):
        with tempfile.TemporaryDirectory() as d:
            out=Path(d)/'failed.pptx'
            bad={**self.connection,'digest':'0'*64}
            with self.assertRaises(ValueError):attach_originals(self.preview,self.report,{'song_6':bad},out)
            self.assertFalse(out.exists())

    def test_unconnected_slots_stay_incomplete(self):
        with tempfile.TemporaryDirectory() as d:
            result=attach_originals(self.preview,self.report,{},Path(d)/'unchanged.pptx')
            self.assertEqual(result['mapping'],self.report['mapping'])
            self.assertFalse(result['complete'])

if __name__=='__main__':unittest.main()
