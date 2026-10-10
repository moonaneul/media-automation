"""Replace reserved song slots using the previously tested OpenXML merger."""
import copy
import hashlib
import importlib.util
import os
from pathlib import Path
import shutil
import tempfile
from pptx import Presentation

module=Path(__file__).parent.parent/'linux-cloud-check/slide_merge_candidate.py'
spec=importlib.util.spec_from_file_location('checked_slide_merger',module)
engine=importlib.util.module_from_spec(spec)
spec.loader.exec_module(engine)


def attach_originals(preview,report,connections,destination):
    preview,destination=Path(preview).resolve(),Path(destination).resolve()
    if destination==preview or destination.exists():
        raise ValueError('Use a new output file; preserve the preview')
    if len(Presentation(preview).slides)!=report['slide_count']:
        raise ValueError('Preview and assembly map do not match')
    output=copy.deepcopy(report)
    destination.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(dir=destination.parent) as directory:
        working=Path(directory)/'assembly.pptx'
        shutil.copyfile(preview,working)
        offset=0
        for entry in output['mapping']:
            start,end=entry['start']+offset,entry['end']+offset
            connection=connections.get(entry['key'],{})
            if entry['kind']=='song' and connection.get('status')=='connected':
                if end!=start or entry['status']!='missing_song_asset':
                    raise ValueError('Expected one unresolved song reservation')
                source=Path(connection['path']).resolve()
                if hashlib.sha256(source.read_bytes()).hexdigest()!=connection['digest']:
                    raise ValueError('Original changed since connection check')
                count=len(Presentation(source).slides)
                if count!=connection['slide_count'] or count==0:
                    raise ValueError('Original slide count changed')
                engine.OpenXmlSlideMerger().insert_all(working,source,after_slide=start-1)
                deck=Presentation(working)
                # Original inserted before the reserved blank; remove that blank only.
                blank=deck.slides[start-1+count]
                if any(s.has_text_frame and s.text.strip() for s in blank.shapes):
                    raise ValueError('Reservation contains text; refuse replacement')
                node=deck.slides._sldIdLst[start-1+count]
                deck.part.drop_rel(node.rId)
                deck.slides._sldIdLst.remove(node)
                deck.save(working)
                end=start+count-1
                offset+=count-1
                entry['status']='original_inserted'
                entry['digest']=connection['digest']
                output['unresolved']=[u for u in output['unresolved'] if u!=entry['key']+':score_asset']
            entry['start'],entry['end']=start,end
        output['slide_count']+=offset
        output['complete']=not output['unresolved']
        if len(Presentation(working).slides)!=output['slide_count']:
            raise ValueError('Assembled count does not match updated map')
        # A partial failure leaves the requested destination absent.
        os.replace(working,destination)
    return output
