"""Resolve only registered, reviewed originals; never inspect past worship decks."""
import hashlib
import zipfile
from pathlib import Path


def normalize(title):
    return ''.join(title.split()).casefold()


def resolve_song(title, songs, file_directory):
    if not isinstance(title,str) or not title.strip():
        return {'status':'missing_song_input'}
    matches=[s for s in songs if normalize(title) in [normalize(t) for t in [s['title'],*s.get('aliases',[])]]]
    if not matches:
        return {'status':'song_not_registered','title':title}
    if len(matches)!=1:
        return {'status':'confirm_song','song_ids':[s['id'] for s in matches]}
    song=matches[0]
    eligible=[]
    rejected=[]
    for asset in song.get('assets',[]):
        review=asset.get('review',{})
        if not (review.get('technical') is True and review.get('church') is True and review.get('permission')=='allowed'):
            rejected.append('review_pending')
            continue
        if Path(asset['filename']).suffix.lower()!='.pptx':
            rejected.append('conversion_required_or_wrong_format')
            continue
        digest=asset.get('digest','')
        if len(digest)!=64 or any(c not in '0123456789abcdef' for c in digest):
            rejected.append('invalid_file_id')
            continue
        file=Path(file_directory)/digest
        if not file.is_file() or hashlib.sha256(file.read_bytes()).hexdigest()!=digest:
            rejected.append('file_missing_or_changed')
            continue
        try:
            with zipfile.ZipFile(file) as z:
                if z.testzip() is not None: raise ValueError('Invalid CRC')
            from pptx import Presentation
            deck=Presentation(file)
            if not len(deck.slides): raise ValueError('Empty deck')
            if abs(deck.slide_width-9144000)>12700 or abs(deck.slide_height-6858000)>12700:
                rejected.append('canvas_mismatch')
                continue
        except Exception:
            rejected.append('invalid_pptx')
            continue
        eligible.append({'digest':digest,'path':str(file.resolve()),'slide_count':len(deck.slides)})
    if not eligible:
        return {'status':'no_eligible_original','song_id':song['id'],'reasons':rejected}
    if len(eligible)>1:
        return {'status':'confirm_original','song_id':song['id'],'candidates':eligible}
    return {'status':'connected','song_id':song['id'],**eligible[0]}


def connect_plan(plan,songs,file_directory):
    return {b['key']:resolve_song(b['content'],songs,file_directory) for b in plan['blocks'] if b['kind']=='song'}
