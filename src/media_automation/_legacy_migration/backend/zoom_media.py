"""Embed explicit weekly MP4s, preserving their exact bytes and playback relationships."""
import hashlib
import io
import tempfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

def attach_videos(file,report,attachments,directory):
    deck=Presentation(file);connections={}
    with tempfile.TemporaryDirectory() as folder:
        poster=Path(folder)/'poster.png';Image.new('RGB',(1280,720),'black').save(poster)
        for entry in report['mapping']:
            if entry['kind']!='song':continue
            attachment=attachments.get(entry['key'])
            if not attachment:
                connections[entry['key']]={'status':'media_not_attached'};continue
            source=Path(directory)/attachment['digest']
            if not source.is_file() or hashlib.sha256(source.read_bytes()).hexdigest()!=attachment['digest']:
                connections[entry['key']]={'status':'file_missing_or_changed'};continue
            if Path(attachment['filename']).suffix.lower()!='.mp4':
                connections[entry['key']]={'status':'unsupported_video_format'};continue
            slide=deck.slides[entry['start']-1]
            movie=slide.shapes.add_movie(io.BytesIO(source.read_bytes()),0,0,deck.slide_width,deck.slide_height,poster_frame_image=str(poster),mime_type='video/mp4')
            movie.name=attachment['filename']
            for condition in slide._element.iter(qn('p:cond')):
                if condition.get('delay')=='indefinite':condition.set('delay','0')
            entry['status']='original_inserted';entry['digest']=attachment['digest']
            report['unresolved']=[u for u in report['unresolved'] if u!=entry['key']+':media_asset']
            connections[entry['key']]={'status':'connected','source':'weekly_attachment','digest':attachment['digest'],'filename':attachment['filename'],'autoplay_configured':True,'actual_playback_verified':False}
        deck.save(file)
    report['connections']=connections
    return report

def normalize_layout_ids(file):
    """Layout IDs share one presentation-wide namespace, including copied masters."""
    from zipfile import ZipFile, ZIP_DEFLATED
    from lxml import etree as E
    with ZipFile(file) as z:parts={n:z.read(n) for n in z.namelist()}
    presentation=E.fromstring(parts['ppt/presentation.xml'])
    number=2147483647
    for node in presentation.iter(qn('p:sldMasterId')):number+=1;node.set('id',str(number))
    parts['ppt/presentation.xml']=E.tostring(presentation,encoding='UTF-8',xml_declaration=True,standalone=True)
    for name in parts:
        if name.startswith('ppt/slideMasters/') and name.endswith('.xml'):
            root=E.fromstring(parts[name])
            for node in root.iter(qn('p:sldLayoutId')):number+=1;node.set('id',str(number))
            parts[name]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    # Some edited references contain appended font fills after typefaces, or two
    # colors inside solidFill. Office repairs these even though renderers accept them.
    order=['ln','noFill','solidFill','gradFill','blipFill','pattFill','grpFill','effectLst','effectDag','highlight','uLnTx','uLn','uFillTx','uFill','latin','ea','cs','sym','hlinkClick','hlinkMouseOver','rtl','extLst']
    fill_tags={'noFill','solidFill','gradFill','blipFill','pattFill','grpFill'}
    color_tags={'scrgbClr','srgbClr','hslClr','sysClr','schemeClr','prstClr'}
    for name,blob in list(parts.items()):
        if not name.startswith('ppt/') or not name.endswith('.xml'):continue
        root=E.fromstring(blob);changed=False
        for fill in root.iter(qn('a:solidFill')):
            colors=[c for c in fill if E.QName(c).localname in color_tags]
            for extra in colors[1:]:fill.remove(extra);changed=True
        for props in root.iter():
            if props.tag not in (qn('a:rPr'),qn('a:defRPr'),qn('a:endParaRPr')):continue
            fills=[c for c in props if E.QName(c).localname in fill_tags]
            for extra in fills[1:]:props.remove(extra);changed=True
            children=list(props)
            sorted_children=sorted(children,key=lambda c:order.index(E.QName(c).localname) if E.QName(c).localname in order else len(order))
            if children!=sorted_children:
                for child in children:props.remove(child)
                for child in sorted_children:props.append(child)
                changed=True
        if changed:parts[name]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    target=Path(str(file)+'.normalized')
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
        for name,blob in parts.items():z.writestr(name,blob)
    target.replace(file)

def attach_background_music(file,report):
    """User-approved reusable Friday Zoom prayer music; independent of weekly songs."""
    import json
    from zipfile import ZipFile, BadZipFile
    import subprocess
    from pptx.util import Inches
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    directory=Path(__file__).with_name('static')/'friday-background-music'
    manifest=directory/'manifest.json'
    if not manifest.is_file():
        directory.mkdir(parents=True,exist_ok=True)
        source=Path(__file__).resolve().parents[3]/'sources/20260911_금요예배(줌).pptx'
        settings={
            'pre_service':('media1.mp3',75758),
            'before_sermon':('media4.mp3',20000),
            'after_sermon':('media7.mp3',53030),
            'intercession':('media9.mp3',20000),
            'personal':('media9.mp3',80303),
        }
        assets={}
        with ZipFile(source) as z:
            for role,(name,volume) in settings.items():
                recovered=False
                try:blob=z.read('ppt/media/'+name)
                except BadZipFile:
                    # Preserve source; recover a separately checked standalone audio copy.
                    with z.open('ppt/media/'+name) as member:
                        member._expected_crc=None;blob=member.read()
                    recovered=True
                target=directory/name
                if not target.exists():target.write_bytes(blob)
                metadata=subprocess.run(['/usr/bin/afinfo',str(target)],capture_output=True,text=True,check=True).stdout
                if 'MPG3' not in metadata:raise ValueError('기존 배경음악 형식을 확인하지 못했습니다.')
                assets[role]={'filename':name,'volume':volume,'digest':hashlib.sha256(blob).hexdigest(),'source_crc_recovery':recovered}
        manifest.write_text(json.dumps({'purpose':'금요 Zoom 고정 기도 배경음악','reuse_authorized':True,'assets':assets},ensure_ascii=False,indent=2))
    assets=json.loads(manifest.read_text())['assets']
    deck=Presentation(file);inserted=[]
    with tempfile.TemporaryDirectory() as folder:
        poster=Path(folder)/'audio.png'
        Image.new('RGB',(64,64),(80,80,80)).save(poster)
        music_groups=[]
        for entry in report['mapping']:
            if entry['kind']=='blank':continue
            key=entry['key']
            role='pre_service' if entry['kind']=='pre_service' else 'before_sermon' if '첫 기도 제목' in key else 'after_sermon' if '말씀 관련 기도 제목' in key else 'intercession' if '공동체·중보 기도 제목' in key else 'personal' if '개인 기도' in key else None
            if role is None:continue
            if music_groups and entry['start']==music_groups[-1]['end']+1 and (role==music_groups[-1]['role'] or role=='personal' and music_groups[-1]['role']=='intercession'):
                music_groups[-1]['end']=entry['end']
            else:music_groups.append({'key':key,'start':entry['start'],'end':entry['end'],'role':role})
        for entry in music_groups:
            key=entry['key'];role=entry['role']
            asset=assets[role];source=directory/asset['filename']
            if hashlib.sha256(source.read_bytes()).hexdigest()!=asset['digest']:raise ValueError('기본 배경음악이 변경되었습니다.')
            slide=deck.slides[entry['start']-1]
            # Remove the historical non-playing icon from the start template.
            for shape in list(slide.shapes):
                if 'mp3' in shape.name.lower():slide.shapes._spTree.remove(shape._element)
            audio=slide.shapes.add_movie(str(source),deck.slide_width-Inches(.65),deck.slide_height-Inches(.65),Inches(.4),Inches(.4),mime_type='audio/mpeg')
            audio.name='기도 배경음악'
            for node in audio._element.iter(qn('a:videoFile')):
                node.tag=qn('a:audioFile')
                rel=slide.part.rels[node.get(qn('r:link'))]
                rel._reltype=RT.AUDIO;rel.__dict__.pop('reltype',None)
            for node in slide._element.iter(qn('p:video')):
                node.tag=qn('p:audio')
            for node in slide._element.iter(qn('p:cMediaNode')):
                node.set('vol',str(asset['volume']));node.set('numSld',str(entry['end']-entry['start']+1))
                for timing in node.iter(qn('p:cTn')):timing.set('repeatCount','indefinite')
            for node in slide._element.iter(qn('p:cond')):
                if node.get('delay')=='indefinite':node.set('delay','0')
            inserted.append({'key':key,'page':entry['start'],'end':entry['end'],'role':role,'digest':asset['digest']})
        deck.save(file)
    report['background_music']={'source':'approved_fixed_friday_music','inserted':inserted,'actual_playback_verified':False}
    return report
