from __future__ import annotations

import re
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from media_automation.bible import BiblePassage

from .base import add_blank_slide


def _display_bible_reference(
    reference: str,
) -> str:
    return re.sub(
        r"(?<=\d)-(?=\d)",
        "~",
        reference,
    )


def _display_bible_text(text: str) -> str:
    """
    교회 표기 원칙:
    개역개정 본문의 '세례'는 화면에서 '침례'로 표시한다.
    """
    return text.replace(
        "세례",
        "침례",
    )


def _body_font_size(text: str) -> int:
    """
    4:3 예배 화면 기준.
    긴 절만 조금 줄이고 짧은 절을 과도하게 키우지 않는다.
    """
    length = len(text)

    if length >= 150:
        return 28

    if length >= 100:
        return 30

    return 32


def _validate_passage(
    passage: BiblePassage,
) -> None:
    if not passage.verses:
        raise ValueError(
            f"성경 본문에 절이 없습니다: "
            f"{passage.reference}"
        )

    numbers = [
        verse.number
        for verse in passage.verses
    ]

    if len(numbers) != len(set(numbers)):
        raise ValueError(
            f"성경 절 번호가 중복되었습니다: "
            f"{passage.reference}"
        )

    if numbers != sorted(numbers):
        raise ValueError(
            f"성경 절 순서가 올바르지 않습니다: "
            f"{passage.reference}"
        )

    for verse in passage.verses:
        if verse.end_number is not None and verse.end_number < verse.number:
            raise ValueError(f"성경 통합 절 범위가 잘못되었습니다: {passage.reference}")

    for previous_verse, current_verse in zip(passage.verses, passage.verses[1:]):
        previous = previous_verse.end_number or previous_verse.number
        current = current_verse.number
        if current != previous + 1:
            raise ValueError(
                f"성경 절이 누락되었습니다: "
                f"{passage.reference} "
                f"({previous}절 다음이 {current}절)"
            )


def add_scripture_passage_slides(
    prs: Presentation,
    passage: BiblePassage,
) -> list:
    """
    수요·주일 기본 방식:
    성경 본문을 한 절당 한 화면으로 만든다.
    """

    _validate_passage(
        passage
    )

    slides = []

    for verse in passage.verses:
        slide = add_blank_slide(
            prs
        )

        # -------------------------
        # 본문 범위
        # -------------------------

        reference_box = (
            slide.shapes.add_textbox(
                Inches(0.8),
                Inches(0.55),
                Inches(8.4),
                Inches(0.65),
            )
        )

        reference_frame = (
            reference_box.text_frame
        )

        reference_frame.clear()
        reference_frame.word_wrap = True

        reference_paragraph = (
            reference_frame.paragraphs[0]
        )

        reference_paragraph.text = (
            passage.reference
        )

        reference_paragraph.alignment = (
            PP_ALIGN.LEFT
        )

        for run in reference_paragraph.runs:
            run.font.size = Pt(28)
            run.font.bold = True
            run.font.color.rgb = RGBColor(
                0,
                0,
                0,
            )

        # -------------------------
        # 절 번호 + 본문
        # -------------------------

        body_box = (
            slide.shapes.add_textbox(
                Inches(0.8),
                Inches(1.35),
                Inches(8.4),
                Inches(5.35),
            )
        )

        body_frame = (
            body_box.text_frame
        )

        body_frame.clear()
        body_frame.word_wrap = True

        number_paragraph = (
            body_frame.paragraphs[0]
        )

        number_paragraph.text = (
            f"{verse.display_number}."
        )

        number_paragraph.alignment = (
            PP_ALIGN.LEFT
        )

        number_paragraph.line_spacing = 1.45

        for run in number_paragraph.runs:
            run.font.size = Pt(30)
            run.font.bold = True
            run.font.color.rgb = RGBColor(
                0,
                0,
                0,
            )

        body_text = _display_bible_text(
            verse.text
        )

        body_paragraph = (
            body_frame.add_paragraph()
        )

        body_paragraph.text = body_text
        body_paragraph.alignment = (
            PP_ALIGN.LEFT
        )

        body_paragraph.line_spacing = 1.45

        font_size = _body_font_size(
            body_text
        )

        for run in body_paragraph.runs:
            run.font.size = Pt(
                font_size
            )
            run.font.bold = True
            run.font.color.rgb = RGBColor(
                0,
                0,
                0,
            )

        slides.append(
            slide
        )

    return slides
