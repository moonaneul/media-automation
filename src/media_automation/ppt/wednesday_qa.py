from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation

from .wednesday_structure import (
    WednesdaySlideRange,
)
from media_automation.bible import BiblePassage
from pptx.enum.text import (
    MSO_ANCHOR,
    PP_ALIGN,
)
from pptx.dml.color import RGBColor
@dataclass(frozen=True, slots=True)
class WednesdayQaIssue:
    code: str
    message: str
    slide_number: int | None = None


@dataclass(frozen=True, slots=True)
class WednesdayQaResult:
    issues: tuple[WednesdayQaIssue, ...]

    @property
    def ok(self) -> bool:
        return not self.issues


def _slide_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
def _iter_shapes(shapes):
    for shape in shapes:
        yield shape

        child_shapes = getattr(
            shape,
            "shapes",
            None,
        )

        if child_shapes is not None:
            yield from _iter_shapes(
                child_shapes
            )


def _visible_url(text: str) -> bool:
    lowered = text.lower()

    return (
        "http://" in lowered
        or "https://" in lowered
        or "www." in lowered
    )
def _text_shapes(slide):
    return [
        shape
        for shape in slide.shapes
        if hasattr(shape, "text_frame")
        and shape.has_text_frame
        and shape.text.strip()
    ]

def validate_wednesday_structure(
    presentation_path: str | Path,
    *,
    slide_ranges: dict[
        str,
        WednesdaySlideRange,
    ],
    scripture_passages: dict[
        str,
        BiblePassage,
    ] | None = None,
) -> WednesdayQaResult:
    path = Path(
        presentation_path
    )

    prs = Presentation(
        path
    )

    issues: list[
        WednesdayQaIssue
    ] = []

    # -------------------------
    # 4:3 비율 확인
    # -------------------------

    ratio = (
        prs.slide_width
        / prs.slide_height
    )

    expected_ratio = 4 / 3

    if abs(
        ratio - expected_ratio
    ) > 0.02:
        issues.append(
            WednesdayQaIssue(
                code="INVALID_ASPECT_RATIO",
                message=(
                    "수요예배 PPT가 약 4:3 비율이 아닙니다. "
                    f"현재 비율: {ratio:.4f}"
                ),
            )
        )

    # -------------------------
    # transition 빈 화면 검사
    # -------------------------

    for key, slide_range in (
        slide_ranges.items()
    ):
        if not key.startswith(
            "transition:"
        ):
            continue

        if (
            slide_range.start
            != slide_range.end
        ):
            issues.append(
                WednesdayQaIssue(
                    code=(
                        "INVALID_TRANSITION_COUNT"
                    ),
                    message=(
                        f"{key}가 빈 화면 "
                        "1장이 아닙니다."
                    ),
                )
            )

            continue

        slide_number = (
            slide_range.start
        )

        slide = prs.slides[
            slide_number - 1
        ]

        text = _slide_text(
            slide
        ).strip()

        if text:
            issues.append(
                WednesdayQaIssue(
                    code=(
                        "TRANSITION_NOT_BLANK"
                    ),
                    message=(
                        f"{key} 빈 화면에 "
                        "텍스트가 들어 있습니다."
                    ),
                    slide_number=(
                        slide_number
                    ),
                )
            )

    # -------------------------
    # 세례 → 침례 검사
    # -------------------------

    for slide_number, slide in enumerate(
        prs.slides,
        start=1,
    ):
        text = _slide_text(
            slide
        )

        if "세례" in text:
            issues.append(
                WednesdayQaIssue(
                    code=(
                        "BAPTISM_TERM_NOT_CONVERTED"
                    ),
                    message=(
                        "'세례'가 화면에 남아 있습니다. "
                        "'침례'로 표시해야 합니다."
                    ),
                    slide_number=(
                        slide_number
                    ),
                )
            )

    # -------------------------
    # 성경 본문 화면 검사
    # -------------------------

    if scripture_passages is not None:
        for key, passage in (
            scripture_passages.items()
        ):
            slide_range = (
                slide_ranges.get(key)
            )

            if slide_range is None:
                issues.append(
                    WednesdayQaIssue(
                        code=(
                            "SCRIPTURE_RANGE_MISSING"
                        ),
                        message=(
                            f"{key}의 슬라이드 범위를 "
                            "찾을 수 없습니다."
                        ),
                    )
                )
                continue

            actual_count = (
                slide_range.end
                - slide_range.start
                + 1
            )

            expected_count = len(
                passage.verses
            )

            if actual_count != expected_count:
                issues.append(
                    WednesdayQaIssue(
                        code=(
                            "SCRIPTURE_SLIDE_COUNT_MISMATCH"
                        ),
                        message=(
                            f"{key}의 본문 슬라이드는 "
                            f"{expected_count}장이 필요하지만 "
                            f"{actual_count}장입니다."
                        ),
                    )
                )

                continue

            for offset, verse in enumerate(
                passage.verses
            ):
                slide_number = (
                    slide_range.start
                    + offset
                )

                slide = prs.slides[
                    slide_number - 1
                ]

                text = _slide_text(
                    slide
                )

                # -------------------------
                # 성경 본문 레이아웃 검사
                # -------------------------

                text_shapes = _text_shapes(
                    slide
                )

                reference_shapes = [
                    shape
                    for shape in text_shapes
                    if passage.reference
                    in shape.text
                ]

                body_shapes = [
                    shape
                    for shape in text_shapes
                    if (
                        f"{verse.number}."
                        in shape.text
                    )
                ]

                # 본문 범위 글자 크기
                if reference_shapes:
                    reference_frame = (
                        reference_shapes[0]
                        .text_frame
                    )

                    reference_sizes = []

                    for paragraph in (
                        reference_frame.paragraphs
                    ):
                        reference_sizes.extend(
                            _run_font_sizes(
                                paragraph
                            )
                        )

                    if (
                        not reference_sizes
                        or any(
                            size < 26
                            or size > 30
                            for size
                            in reference_sizes
                        )
                    ):
                        issues.append(
                            WednesdayQaIssue(
                                code=(
                                    "SCRIPTURE_REFERENCE_FONT_SIZE_INVALID"
                                ),
                                message=(
                                    "본문 범위 표기 글자 크기가 "
                                    "26~30pt 범위를 벗어났습니다."
                                ),
                                slide_number=(
                                    slide_number
                                ),
                            )
                        )

                # 절 번호 + 본문
                if body_shapes:
                    body_frame = (
                        body_shapes[0]
                        .text_frame
                    )

                    for paragraph in (
                        body_frame.paragraphs
                    ):
                        if not paragraph.text.strip():
                            continue

                        sizes = (
                            _run_font_sizes(
                                paragraph
                            )
                        )

                        if (
                            not sizes
                            or any(
                                size < 28
                                or size > 34
                                for size in sizes
                            )
                        ):
                            issues.append(
                                WednesdayQaIssue(
                                    code=(
                                        "SCRIPTURE_BODY_FONT_SIZE_INVALID"
                                    ),
                                    message=(
                                        "성경 본문 글자 크기가 "
                                        "28~34pt 범위를 벗어났습니다."
                                    ),
                                    slide_number=(
                                        slide_number
                                    ),
                                )
                            )

                        spacing = (
                            paragraph.line_spacing
                        )

                        if (
                            spacing is None
                            or spacing < 1.4
                            or spacing > 1.5
                        ):
                            issues.append(
                                WednesdayQaIssue(
                                    code=(
                                        "SCRIPTURE_LINE_SPACING_INVALID"
                                    ),
                                    message=(
                                        "성경 본문 줄간격이 "
                                        "1.4~1.5 범위가 아닙니다."
                                    ),
                                    slide_number=(
                                        slide_number
                                    ),
                                )
                            )

                        for run in paragraph.runs:
                            if (
                                run.text.strip()
                                and not _is_black(run)
                            ):
                                issues.append(
                                    WednesdayQaIssue(
                                        code=(
                                            "SCRIPTURE_TEXT_COLOR_INVALID"
                                        ),
                                        message=(
                                            "성경 본문 글자색이 "
                                            "검정이 아닙니다."
                                        ),
                                        slide_number=(
                                            slide_number
                                        ),
                                    )
                                )

                if passage.reference not in text:
                    issues.append(
                        WednesdayQaIssue(
                            code=(
                                "SCRIPTURE_REFERENCE_MISMATCH"
                            ),
                            message=(
                                f"{key}의 본문 범위 표기가 "
                                f"'{passage.reference}'와 "
                                "일치하지 않습니다."
                            ),
                            slide_number=(
                                slide_number
                            ),
                        )
                    )

                expected_number = (
                    f"{verse.number}."
                )

                if expected_number not in text:
                    issues.append(
                        WednesdayQaIssue(
                            code=(
                                "SCRIPTURE_VERSE_NUMBER_MISMATCH"
                            ),
                            message=(
                                f"{key}에서 "
                                f"{verse.number}절 번호를 "
                                "찾을 수 없습니다."
                            ),
                            slide_number=(
                                slide_number
                            ),
                        )
                    )
    # -------------------------
    # 단독 안내 화면 정렬 검사
    # -------------------------

    for key in (
        "prayer",
        "sermon_title",
    ):
        slide_range = slide_ranges.get(
            key
        )

        if slide_range is None:
            continue

        if (
            slide_range.start
            != slide_range.end
        ):
            issues.append(
                WednesdayQaIssue(
                    code=(
                        "STANDALONE_SLIDE_COUNT_INVALID"
                    ),
                    message=(
                        f"{key}는 한 화면이어야 합니다."
                    ),
                )
            )
            continue

        slide_number = (
            slide_range.start
        )

        slide = prs.slides[
            slide_number - 1
        ]

        shapes = _text_shapes(
            slide
        )

        if not shapes:
            issues.append(
                WednesdayQaIssue(
                    code=(
                        "STANDALONE_TEXT_MISSING"
                    ),
                    message=(
                        f"{key} 화면에 텍스트가 없습니다."
                    ),
                    slide_number=(
                        slide_number
                    ),
                )
            )
            continue

        for shape in shapes:
            frame = shape.text_frame

            if (
                frame.vertical_anchor
                != MSO_ANCHOR.MIDDLE
            ):
                issues.append(
                    WednesdayQaIssue(
                        code=(
                            "STANDALONE_NOT_VERTICALLY_CENTERED"
                        ),
                        message=(
                            f"{key} 텍스트가 "
                            "세로 가운데 정렬이 아닙니다."
                        ),
                        slide_number=(
                            slide_number
                        ),
                    )
                )

            for paragraph in (
                frame.paragraphs
            ):
                if (
                    paragraph.text.strip()
                    and paragraph.alignment
                    != PP_ALIGN.CENTER
                ):
                    issues.append(
                        WednesdayQaIssue(
                            code=(
                                "STANDALONE_NOT_HORIZONTALLY_CENTERED"
                            ),
                            message=(
                                f"{key} 텍스트가 "
                                "가로 가운데 정렬이 아닙니다."
                            ),
                            slide_number=(
                                slide_number
                            ),
                        )
                    )
    # -------------------------
    # 생성 화면 파란 배경 도형 검사
    # -------------------------

    generated_keys = (
        "prayer",
        "scripture",
        "sermon_title",
        "additional_scripture",
    )

    for key in generated_keys:
        slide_range = slide_ranges.get(key)

        if slide_range is None:
            continue

        for slide_number in range(
            slide_range.start,
            slide_range.end + 1,
        ):
            slide = prs.slides[
                slide_number - 1
            ]

            for shape in slide.shapes:
                if _is_blue_fill(shape):
                    issues.append(
                        WednesdayQaIssue(
                            code="BLUE_BACKGROUND_SHAPE_FOUND",
                            message=(
                                f"{key} 화면에 "
                                "파란색 배경 도형이 있습니다."
                            ),
                            slide_number=slide_number,
                        )
                    )

    # -------------------------
    # URL / 하이퍼링크 검사
    # -------------------------

    for slide_number, slide in enumerate(
        prs.slides,
        start=1,
    ):
        for shape in _iter_shapes(
            slide.shapes
        ):
            # 도형 클릭 링크
            try:
                address = (
                    shape.click_action
                    .hyperlink
                    .address
                )
            except (
                AttributeError,
                TypeError,
                ValueError,
            ):
                address = None

            if address:
                issues.append(
                    WednesdayQaIssue(
                        code=(
                            "HYPERLINK_FOUND"
                        ),
                        message=(
                            "불필요한 도형 "
                            f"하이퍼링크가 있습니다: "
                            f"{address}"
                        ),
                        slide_number=(
                            slide_number
                        ),
                    )
                )

            if not getattr(
                shape,
                "has_text_frame",
                False,
            ):
                continue

            # 보이는 URL
            if _visible_url(
                shape.text
            ):
                issues.append(
                    WednesdayQaIssue(
                        code="VISIBLE_URL_FOUND",
                        message=(
                            "화면 텍스트에 URL이 "
                            "포함되어 있습니다."
                        ),
                        slide_number=(
                            slide_number
                        ),
                    )
                )

            # 텍스트에 걸린 링크
            for paragraph in (
                shape.text_frame.paragraphs
            ):
                for run in paragraph.runs:
                    try:
                        address = (
                            run.hyperlink.address
                        )
                    except (
                        AttributeError,
                        TypeError,
                        ValueError,
                    ):
                        address = None

                    if address:
                        issues.append(
                            WednesdayQaIssue(
                                code=(
                                    "HYPERLINK_FOUND"
                                ),
                                message=(
                                    "텍스트에 "
                                    "하이퍼링크가 있습니다: "
                                    f"{address}"
                                ),
                                slide_number=(
                                    slide_number
                                ),
                            )
                        )

    return WednesdayQaResult(
        issues=tuple(
            issues
        )
    )
def _run_font_sizes(paragraph) -> list[float]:
    sizes = []

    for run in paragraph.runs:
        if run.font.size is not None:
            sizes.append(
                run.font.size.pt
            )

    return sizes


def _is_black(run) -> bool:
    color = run.font.color

    if color.type is None:
        return False

    try:
        return (
            color.rgb
            == RGBColor(0, 0, 0)
        )
    except AttributeError:
        return False
def _is_blue_fill(shape) -> bool:
    try:
        fill = shape.fill
    except AttributeError:
        return False

    if fill.type is None:
        return False

    try:
        rgb = fill.fore_color.rgb
    except (AttributeError, TypeError):
        return False

    if rgb is None:
        return False

    red = rgb[0]
    green = rgb[1]
    blue = rgb[2]

    return (
        blue > red + 40
        and blue > green + 20
    )