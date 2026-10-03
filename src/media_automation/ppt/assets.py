from __future__ import annotations

from pathlib import Path
from typing import Protocol

from pptx import Presentation


class PreServiceSlideProvider(Protocol):
    def create_presentation(
        self,
        key: str,
    ) -> Presentation:
        ...


class FilePreServiceSlideProvider:
    def __init__(
        self,
        sources: dict[str, Path],
        *,
        slide_counts: dict[str, int] | None = None,
    ) -> None:
        self._sources = {
            key: Path(path)
            for key, path in sources.items()
        }

        self._slide_counts = (
            dict(slide_counts)
            if slide_counts is not None
            else {}
        )

    def create_presentation(
        self,
        key: str,
    ) -> Presentation:
        source = self._sources.get(
            key
        )

        if source is None:
            raise KeyError(
                f"등록된 예배 전 안내 원본이 없습니다: "
                f"{key}"
            )

        if not source.exists():
            raise FileNotFoundError(
                f"예배 전 안내 원본 파일이 없습니다: "
                f"{source}"
            )

        prs = Presentation(
            source
        )

        keep_count = self._slide_counts.get(
            key,
            len(prs.slides),
        )

        if keep_count < 0:
            raise ValueError(
                "slide_count는 0 이상이어야 합니다."
            )

        while len(prs.slides) > keep_count:
            slide_id = (
                prs.slides._sldIdLst[-1]
            )

            prs.part.drop_rel(
                slide_id.rId
            )

            prs.slides._sldIdLst.remove(
                slide_id
            )

        return prs