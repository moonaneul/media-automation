from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation

from media_automation.bible import BibleProvider
from media_automation.planning import (
    BlockKind,
    WorshipBlock,
)

from .assets import PreServiceSlideProvider
from .base import (
    add_blank_slide,
    create_4x3_presentation,
)
from .renderer import (
    MissingRenderDependencyError,
    render_block,
)
from .song_assets import (
    MissingSongAssetError,
    SongAsset,
    SongAssetProvider,
)


@dataclass(frozen=True, slots=True)
class WednesdaySlideRange:
    start: int
    end: int


@dataclass(frozen=True, slots=True)
class WednesdayStructureResult:
    output_path: Path
    slide_ranges: dict[
        str,
        WednesdaySlideRange,
    ]


def _get_song_slide_count(
    asset: SongAsset,
) -> int:
    if (
        asset.start_slide is not None
        and asset.end_slide is not None
    ):
        return (
            asset.end_slide
            - asset.start_slide
            + 1
        )

    source = Presentation(
        asset.path
    )

    return len(source.slides)


def build_wednesday_structure(
    plan: list[WorshipBlock],
    output_path: str | Path,
    *,
    bible_provider: BibleProvider,
    pre_service_provider: PreServiceSlideProvider,
    song_asset_provider: SongAssetProvider,
    skip_missing_songs: bool = True,
) -> WednesdayStructureResult:
    """
    Mac에서도 확인할 수 있는
    수요예배 구조 프리뷰를 만든다.

    예배 전 안내:
        실제 기존 PPT의 지정 슬라이드를 사용한다.

    찬양:
        악보 PPT의 실제 슬라이드 수만 계산하고
        동일한 장수의 빈 placeholder를 만든다.

    나머지:
        실제 생성 renderer를 사용한다.

    따라서 악보 자체는 보이지 않지만
    최종 슬라이드 번호와 순서 구조는
    실제 제작 결과와 동일하게 계산할 수 있다.
    """

    output_path = Path(
        output_path
    )

    resolved_song_assets: dict[
        int,
        SongAsset,
    ] = {}

    # -------------------------
    # 찬양 자산 사전 확인
    # -------------------------

    for index, block in enumerate(
        plan
    ):
        if block.kind != BlockKind.SONG:
            continue

        title = block.value.title

        try:
            resolved_song_assets[
                index
            ] = (
                song_asset_provider
                .get_song_asset(
                    title
                )
            )

        except MissingSongAssetError:
            if not skip_missing_songs:
                raise

    # -------------------------
    # 누락 찬양 반영 후
    # 실제 plan 정리
    # -------------------------

    effective_plan: list[
        tuple[int, WorshipBlock]
    ] = []

    for index, block in enumerate(
        plan
    ):
        if (
            block.kind
            == BlockKind.SONG
            and index
            not in resolved_song_assets
            and skip_missing_songs
        ):
            continue

        if block.kind == BlockKind.BLANK:
            if not effective_plan:
                continue

            previous = (
                effective_plan[-1][1]
            )

            if previous.kind in {
                BlockKind.BLANK,
                BlockKind.PRE_SERVICE,
            }:
                continue

        effective_plan.append(
            (
                index,
                block,
            )
        )

    if (
        effective_plan
        and effective_plan[-1][1].kind
        == BlockKind.BLANK
    ):
        effective_plan.pop()

    # -------------------------
    # 예배 전 안내
    # -------------------------

    pre_service_blocks = [
        block
        for _, block in effective_plan
        if block.kind
        == BlockKind.PRE_SERVICE
    ]

    if len(pre_service_blocks) > 1:
        raise ValueError(
            "PRE_SERVICE 블록은 "
            "하나만 사용할 수 있습니다."
        )

    slide_ranges: dict[
        str,
        WednesdaySlideRange,
    ] = {}

    if pre_service_blocks:
        block = pre_service_blocks[0]

        prs = (
            pre_service_provider
            .create_presentation(
                block.key
            )
        )

        count = len(
            prs.slides
        )

        if count > 0:
            slide_ranges[
                block.key
            ] = WednesdaySlideRange(
                start=1,
                end=count,
            )

    else:
        prs = (
            create_4x3_presentation()
        )

    # -------------------------
    # 나머지 순서 생성
    # -------------------------

    for index, block in effective_plan:
        if (
            block.kind
            == BlockKind.PRE_SERVICE
        ):
            continue

        before = len(
            prs.slides
        )

        if block.kind == BlockKind.SONG:
            asset = (
                resolved_song_assets[
                    index
                ]
            )

            slide_count = (
                _get_song_slide_count(
                    asset
                )
            )

            for _ in range(
                slide_count
            ):
                add_blank_slide(
                    prs
                )

        else:
            render_block(
                prs,
                block,
                bible_provider=(
                    bible_provider
                ),
            )

        after = len(
            prs.slides
        )

        if after > before:
            slide_ranges[
                block.key
            ] = WednesdaySlideRange(
                start=before + 1,
                end=after,
            )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    prs.save(
        output_path
    )

    return WednesdayStructureResult(
        output_path=output_path,
        slide_ranges=slide_ranges,
    )