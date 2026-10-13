"""Reuse the approved original Zoom slides; replace weekly text only."""
from copy import deepcopy
from pathlib import Path
import tempfile
import re
import unicodedata
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from .slide_merge import OpenXmlSlideMerger

VERSION='friday-original-layout-readable-centered-v7'


def replace_paragraph(paragraph, text):
    # Retain paragraph and first-run formatting, including inherited styles.
    prototype=next(iter(paragraph.runs),None)
    props=deepcopy(prototype._r.find(qn('a:rPr'))) if prototype is not None and prototype._r.find(qn('a:rPr')) is not None else None
    for child in list(paragraph._p):
        if child.tag!=qn('a:pPr'):paragraph._p.remove(child)
    paragraph.text=text
    for run in paragraph.runs:
        if props is not None:run._r.insert(0,deepcopy(props))
    for br in paragraph._p.findall(qn('a:br')):
        if props is not None:br.append(deepcopy(props))


def wrap_title(text, width=48):
    lines=[];line=''
    for word in text.split():
        candidate=(line+' '+word).strip()
        units=sum(2 if unicodedata.east_asian_width(c) in ('W','F') else 1 for c in candidate)
        if line and units>width:lines.append(line);line=word
        else:line=candidate
    return '\v'.join(lines+[line])


def prayer_body_text(body):
    from media_automation.input_text import clean_leading_marks
    return '\n'.join('- '+line for line in clean_leading_marks(body).splitlines() if line.strip())


def replace_prayer(shape, title, body):
    frame=shape.text_frame
    source=frame.paragraphs[0]
    title_props=next((deepcopy(r._r.find(qn('a:rPr'))) for r in source.runs if r._r.find(qn('a:rPr')) is not None),None)
    body_props=next((deepcopy(r._r.find(qn('a:rPr'))) for p in frame.paragraphs for r in p.runs if r.font.size and r.font.size.pt==32),title_props)
    replace_paragraph(source,wrap_title(title,width=40 if title_props is not None and title_props.get('sz')=='3800' else 44))
    source._p.get_or_add_pPr().set('eaLnBrk','0')
    for p in list(frame._txBody.findall(qn('a:p')))[1:]:frame._txBody.remove(p)
    # Original prayer separator consists of 32pt, 8pt, and 4pt breaks.
    for size in (3200,800,400):
        br=OxmlElement('a:br');rpr=OxmlElement('a:rPr');rpr.set('sz',str(size));br.append(rpr);source._p.append(br)
    body_lines=[wrapped for line in prayer_body_text(body).splitlines()
                for wrapped in wrap_title(line,width=46).split('\v')]
    for index,line in enumerate(body_lines):
        if index:
            br=OxmlElement('a:br')
            if body_props is not None:br.append(deepcopy(body_props))
            source._p.append(br)
        run=source.add_run();run.text=('    ' if index else '  ')+line
        if body_props is not None:run._r.insert(0,deepcopy(body_props))
    shape.width=Inches(10.78)
    shape.left=Inches(1.65)
    shape.height=Inches(4.8)



def background_only_slide(slide):
    """Keep the original theme background without weekly text or decorations."""
    # Resolve inherited backgrounds before suppressing master graphics.
    inherited = next((node for node in (
        slide._element.find(qn('p:cSld')).find(qn('p:bg')),
        slide.slide_layout._element.find(qn('p:cSld')).find(qn('p:bg')),
        slide.slide_layout.slide_master._element.find(qn('p:cSld')).find(qn('p:bg')),
    ) if node is not None), None)
    csld = slide._element.find(qn('p:cSld'))
    existing = csld.find(qn('p:bg'))
    if existing is not None: csld.remove(existing)
    if inherited is not None: csld.insert(0, deepcopy(inherited))
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    # Layouts may contain visible inherited shapes too; this is a private copy.
    for shape in list(slide.slide_layout.shapes):
        # Original cover's full-slide gradient rectangle is the visible background.
        if shape.left == 0 and shape.top == 0 and shape.width >= slide.part.package.presentation_part.presentation.slide_width and shape.height >= slide.part.package.presentation_part.presentation.slide_height:
            slide.shapes._spTree.append(deepcopy(shape._element))
        shape._element.getparent().remove(shape._element)
    slide._element.set('showMasterSp', '0')


def normalize_reference(text):
    text=re.sub(r'\(\s*', '(', text.strip())
    return re.sub(r'\s*\)', ')', text)


def center_personal(slide, height):
    items=[s for s in slide.shapes if not (s.has_text_frame and s.text.strip().isdigit())]
    top=min(s.top for s in items);bottom=max(s.top+s.height for s in items)
    shift=(height-(bottom-top))//2-top
    for s in items:s.top+=shift


def airy_bible(target, ps):
    # Keep original fonts, but replace spacer paragraphs with explicit spacing.
    for index in (0,3,4,5,6,9,12):
        p=ps[index]._p;p.getparent().remove(p)
    ps[1].space_after=Pt(8)
    ps[2].space_after=Pt(12)
    for number,body in ((7,8),(10,11)):
        ps[number].space_before=Pt(8);ps[number].space_after=Pt(4)
        ps[body].line_spacing=1.35;ps[body].space_after=Pt(8)
    target.height=Inches(6.65)


def apply_reference_layout(path, mapping, date, directory):
    directory=Path(directory)
    if not (directory/'cover.pptx').is_file():return False
    original=Presentation(path)
    targets={1:('cover',None)}
    for entry in mapping:
        if entry['kind']=='blank' or entry['status'] in ('missing_song_asset','input_pending','missing_verified_bible_text'):
            for page in range(entry['start'],entry['end']+1):targets[page]=('blank',entry)
            continue
        if entry['status'] not in ('rendered','bible_inserted'):continue
        for page in range(entry['start'],entry['end']+1):
            if not original.slides[page-1].shapes:targets[page]=('blank',entry)
        key=entry['key']
        role='bible' if entry['status']=='bible_inserted' else 'sermon' if key=='설교 제목' else 'personal' if '개인 기도' in key else 'community-prayer' if '공동체' in key else 'word-prayer' if '말씀 관련' in key else 'prayer' if '기도' in key else None
        if role:
            for page in range(entry['start'],entry['end']+1):targets[page]=(role,entry)
    with tempfile.TemporaryDirectory(dir=Path(path).parent) as scratch:
        for page,(role,entry) in sorted(targets.items()):
            src=directory/(('cover' if role=='blank' else role)+'.pptx')
            if not src.is_file():continue
            ref=Presentation(src);slide=ref.slides[0]
            old=original.slides[page-1]
            texts=[s.text for s in old.shapes if s.has_text_frame and s.text]
            shapes=[s for s in slide.shapes if s.has_text_frame and s.text]
            # Always update visible dates and page numbers in reused footer objects.
            for s in shapes:
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        for old_date in ('2026.09.18','2026.09.11'):
                            if old_date in r.text:r.text=r.text.replace(old_date,date.replace('-','.'))
                if s.text.strip().isdigit():replace_paragraph(s.text_frame.paragraphs[0],str(page))
            if role=='blank':background_only_slide(slide)
            elif role=='cover':pass
            elif role in ('prayer','word-prayer','community-prayer'):
                candidates=[s for s in shapes if s.text.lstrip().startswith(('1.','2.','3.'))]
                keep=next((s for s in candidates if s.text.lstrip().startswith('1.')),candidates[0])
                for s in candidates:
                    if s is not keep:s._element.getparent().remove(s._element)
                content=texts[0].splitlines();title=content[0].replace('\v',' ')
                # Generated soft breaks are layout artifacts, not content changes.
                title=' '.join(texts[0].split('\n\n')[0].split())
                body=' '.join(texts[0].split('\n\n',1)[1].split()) if '\n\n' in texts[0] else ''
                replace_prayer(keep,title,body)
            elif role=='sermon':
                target=shapes[0];parts=target.text_frame.paragraphs
                replace_paragraph(parts[0],texts[0].strip());replace_paragraph(parts[1],normalize_reference(texts[1]) if len(texts)>1 else '')
                for p in parts:p.alignment=PP_ALIGN.CENTER
                target.top=(ref.slide_height-target.height)//2
            elif role=='personal':
                target=next(s for s in shapes if s.text.startswith('지금은'))
                replace_paragraph(target.text_frame.paragraphs[0],texts[0])
                target=next(s for s in shapes if s.text.startswith('더 깊은'))
                replace_paragraph(target.text_frame.paragraphs[0],texts[1].replace('나아가시기 바랍니다.','나아가시기\v바랍니다.'))
                center_personal(slide,ref.slide_height)
            elif role=='bible':
                target=next(s for s in shapes if '왕상 9:' in s.text)
                ps=target.text_frame.paragraphs
                replace_paragraph(ps[1],texts[0])
                verse=[p.text for s in old.shapes if s.has_text_frame and s.text and s.text!=texts[0] for p in s.text_frame.paragraphs if p.text]
                for i,j in enumerate([7,8,10,11]):replace_paragraph(ps[j],verse[i] if i<len(verse) else '')
            if role in ('bible','sermon'):
                target.width=Inches(10.78)
                target.left=Inches(1.65) if role=='bible' else (ref.slide_width-target.width)//2
                if role=='bible':
                    for j in (8,11):replace_paragraph(ps[j],wrap_title(ps[j].text,width=44))
                    airy_bible(target,ps)
            tmp=Path(scratch)/'reference.pptx';ref.save(tmp)
            OpenXmlSlideMerger().insert_all(path,tmp,after_slide=page-1)
            deck=Presentation(path);element=deck.slides._sldIdLst[page];deck.part.drop_rel(element.rId);deck.slides._sldIdLst.remove(element);deck.save(path)
    return True
