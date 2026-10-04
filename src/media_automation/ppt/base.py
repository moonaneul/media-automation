from __future__ import annotations

from pptx import Presentation
from pptx.enum.text import (
    MSO_ANCHOR,
    PP_ALIGN,
)
from pptx.util import Inches, Pt


def create_4x3_presentation() -> Presentation:
    prs = Presentation()

    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # python-pptx 기본 Presentation에는
    # 첫 슬라이드가 없으므로 그대로 사용한다.
    return prs

def create_16x9_presentation() -> Presentation:
    prs = Presentation()

    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    return prs
    
def get_blank_layout(prs: Presentation):
    """
    특정 템플릿에서 blank layout의 위치가 달라도
    최대한 안전하게 빈 layout을 찾는다.
    """

    preferred_names = {
        "blank",
        "빈 화면",
        "빈 슬라이드",
    }

    for layout in prs.slide_layouts:
        name = (
            getattr(layout, "name", "")
            or ""
        ).strip().lower()

        if name in preferred_names:
            return layout

    # 이름으로 찾지 못하면
    # placeholder가 하나도 없는 layout을 사용한다.
    for layout in prs.slide_layouts:
        if len(layout.placeholders) == 0:
            return layout

    raise ValueError(
        "빈 슬라이드 레이아웃을 찾을 수 없습니다."
    )


def add_blank_slide(
    prs: Presentation,
):
    layout = get_blank_layout(prs)
    return prs.slides.add_slide(layout)


def add_centered_text_slide(
    prs: Presentation,
    text: str,
    *,
    font_size: int = 34,
):
    slide = add_blank_slide(prs)

    textbox = slide.shapes.add_textbox(
        Inches(0.8),
        Inches(0.8),
        Inches(8.4),
        Inches(5.9),
    )

    frame = textbox.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.word_wrap = True

    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.CENTER

    for run in paragraph.runs:
        run.font.size = Pt(font_size)

    return slide


def add_sermon_title_slide(
    prs: Presentation,
    title: str,
    scripture_reference: str,
):
    slide = add_blank_slide(prs)

    textbox = slide.shapes.add_textbox(
        Inches(0.8),
        Inches(0.8),
        Inches(8.4),
        Inches(5.9),
    )

    frame = textbox.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.word_wrap = True

    reference_paragraph = frame.paragraphs[0]
    reference_paragraph.text = scripture_reference
    reference_paragraph.alignment = PP_ALIGN.CENTER

    for run in reference_paragraph.runs:
        run.font.size = Pt(28)

    title_paragraph = frame.add_paragraph()
    title_paragraph.text = title
    title_paragraph.alignment = PP_ALIGN.CENTER
    title_paragraph.space_before = Pt(18)

    for run in title_paragraph.runs:
        run.font.size = Pt(34)

    return slide