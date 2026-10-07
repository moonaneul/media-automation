from __future__ import annotations

from functools import lru_cache
from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import (
    MSO_ANCHOR,
    PP_ALIGN,
)
from pptx.util import Inches, Pt

from media_automation.bible import (
    BiblePassage,
    BibleProvider,
)
from media_automation.planning import (
    BlockKind,
    WorshipBlock,
)

from .base import add_blank_slide
from .renderer import (
    MissingRenderDependencyError,
)
from .scripture import (
    _display_bible_text,
    _validate_passage,
)
from .scripture import _display_bible_reference

@lru_cache(maxsize=1)
def _zoom_background_image() -> bytes:
    """
    기존 금요 Zoom 자료의
    청록 → 적갈색 → 금색 계열 배경을
    간단한 4-corner gradient로 재현한다.
    """

    width = 640
    height = 360

    image = Image.new(
        "RGB",
        (width, height),
    )

    pixels = image.load()

    top_left = (119, 78, 83)
    top_right = (190, 151, 78)
    bottom_left = (48, 91, 94)
    bottom_right = (164, 103, 82)

    for y in range(height):
        ty = y / (height - 1)

        for x in range(width):
            tx = x / (width - 1)

            color = []

            for index in range(3):
                top = (
                    top_left[index]
                    * (1 - tx)
                    + top_right[index]
                    * tx
                )

                bottom = (
                    bottom_left[index]
                    * (1 - tx)
                    + bottom_right[index]
                    * tx
                )

                value = (
                    top * (1 - ty)
                    + bottom * ty
                )

                color.append(int(value))

            pixels[x, y] = tuple(color)

    buffer = BytesIO()

    image.save(
        buffer,
        format="PNG",
    )

    return buffer.getvalue()


def add_zoom_background(
    prs: Presentation,
    slide,
):
    slide.shapes.add_picture(
        BytesIO(
            _zoom_background_image()
        ),
        0,
        0,
        width=prs.slide_width,
        height=prs.slide_height,
    )


def add_zoom_rule(
    slide,
    *,
    left: float,
    top: float,
    width: float,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(0.025),
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(
        235,
        235,
        235,
    )

    shape.line.fill.background()

    return shape

def add_zoom_centered_text_slide(
    prs: Presentation,
    text: str,
    *,
    font_size: int = 36,
):
    slide = add_blank_slide(prs)
    add_zoom_background(
        prs,
        slide,
    )
    textbox = slide.shapes.add_textbox(
        Inches(1.2),
        Inches(0.8),
        Inches(10.93),
        Inches(5.9),
    )

    frame = textbox.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.word_wrap = True

    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = PP_ALIGN.CENTER

    for run in paragraph.runs:
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(
            255,
            255,
            255,
        )

    return slide


def add_prayer_topics_slide(
    prs: Presentation,
    topics: list[str],
):
    slide = add_blank_slide(prs)

    add_zoom_background(
        prs,
        slide,
    )

    for index, topic in enumerate(
        topics,
        start=1,
    ):
        # ?? ?? Zoom ????
        # ? ?? ??? ?? ??? ?? ??.
        # ?? ????? PowerPoint animation??
        # ??? ?? ????.
        textbox = slide.shapes.add_textbox(
            Inches(1.40),
            Inches(1.82),
            Inches(10.45),
            Inches(4.35),
        )

        textbox.name = (
            f"FridayZoomPrayerTopic{index}"
        )

        frame = textbox.text_frame
        frame.clear()
        frame.vertical_anchor = (
            MSO_ANCHOR.MIDDLE
        )
        frame.word_wrap = True

        paragraph = frame.paragraphs[0]

        paragraph.text = (
            f"{index}. {topic}"
        )

        paragraph.alignment = (
            PP_ALIGN.LEFT
        )

        paragraph.line_spacing = 1.45

        for run in paragraph.runs:
            run.font.size = Pt(
                36 if len(topics) >= 3 else 38
            )

            run.font.color.rgb = RGBColor(
                255,
                255,
                255,
            )

    return slide




def add_zoom_sermon_title_slide(
    prs: Presentation,
    title: str,
    scripture_reference: str,
):
    scripture_reference = _display_bible_reference(
        scripture_reference
    )

    slide = add_blank_slide(prs)

    add_zoom_background(
        prs,
        slide,
    )

    # ?? ??
    title_box = slide.shapes.add_textbox(
        Inches(1.1),
        Inches(2.87),
        Inches(11.13),
        Inches(0.85),
    )

    title_frame = title_box.text_frame
    title_frame.clear()
    title_frame.vertical_anchor = (
        MSO_ANCHOR.MIDDLE
    )
    title_frame.word_wrap = True

    title_paragraph = (
        title_frame.paragraphs[0]
    )

    title_paragraph.text = title
    title_paragraph.alignment = (
        PP_ALIGN.CENTER
    )

    for run in title_paragraph.runs:
        run.font.size = Pt(42)
        run.font.color.rgb = RGBColor(
            255,
            255,
            255,
        )

    # ?? ??
    reference_box = (
        slide.shapes.add_textbox(
            Inches(3.0),
            Inches(3.93),
            Inches(7.33),
            Inches(0.55),
        )
    )

    reference_frame = (
        reference_box.text_frame
    )

    reference_frame.clear()

    reference_paragraph = (
        reference_frame.paragraphs[0]
    )

    reference_paragraph.text = (
        f"({scripture_reference})"
    )

    reference_paragraph.alignment = (
        PP_ALIGN.CENTER
    )

    for run in reference_paragraph.runs:
        run.font.size = Pt(25)
        run.font.color.rgb = RGBColor(
            255,
            255,
            255,
        )

    return slide



def _group_zoom_verses(
    passage: BiblePassage,
    *,
    max_verses: int = 3,
    max_lines: int = 7,
    characters_per_line: int = 30,
):
    """
    금요 Zoom 본문을 실제 화면 분량에 맞춰 묶는다.

    2026-09-04 / 09-11 / 09-18 실제 PPT를 기준으로
    한 화면에 최대 3절까지 허용하되,
    예상 줄 수가 너무 많으면 먼저 분리한다.

    실제 렌더링 QA는 별도 수행한다.
    """

    groups = []
    current = []
    current_lines = 0

    for verse in passage.verses:
        display_text = _display_bible_text(
            verse.text
        )

        estimated_lines = max(
            1,
            (
                len(display_text)
                + characters_per_line
                - 1
            )
            // characters_per_line,
        )

        too_many_verses = (
            current
            and len(current) >= max_verses
        )

        too_many_lines = (
            current
            and (
                current_lines
                + estimated_lines
                > max_lines
            )
        )

        if (
            too_many_verses
            or too_many_lines
        ):
            groups.append(current)
            current = []
            current_lines = 0

        current.append(verse)
        current_lines += estimated_lines

    if current:
        groups.append(current)

    return groups

def add_zoom_scripture_reference_slide(
    prs: Presentation,
    reference: str,
):
    """
    기존 금요 Zoom PPT의 본문 안내 화면.

    실제 자료처럼 본문 범위만
    한 화면에 크게 표시한다.
    """

    return add_zoom_centered_text_slide(
        prs,
        reference,
        font_size=54,
    )

def add_zoom_scripture_passage_slides(
    prs: Presentation,
    passage: BiblePassage,
):
    _validate_passage(
        passage
    )

    groups = _group_zoom_verses(
        passage
    )

    slides = []

    for group in groups:
        slide = add_blank_slide(prs)

        add_zoom_background(
            prs,
            slide,
        )

        # ----------------------------------------
        # ?? ??
        # ----------------------------------------

        reference_box = (
            slide.shapes.add_textbox(
                Inches(4.0),
                Inches(0.38),
                Inches(5.3),
                Inches(0.58),
            )
        )

        reference_frame = (
            reference_box.text_frame
        )

        reference_frame.clear()
        reference_frame.word_wrap = True

        reference_paragraph = (
            reference_frame.paragraphs[0]
        )

        reference_paragraph.text = (
            _display_bible_reference(
                passage.reference
            )
        )

        reference_paragraph.alignment = (
            PP_ALIGN.CENTER
        )

        for run in reference_paragraph.runs:
            run.font.size = Pt(28)

            run.font.color.rgb = (
                RGBColor(
                    255,
                    255,
                    255,
                )
            )

        add_zoom_rule(
            slide,
            left=3.95,
            top=1.00,
            width=5.4,
        )

        # ----------------------------------------
        # ??
        # ----------------------------------------

        body_box = (
            slide.shapes.add_textbox(
                Inches(1.05),
                Inches(1.42),
                Inches(11.20),
                Inches(5.45),
            )
        )

        body_frame = (
            body_box.text_frame
        )

        body_frame.clear()
        body_frame.word_wrap = True

        total_length = sum(
            len(
                _display_bible_text(
                    verse.text
                )
            )
            for verse in group
        )

        # ?? 26~28pt?? ??.
        # ? ??? ???? ?? ???.
        if total_length >= 260:
            body_font_size = 27
        elif total_length >= 190:
            body_font_size = 29
        else:
            body_font_size = 31

        first_paragraph = True

        for verse_index, verse in enumerate(
            group
        ):
            # ? ??
            number_paragraph = (
                body_frame.paragraphs[0]
                if first_paragraph
                else body_frame.add_paragraph()
            )

            first_paragraph = False

            number_paragraph.text = (
                verse.display_number
            )

            number_paragraph.alignment = (
                PP_ALIGN.LEFT
            )

            number_paragraph.space_before = (
                Pt(
                    8
                    if verse_index > 0
                    else 0
                )
            )

            number_paragraph.space_after = (
                Pt(1)
            )

            for run in number_paragraph.runs:
                run.font.size = Pt(18)

                run.font.color.rgb = (
                    RGBColor(
                        255,
                        255,
                        255,
                    )
                )

            # ??
            paragraph = (
                body_frame.add_paragraph()
            )

            paragraph.text = (
                _display_bible_text(
                    verse.text
                )
            )

            paragraph.alignment = (
                PP_ALIGN.LEFT
            )

            paragraph.line_spacing = 1.18
            paragraph.space_after = Pt(8)

            for run in paragraph.runs:
                run.font.size = Pt(
                    body_font_size
                )

                run.font.color.rgb = (
                    RGBColor(
                        255,
                        255,
                        255,
                    )
                )

        slides.append(
            slide
        )

    return slides


def add_personal_prayer_slide(
    prs: Presentation,
    text: str | None = None,
):
    slide = add_blank_slide(prs)

    add_zoom_background(
        prs,
        slide,
    )

    add_zoom_rule(
        slide,
        left=4.4,
        top=2.65,
        width=4.55,
    )

    textbox = slide.shapes.add_textbox(
        Inches(4.0),
        Inches(2.9),
        Inches(5.35),
        Inches(1.0),
    )

    frame = textbox.text_frame
    frame.clear()
    frame.vertical_anchor = (
        MSO_ANCHOR.MIDDLE
    )

    paragraph = frame.paragraphs[0]

    paragraph.text = text or "\uac1c\uc778 \uae30\ub3c4 \uc2dc\uac04"

    paragraph.alignment = (
        PP_ALIGN.CENTER
    )

    for run in paragraph.runs:
        run.font.size = Pt(30)
        run.font.color.rgb = RGBColor(
            255,
            255,
            255,
        )

    add_zoom_rule(
        slide,
        left=4.4,
        top=3.95,
        width=4.55,
    )

    return slide

def render_friday_zoom_block(
    prs: Presentation,
    block: WorshipBlock,
    *,
    bible_provider: BibleProvider | None = None,
    include_scripture_reference_slide: bool = False,
):
    if block.kind == BlockKind.BLANK:
        return add_blank_slide(prs)

    if block.kind == BlockKind.PRAYER_TOPICS:
        return add_prayer_topics_slide(
            prs,
            block.value.topics,
        )

    if block.kind == BlockKind.PERSONAL_PRAYER:
        return add_personal_prayer_slide(
            prs,
            block.value.text,
        )

    if block.kind == BlockKind.SERMON_TITLE:
        return add_zoom_sermon_title_slide(
            prs,
            block.value.title,
            block.value.scripture_reference,
        )

    if block.kind == BlockKind.SCRIPTURE:
        if bible_provider is None:
            raise MissingRenderDependencyError(
                "금요 Zoom SCRIPTURE 블록을 "
                "렌더링하려면 BibleProvider가 "
                "필요합니다."
            )

        from media_automation.bible.json_source import split_references

        slides = []
        for reference in split_references(block.value.reference):
            passage = bible_provider.get_passage(reference)
            if include_scripture_reference_slide:
                slides.append(add_zoom_scripture_reference_slide(prs, passage.reference))
            slides.extend(add_zoom_scripture_passage_slides(prs, passage))
        return slides

    raise ValueError(
        "금요 Zoom renderer가 "
        f"아직 지원하지 않는 블록입니다: "
        f"{block.kind}"
    )