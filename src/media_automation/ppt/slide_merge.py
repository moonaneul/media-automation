from __future__ import annotations

import platform
import subprocess
from pathlib import Path
from typing import Protocol

from pptx import Presentation


class SlideMerger(Protocol):
    def insert_all(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
    ) -> None:
        ...

    def insert_range(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
        start_slide: int,
        end_slide: int,
    ) -> None:
        ...


class PowerPointComSlideMerger:
    """
    Windows + Microsoft PowerPoint용 슬라이드 병합기.
    """

    def _require_windows(self) -> None:
        if platform.system() != "Windows":
            raise RuntimeError(
                "PowerPoint COM 슬라이드 병합은 "
                "Windows에서만 지원합니다."
            )

    def insert_all(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
    ) -> None:
        source = Path(source)

        source_prs = Presentation(source)
        slide_count = len(source_prs.slides)

        if slide_count == 0:
            return

        self.insert_range(
            destination,
            source,
            after_slide=after_slide,
            start_slide=1,
            end_slide=slide_count,
        )

    def insert_range(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
        start_slide: int,
        end_slide: int,
    ) -> None:
        self._require_windows()

        destination = Path(destination).resolve()
        source = Path(source).resolve()

        if start_slide < 1:
            raise ValueError(
                "start_slide은 1 이상이어야 합니다."
            )

        if end_slide < start_slide:
            raise ValueError(
                "end_slide은 start_slide보다 "
                "작을 수 없습니다."
            )

        if not destination.exists():
            raise FileNotFoundError(
                f"대상 PPT가 없습니다: {destination}"
            )

        if not source.exists():
            raise FileNotFoundError(
                f"삽입할 PPT가 없습니다: {source}"
            )

        try:
            import win32com.client
        except ImportError as error:
            raise RuntimeError(
                "Windows PowerPoint 병합에는 "
                "pywin32가 필요합니다."
            ) from error

        powerpoint = None
        presentation = None

        try:
            powerpoint = (
                win32com.client.DispatchEx(
                    "PowerPoint.Application"
                )
            )

            presentation = (
                powerpoint.Presentations.Open(
                    str(destination),
                    False,
                    False,
                    False,
                )
            )

            presentation.Slides.InsertFromFile(
                str(source),
                after_slide,
                start_slide,
                end_slide,
            )

            presentation.Save()

        finally:
            if presentation is not None:
                presentation.Close()

            if powerpoint is not None:
                powerpoint.Quit()


class MacPowerPointAppleScriptSlideMerger:
    """
    macOS + Microsoft PowerPoint용 슬라이드 병합기.

    PowerPoint 자체의 copy/paste 기능을 AppleScript로 호출해서
    python-pptx가 원본 악보 슬라이드를 재구성하지 않도록 한다.
    """

    def _require_macos(self) -> None:
        if platform.system() != "Darwin":
            raise RuntimeError(
                "Mac PowerPoint 슬라이드 병합은 "
                "macOS에서만 지원합니다."
            )

    def insert_all(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
    ) -> None:
        source = Path(source)

        source_prs = Presentation(source)
        slide_count = len(source_prs.slides)

        if slide_count == 0:
            return

        self.insert_range(
            destination,
            source,
            after_slide=after_slide,
            start_slide=1,
            end_slide=slide_count,
        )

    def insert_range(
        self,
        destination: str | Path,
        source: str | Path,
        *,
        after_slide: int,
        start_slide: int,
        end_slide: int,
    ) -> None:
        self._require_macos()

        destination = Path(destination).resolve()
        source = Path(source).resolve()

        if start_slide < 1:
            raise ValueError(
                "start_slide은 1 이상이어야 합니다."
            )

        if end_slide < start_slide:
            raise ValueError(
                "end_slide은 start_slide보다 "
                "작을 수 없습니다."
            )

        # 현재 실제 예배 PPT는 PRE_SERVICE 뒤에 찬양이 들어가므로
        # after_slide >= 1이다.
        # 0번 위치 삽입은 별도 구현 전까지 명시적으로 막는다.
        if after_slide < 1:
            raise ValueError(
                "Mac PowerPoint 병합은 현재 "
                "첫 슬라이드 앞 삽입을 지원하지 않습니다."
            )

        if not destination.exists():
            raise FileNotFoundError(
                f"대상 PPT가 없습니다: {destination}"
            )

        if not source.exists():
            raise FileNotFoundError(
                f"삽입할 PPT가 없습니다: {source}"
            )

        # Mac PowerPoint에서는 두 번째 파일을 연 뒤 이전의
        # `active presentation` 객체 참조가 불안정해질 수 있다.
        # 따라서 열린 프레젠테이션을 파일명으로 다시 찾아서
        # copy/select/paste 한다.
        script = r'''
on run argv
    set destinationPath to item 1 of argv
    set sourcePath to item 2 of argv
    set afterSlideNumber to (item 3 of argv) as integer
    set startSlideNumber to (item 4 of argv) as integer
    set endSlideNumber to (item 5 of argv) as integer
    set destinationName to item 6 of argv
    set sourceName to item 7 of argv

    set destinationFile to POSIX file destinationPath as alias
    set sourceFile to POSIX file sourcePath as alias

    tell application "Microsoft PowerPoint"
        activate

        open sourceFile
        delay 0.2
        open destinationFile
        delay 0.2

        set insertionPoint to afterSlideNumber

        repeat with sourceSlideNumber from startSlideNumber to endSlideNumber
            set beforeCount to count slides of presentation destinationName

            copy object slide sourceSlideNumber of presentation sourceName
            select slide insertionPoint of presentation destinationName

            tell active window
                set view type to slide sorter view
                paste object its view
                set view type to normal view
            end tell

            delay 0.1

            set afterCount to count slides of presentation destinationName
            if afterCount is not (beforeCount + 1) then
                error "슬라이드 붙여넣기 후 대상 PPT의 슬라이드 수가 증가하지 않았습니다."
            end if

            set insertionPoint to insertionPoint + 1
        end repeat

        close presentation sourceName saving no
        close presentation destinationName saving yes
    end tell
end run
'''

        result = subprocess.run(
            [
                "osascript",
                "-e",
                script,
                str(destination),
                str(source),
                str(after_slide),
                str(start_slide),
                str(end_slide),
                destination.name,
                source.name,
            ],
            text=True,
            capture_output=True,
        )

        if result.returncode != 0:
            message = (
                result.stderr.strip()
                or result.stdout.strip()
                or "알 수 없는 AppleScript 오류"
            )

            raise RuntimeError(
                "Mac PowerPoint 슬라이드 병합에 "
                f"실패했습니다: {message}"
            )


def create_platform_slide_merger() -> SlideMerger:
    system = platform.system()

    if system == "Windows":
        return PowerPointComSlideMerger()

    if system == "Darwin":
        return MacPowerPointAppleScriptSlideMerger()

    raise RuntimeError(
        f"지원하지 않는 운영체제입니다: {system}"
    )
