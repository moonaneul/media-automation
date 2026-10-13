"""Original UI adapter to the supported media_automation PPT engine."""
from pathlib import Path
import json
import os
import re
import uuid
import hashlib
import unicodedata
from zipfile import BadZipFile
from lxml.etree import XMLSyntaxError
from pptx.exc import PackageNotFoundError
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from media_automation.ppt.prayer_audio import embed_prayer_audio
from media_automation.bible.provider import BiblePassage, BibleVerse
from media_automation.ppt.base import create_4x3_presentation, create_16x9_presentation, add_blank_slide, add_centered_text_slide, add_sermon_title_slide
from media_automation.ppt.scripture import add_scripture_passage_slides
from media_automation.ppt.friday_zoom import add_zoom_scripture_passage_slides, add_zoom_centered_text_slide, add_zoom_sermon_title_slide
from media_automation.ppt.slide_merge import OpenXmlSlideMerger
from media_automation.ppt.hyperlinks import remove_presentation_hyperlinks
from media_automation.ppt.normalization import normalize_layout_ids
from media_automation.ppt.sunday_typography import apply_sunday_typography, VERSION as SUNDAY_TYPOGRAPHY_VERSION
from media_automation.ppt.friday_typography import apply_friday_typography, VERSION as FRIDAY_TYPOGRAPHY_VERSION
from .production_order import build_order, combine_friday_prayers
from .bible_connection import resolve
from .ppt_conversion import convert_original

BUILD_VERSION = 'official-web-engine-v6-invalid-score-report'
ROOT = Path(__file__).resolve().parents[4]

def compatible_original(filename, service, mode, slot):
    suffix = Path(filename).suffix.lower()
    if slot.startswith('배경음악'): return suffix == '.mp3'
    return suffix == '.mp4' if service == 'friday' and mode == 'zoom' else suffix in ('.ppt', '.pptx')

def _wrap_prayer_words(text, width=50):
    """Keep Korean words together within the existing 26pt prayer text box."""
    lines=[];line=''
    for word in text.split():
        candidate=(line+' '+word).strip()
        units=sum(2 if unicodedata.east_asian_width(c) in ('W','F') else 1 for c in candidate)
        if line and units>width:lines.append(line);line=word
        else:line=candidate
    return '\n'.join(lines+([line] if line else []))


def render_configuration(data):
    """Choices outside fields also affect the produced presentation."""
    return {key: data.get(key) for key in
            ('fridayMode', 'songChoices', 'songEditions', 'assetChoices', 'extraSongs',
             'scriptureIds', 'fridayNoticeGroups', 'noticePrayerGroups', 'prayerTopicIds')}


def current_input(store, body):
    current = store.get(body['date'], body['service'])
    if current['revision'] != body['revision']:
        raise ValueError('입력이 변경됐습니다. 최신 입력으로 제작하세요.')
    if current['data'].get('noticeDate') not in (None,body['date']):
        raise ValueError('안내문 날짜와 예배 날짜가 다릅니다.')
    return current['data']


def _asset(store, body, block):
    data = store.get(body['date'], body['service'])['data']
    compatible = lambda a: compatible_original(a['filename'], body['service'], data.get('fridayMode','zoom'), block['key'])
    attachments = {a['slot']: a for a in store.attachments(body['date'], body['service'])}
    if block['key'] in attachments:
        asset = attachments[block['key']]
        if not compatible(asset): return None
    else:
        title = re.sub(r'\s+', '', block.get('content') or '')
        data = store.get(body['date'], body['service'])['data']
        selected = data.get('songChoices', {}).get(block['key'])
        number = re.fullmatch(r'(?:새찬송가|구찬송가)?(\d+)장?', title)
        edition = data.get('songEditions', {}).get(block['key'])
        songs = [song for song in store.songs() if
            (song['id'] == selected if selected else
             song.get('number') == int(number[1]) and song.get('edition') == edition if number and edition else
             title in [re.sub(r'\s+', '', v) for v in [song['title'], *song['aliases']]])]
        chosen_asset = data.get('assetChoices', {}).get(block['key'])
        assets = [asset for song in songs for asset in song['assets'] if asset.get('eligible') and compatible(asset)
                  and (not chosen_asset or asset['digest'] == chosen_asset)]
        if len(assets) != 1: return None
        asset = assets[0]
    path = Path(store.path).with_suffix('.files') / asset['digest']
    if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != asset['digest']: return None
    suffix = Path(asset['filename']).suffix.lower()
    if suffix in ('.mp4', '.mp3'):
        # python-pptx derives package media extensions from the source name.
        # Content-addressed uploads have no suffix; retain the real format for Office.
        cache = path.parent / '.media-cache'
        cache.mkdir(exist_ok=True)
        named = cache / (asset['digest'] + suffix)
        try:
            os.link(path, named)
        except FileExistsError:
            if hashlib.sha256(named.read_bytes()).hexdigest() != asset['digest']:
                raise ValueError('미디어 검증 캐시가 원본과 다릅니다.')
        path = named
    if suffix == '.ppt': path = convert_original(path)
    return path, suffix, asset


def generate(store, body):
    data = current_input(store, body)
    fields = data.get('fields', {})
    for label in ('성경 본문', '설교 제목'):
        if fields.get(label, {}).get('state') != 'VALUE':
            raise ValueError(label + '을 해당 예배 안내로 입력하세요.')
    service = body['service']
    zoom = service == 'friday' and data.get('fridayMode', 'zoom') == 'zoom'
    order = build_order(body['groups'], fields, service)
    bible = resolve(order)
    if service == 'friday': order = combine_friday_prayers(order, bible, include_personal=zoom)
    directory = Path(store.path).with_suffix('.generated')
    directory.mkdir(parents=True, exist_ok=True)
    identifier = uuid.uuid4().hex
    path = directory / (identifier + '.pptx')
    template_key = 'friday-zoom' if zoom else 'friday-in-person' if service == 'friday' else service
    template_dir = Path(os.environ.get('MEDIA_WORSHIP_TEMPLATES', str(ROOT / 'assets/worship')))
    start = template_dir / (template_key + '-start.pptx')
    prs = create_16x9_presentation() if zoom else create_4x3_presentation()
    if start.is_file():
        reference = Presentation(start)
        prs.slide_width, prs.slide_height = reference.slide_width, reference.slide_height
        prs.save(path)
        OpenXmlSlideMerger().insert_all(path, start, after_slide=0)
        prs = Presentation(path)
    expected_ratio = 16/9 if zoom else 4/3
    if abs(prs.slide_width/prs.slide_height-expected_ratio) > .02:
        raise ValueError('예배별 시작 템플릿 화면 비율을 확인하세요.')
    mapping, unresolved, connections = [], [], {}
    if start.is_file():
        mapping.append(dict(kind='pre_service', key='예배 전 안내', start=1, end=len(prs.slides), status='original_inserted'))
        add_blank_slide(prs)
        mapping.append(dict(kind='blank', key='boundary:start', start=len(prs.slides), end=len(prs.slides), status='rendered'))
    else: unresolved.append('예배 전 안내:template')
    try:
        for block in order:
            before = len(prs.slides)
            kind, key = block['kind'], block['key']
            status = 'rendered'
            if kind == 'blank': add_blank_slide(prs)
            elif kind == 'song':
                attached = next((a for a in store.attachments(body['date'],service) if a['slot']==key), None)
                if attached and not compatible_original(attached['filename'], service, data.get('fridayMode','zoom'), key):
                    connections[key] = dict(status='incompatible_format', filename=attached['filename'])
                asset = _asset(store, body, block) if block.get('state') == 'VALUE' else None
                if asset:
                    source, suffix, metadata = asset
                    if not zoom and suffix in ('.ppt', '.pptx'):
                        try:
                            # Check the source before the merger can touch the output.
                            Presentation(source)
                        except (BadZipFile, PackageNotFoundError, XMLSyntaxError, KeyError, OSError):
                            status = 'missing_song_asset'
                            connections[key] = dict(metadata, status='invalid_ppt_package',
                                error='악보 원본을 정상 PPTX로 읽을 수 없습니다. 원본 파일을 확인하고 다시 첨부하세요.')
                        else:
                            prs.save(path)
                            try:
                                OpenXmlSlideMerger().insert_all(path, source, after_slide=before)
                                prs = Presentation(path)
                                status = 'original_inserted'
                                connections[key] = dict(metadata, status='connected')
                            except ValueError:
                                status = 'missing_song_asset'
                                connections[key] = dict(metadata, status='canvas_mismatch')
                    elif zoom and suffix == '.mp4':
                        slide = add_blank_slide(prs)
                        from media_automation.ppt.video_poster import video_poster
                        slide.shapes.add_movie(str(source), 0, 0, prs.slide_width, prs.slide_height, poster_frame_image=str(video_poster(source)), mime_type='video/mp4')
                        for condition in slide._element.iter(qn('p:cond')):
                            if condition.get('delay')=='indefinite':condition.set('delay','0')
                        status = 'original_inserted'
                        connections[key] = dict(metadata, status='connected', autoplay_configured=True, playback_review_pending=True)
                    else: status = 'missing_song_asset'
                else: status = 'missing_song_asset'
                if status == 'missing_song_asset':
                    add_blank_slide(prs)
                    if block.get('state') != 'VALUE':
                        status = 'input_pending'
                        unresolved.append(key + ':input')
                    else:
                        unresolved.append(key + ':score_asset')
            elif block.get('state') != 'VALUE':
                add_blank_slide(prs); unresolved.append(key + ':input')
            elif re.search('성경 본문|추가 말씀', key):
                result = bible.get(key, {})
                if not result.get('units'):
                    add_blank_slide(prs); status = 'missing_verified_bible_text'; unresolved.append(key + ':bible_text')
                else:
                    passages = {}
                    for unit in result['units']:
                        passages.setdefault(unit['passage_reference'], []).append(BibleVerse(unit['start_verse'], unit['text'], unit['end_verse']))
                    for reference, verses in passages.items():
                        passage = BiblePassage(reference, tuple(verses))
                        (add_zoom_scripture_passage_slides if zoom else add_scripture_passage_slides)(prs, passage)
                    status = 'bible_inserted'
                    if result.get('text_review_pending'): unresolved.append(key + ':bible_review')
            elif key == '설교 제목':
                (add_zoom_sermon_title_slide if zoom else add_sermon_title_slide)(prs, block['content'], block.get('subtitle', ''))
            else:
                text = '교회 소식' if key == '교회 소식' else block['content']
                if zoom and '기도' in key:
                    from media_automation.ppt.friday_zoom import add_prayer_topics_slide, add_personal_prayer_slide
                    if '개인 기도' in key:
                        lines=text.splitlines()
                        slide=add_personal_prayer_slide(prs, lines[0])
                        for personal_shape in slide.shapes:
                            if personal_shape.has_text_frame and personal_shape.text:
                                personal_shape.left=Inches(1.5);personal_shape.width=Inches(10.3)
                        if len(lines)>1:
                            from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
                            from pptx.util import Pt
                            shape=slide.shapes.add_textbox(Inches(3.3), Inches(4.2), Inches(6.7), Inches(1.2))
                            shape.name='개인 기도 보조 안내'
                            frame=shape.text_frame;frame.word_wrap=True;frame.vertical_anchor=MSO_ANCHOR.MIDDLE
                            frame.text='\n'.join(lines[1:])
                            for paragraph in frame.paragraphs:
                                paragraph.alignment=PP_ALIGN.CENTER
                                for run in paragraph.runs:run.font.size=Pt(20);run.font.color.rgb=RGBColor(255,255,255)
                    else:
                        topic=re.sub(r'^\d+\.\s*','',text)
                        slide=add_prayer_topics_slide(prs,[topic])
                        from pptx.util import Pt
                        for shape in slide.shapes:
                            if shape.has_text_frame and shape.name.startswith('FridayZoomPrayerTopic'):
                                shape.width=Inches(11.84)
                                # Keep the order's prayer numbering; one topic per screen.
                                shape.text_frame.paragraphs[0].runs[0].text=text.splitlines()[0]
                                for paragraph in shape.text_frame.paragraphs:
                                    paragraph.text=_wrap_prayer_words(paragraph.text, width=58)
                                    # Korean prayer text wraps at spaces, keeping words intact.
                                    paragraph._p.get_or_add_pPr().set('eaLnBrk', '0')
                                    for run in paragraph.runs:
                                        run.font.size=Pt(24 if len(text)>200 else 26)
                                        run.font.color.rgb=RGBColor(255,255,255)
                        if block.get('supporting_review_pending'):unresolved.append(block['supporting_key']+':bible_review')
                else:
                    (add_zoom_centered_text_slide if zoom else add_centered_text_slide)(prs, text)
            after = len(prs.slides)
            if after > before: mapping.append(dict(kind=kind, key=key, start=before+1, end=after, status=status))
        reference_applied=False
        if zoom:
            prs.save(path)
            from media_automation.ppt.friday_reference import apply_reference_layout
            reference_applied=apply_reference_layout(path,mapping,body['date'],Path(os.environ.get('MEDIA_FRIDAY_REFERENCE',str(ROOT/'assets/private/friday_reference'))))
            prs=Presentation(path)
            if not reference_applied:unresolved.append('금요 원본 디자인:template')
        background_music = {'inserted': [], 'actual_playback_verified': False}
        if zoom:
            roles = [('pre_service','배경음악 예배 준비',lambda e:e['kind']=='pre_service'),
                     ('before_sermon','배경음악 첫 기도',lambda e:'첫 기도' in e['key']),
                     ('after_sermon','배경음악 말씀 기도',lambda e:'말씀 관련' in e['key'] and '기도' in e['key']),
                     ('intercession','배경음악 공동체 기도',lambda e:'공동체' in e['key'] or '개인 기도' in e['key'])]
            for role, slot, predicate in roles:
                entries = [entry for entry in mapping if entry['kind'] != 'blank' and predicate(entry)]
                if not entries: continue
                asset = _asset(store, body, {'key':slot, 'content':''})
                if not asset:
                    music=Path(os.environ.get('MEDIA_FRIDAY_MUSIC',str(ROOT/'assets/private/friday_zoom')))
                    manifest=music/'music.json'
                    config=json.loads(manifest.read_text()) if manifest.is_file() else {}
                    entry=config.get(role,{})
                    source=music/entry.get('filename','missing')
                    if entry.get('confirmed') is True and source.is_file() and source.suffix=='.mp3' and source.parent==music:
                        digest=hashlib.sha256(source.read_bytes()).hexdigest()
                        if digest==entry.get('digest'):
                            asset=(source,'.mp3',{'digest':digest,'source':'confirmed_fixed_prayer_music'})
                if not asset or asset[1] != '.mp3':
                    unresolved.append(slot+':audio_asset'); continue
                source, _, metadata = asset
                # One audio object per contiguous prayer range; never play over a song.
                groups = []
                for entry in entries:
                    if groups and entry['start'] == groups[-1]['end']+1:groups[-1]['end']=entry['end']
                    else:groups.append(dict(entry))
                for entry in groups:
                    embed_prayer_audio(prs, source, slide_number=entry['start'], slide_count=entry['end']-entry['start']+1)
                    background_music['inserted'].append(dict(role=role,page=entry['start'],end=entry['end'],digest=metadata['digest']))
        prs.save(path)
        remove_presentation_hyperlinks(path)
        if service in ('sunday', 'wednesday'):
            apply_sunday_typography(path, mapping)
        if service == 'friday' and not reference_applied:
            apply_friday_typography(path, mapping, zoom=zoom)
        normalize_layout_ids(path)
        report = dict(build_version=BUILD_VERSION, typography_version=SUNDAY_TYPOGRAPHY_VERSION if service in ('sunday', 'wednesday') else FRIDAY_TYPOGRAPHY_VERSION if service == 'friday' else None, date=body['date'], service=service,
            fields=fields, attachments=store.attachments(body['date'],service), format='pptx', slide_count=len(prs.slides), mapping=mapping,
            unresolved=unresolved, connections=connections, complete=False, visual_review_pending=True,
            media_playback_review_pending=zoom, reference_layout_applied=reference_applied, background_music=background_music,
            render_configuration=render_configuration(data))
        (directory / (identifier + '.json')).write_text(json.dumps(report, ensure_ascii=False))
        return dict(engine_version=BUILD_VERSION, typography_version=SUNDAY_TYPOGRAPHY_VERSION if service in ('sunday', 'wednesday') else FRIDAY_TYPOGRAPHY_VERSION if service == 'friday' else None, file_id=identifier, filename=f"{body['date']}-{service}-검수전.pptx",
            slide_count=len(prs.slides), unresolved=unresolved, complete=False, visual_review_pending=True)
    except Exception:
        path.unlink(missing_ok=True)
        raise
