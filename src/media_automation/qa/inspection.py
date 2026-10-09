"""Read-only structural inspection of PowerPoint and folded bulletin PDF outputs.

No claim of visual, copyright, or playback verification is made.
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
import json
import re
from typing import Any

from pptx import Presentation
from pypdf import PdfReader

_URL = re.compile(r"https?://[^\s<>]+", re.I)
_REF = re.compile(r"(?P<book>[가-힣]{1,6})\s*(?P<chapter>\d+):(?P<start>\d+)(?:\s*[~\-]\s*(?P<end>\d+))?")
_VERSE = re.compile(r"^\s*(\d{1,3})\s*[.．]")
_A4_PT = (841.89, 595.28)


def issue(severity: str, code: str, message: str, *, slide: int | None = None) -> dict:
    item = {"severity": severity, "code": code, "message": message}
    if slide is not None:
        item["slide"] = slide
    return item


def _pptx(path: Path, service: str, expected_reference: str | None) -> tuple[dict, list[dict]]:
    prs = Presentation(str(path))
    items: list[dict] = []
    media = []
    visible_text: list[tuple[int, str]] = []
    for number, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    texts.extend(cell.text for cell in row.cells)
        combined = "\n".join(texts)
        visible_text.append((number, combined))
        for value in _URL.findall(combined):
            items.append(issue("review", "external_url", f"URL 확인 필요: {value}", slide=number))
        if "세례" in combined:
            items.append(issue("error", "baptism_word", "표시 텍스트에 '세례'가 남아 있습니다.", slide=number))
        for rel in slide.part.rels.values():
            if rel.is_external:
                target = str(rel.target_ref)
                if target.startswith(("http://", "https://")):
                    items.append(issue("review", "external_relationship", f"외부 연결 확인 필요: {target}", slide=number))
            else:
                target = str(rel.target_ref)
                if "/media/" in target:
                    media.append({"slide": number, "target": target})

    if expected_reference:
        expected = _REF.fullmatch(expected_reference.strip())
        if not expected:
            items.append(issue("review", "unparsed_reference", "본문 범위를 해석하지 못하여 수동 대조가 필요합니다."))
        else:
            start = int(expected.group("start"))
            end = int(expected.group("end") or start)
            if end < start or end - start > 200:
                items.append(issue("error", "invalid_reference", "본문 절 범위가 유효하지 않습니다."))
            else:
                matching: list[tuple[int, int]] = []
                for n, txt in visible_text:
                    if expected_reference.strip() not in txt:
                        continue
                    values = [int(m.group(1)) for line in txt.splitlines() if (m := _VERSE.match(line))]
                    matching.extend((n, v) for v in values)
                if not matching:
                    items.append(issue("review", "reference_not_detected", "기대 본문·절 번호를 텍스트에서 찾지 못했습니다. 이미지 본문은 별도 확인이 필요합니다."))
                else:
                    counts = {v: sum(x == v for _, x in matching) for v in range(start, end + 1)}
                    for verse, count in counts.items():
                        if count == 0:
                            items.append(issue("error", "verse_missing", f"{verse}절 번호가 발견되지 않았습니다."))
                        elif count > 1:
                            items.append(issue("review", "verse_duplicate", f"{verse}절 번호가 {count}회 나타납니다."))
                    if service in {"wednesday", "sunday"}:
                        by_slide = {}
                        for n, v in matching:
                            by_slide.setdefault(n, []).append(v)
                        for n, verses in by_slide.items():
                            if len(verses) > 1:
                                items.append(issue("review", "multi_verse_slide", "한 화면에 여러 절 번호가 있습니다.", slide=n))
    items.append(issue("review", "visual_review_required", "전체 슬라이드 렌더링, 이미지 안 성경 및 악보 잘림은 사람이 확인해야 합니다."))
    if media:
        items.append(issue("review", "media_playback_required", "내장 미디어의 실제 재생과 사용 허락은 별도로 확인해야 합니다."))
    return {"slides": len(prs.slides), "media_relationships": len(media),
            "width_emu": prs.slide_width, "height_emu": prs.slide_height}, items


def _pdf(path: Path) -> tuple[dict, list[dict]]:
    reader = PdfReader(str(path))
    items = []
    count = len(reader.pages)
    if count != 2:
        items.append(issue("error", "folded_page_count", f"접지 주보는 PDF 2페이지여야 합니다. 실제 {count}페이지."))
    for n, page in enumerate(reader.pages, 1):
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        if abs(width - _A4_PT[0]) > 4 or abs(height - _A4_PT[1]) > 4:
            items.append(issue("error", "page_size", f"{n}페이지가 A4 가로가 아닙니다: {width:.1f} x {height:.1f}pt"))
        text = page.extract_text() or ""
        if "세례" in text:
            items.append(issue("error", "baptism_word", f"PDF {n}페이지에 '세례'가 있습니다."))
        for url in _URL.findall(text):
            items.append(issue("review", "external_url", f"PDF {n}페이지 URL 확인 필요: {url}"))
    items.append(issue("review", "folding_visual_review", "논리 쪽 [4|1], [2|3] 배치와 실제 렌더링·양면 접지는 별도 확인해야 합니다."))
    return {"pages": count}, items


def inspect_file(path: str | Path, *, service: str = "wednesday", expected_reference: str | None = None) -> dict[str, Any]:
    """Inspect without changing source; unknown or visual checks remain review."""
    file_path = Path(path)
    if service not in {"wednesday", "sunday", "friday_zoom", "friday_in_person", "bulletin"}:
        raise ValueError("Unsupported service type")
    if not file_path.is_file():
        raise FileNotFoundError(file_path)
    suffix = file_path.suffix.lower()
    if suffix == ".pptx":
        stats, issues = _pptx(file_path, service, expected_reference)
    elif suffix == ".pdf" and service == "bulletin":
        stats, issues = _pdf(file_path)
    else:
        raise ValueError("Expected .pptx, or a .pdf with service='bulletin'")
    return {"schema_version": 1, "file": str(file_path), "service": service,
            "inspected_at": datetime.now(timezone.utc).isoformat(),
            "stats": stats, "issues": issues,
            "summary": {level: sum(x["severity"] == level for x in issues) for level in ("error", "review", "pass")},
            "rendered": False, "played_media": False, "checked_permissions": False}


def save_report(report: dict, target: str | Path) -> Path:
    destination = Path(target)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return destination
